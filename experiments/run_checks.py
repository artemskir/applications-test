"""Offline XLSX gap experiments; not a production normalizer or input adapter."""
import argparse
import base64
from collections import Counter
from datetime import datetime, timezone
import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import time
import tracemalloc
from unittest.mock import patch
import warnings
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_STORED, BadZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from fixtures.scenarios import CASES, HEADERS
from tools.generate_data import clean_rows, generate, write_workbook
from research.run_experiments import candidate_shape, strict_calendar_probe, phone_probe, date_probe, name_probe
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.worksheet.datavalidation import DataValidation

NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
SHEET = 'xl/worksheets/sheet1.xml'


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def dump(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')


def container_guard(path):
    """Experimental subset only; not full OPC validation or archive resource limits."""
    try:
        with ZipFile(path) as archive:
            names = [item.filename for item in archive.infolist()]
            if len(names) != len(set(names)):
                return 'duplicate_parts'
            if archive.testzip() is not None:
                return 'crc_error'
            required = {'[Content_Types].xml', '_rels/.rels', 'xl/workbook.xml', SHEET}
            if not required.issubset(names):
                return 'missing_parts'
            for name in names:
                if name.endswith('.xml') or name.endswith('.rels'):
                    ET.fromstring(archive.read(name))
    except (BadZipFile, ET.ParseError, RuntimeError, OSError):
        return 'invalid_container'
    return 'ok'


def rewrite_zip(source, target, transform=None):
    with ZipFile(source) as old, ZipFile(target, 'w', compression=ZIP_STORED) as new:
        for info in old.infolist():
            value = old.read(info.filename)
            if transform:
                value = transform(info.filename, value)
            new.writestr(info.filename, value)


def rows_read(path, read_only=False, reset=False):
    wb = load_workbook(path, read_only=read_only, data_only=False)
    try:
        if reset:
            wb.active.reset_dimensions()
        return list(wb.active.values)
    finally:
        wb.close()


def first_model(rows):
    """Small identity-phone model for metamorphic checks, never used to clean files."""
    seen, kept = set(), []
    for row_id, phone in rows:
        if phone is None or phone not in seen:
            kept.append(row_id)
        if phone is not None:
            seen.add(phone)
    return kept


def gap_coverage():
    """Inventory every domain table row; descriptions are never invented as literals."""
    records, domain = [], None
    domains = {'## Телефоны': 'Телефон', '## Даты': 'Дата заявки', '## Имена': 'Имя'}
    for line_number, line in enumerate((ROOT / 'experiments/gap-research.md').read_text(encoding='utf-8').splitlines(), 1):
        if line.startswith('## '):
            domain = domains.get(line)
        if not domain or not line.startswith('|'):
            continue
        cells = [cell.strip() for cell in line.strip('|').split('|')]
        if len(cells) < 3 or cells[0] == 'Вход' or cells[0].startswith('---'):
            continue
        literals = re.findall(r'`([^`]*)`', cells[0])
        description = not literals or cells[0].startswith('serial ') or 'либо' in cells[0]
        record = {'source_line': line_number, 'field': domain, 'source_input': cells[0],
                  'source_expectation': cells[1], 'kind': 'description' if description else 'concrete_text', 'probes': []}
        if not description:
            probe = phone_probe if domain == 'Телефон' else date_probe if domain == 'Дата заявки' else name_probe
            for value in literals:
                result = probe(value)
                record['probes'].append({'input': value, 'observation': result,
                                         'exact_catalog_case_ids': [c['id'] for c in CASES if c['field'] == domain and c['input'] == value]})
        records.append(record)
    return records


def worker(destination, ready, gate, seed):
    Path(ready).write_text('ready', encoding='ascii')
    deadline = time.monotonic() + 15
    while not Path(gate).exists():
        if time.monotonic() > deadline:
            return 3
        time.sleep(0.02)
    try:
        generate(destination, mode='clean', rows=100, seed=seed)
        return 0
    except FileExistsError:
        return 2


def run(stage, scale_rows):
    checks, observations = [], {}
    def check(name, condition, detail=None):
        checks.append({'name': name, 'passed': bool(condition), 'detail': detail})

    baseline = stage / 'baseline.xlsx'
    write_workbook(baseline, clean_rows(48, 19))
    baseline_hash = digest(baseline)
    # X01: a real worksheet carries all 48 records but lies about its dimensions.
    truncated = stage / 'bad_dimension.xlsx'
    def shrink(name, value):
        return re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', value) if name == SHEET else value
    rewrite_zip(baseline, truncated, shrink)
    with ZipFile(truncated) as archive:
        xml_count = len(ET.fromstring(archive.read(SHEET)).findall('m:sheetData/m:row', NS))
    normal_count = len(rows_read(truncated))
    streaming_count = len(rows_read(truncated, True))
    reset_count = len(rows_read(truncated, True, True))
    check('X01-actual-rows-ignore-dimension', xml_count == normal_count == reset_count == 49,
          {'xml': xml_count, 'normal': normal_count, 'read_only_default': streaming_count, 'read_only_reset': reset_count})
    check('X01-reproduce-default-streaming-loss', streaming_count == 1)

    duplicate = stage / 'duplicate_parts.xlsx'
    shutil.copyfile(baseline, duplicate)
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        with ZipFile(duplicate, 'a') as archive:
            archive.writestr(SHEET, b'<different/>')
    check('X02-duplicate-parts-rejected', container_guard(duplicate) == 'duplicate_parts')

    corrupt = stage / 'bad_crc.xlsx'
    rewrite_zip(baseline, corrupt)
    with ZipFile(corrupt) as archive:
        info = archive.getinfo(SHEET)
    payload = bytearray(corrupt.read_bytes())
    header = struct.unpack_from('<IHHHHHIIIHH', payload, info.header_offset)
    data_start = info.header_offset + 30 + header[-2] + header[-1]
    payload[data_start + info.file_size - 20] ^= 1
    corrupt.write_bytes(payload)
    check('X03-late-payload-crc-detected', container_guard(corrupt) in {'crc_error', 'invalid_container'})
    foreign = stage / 'foreign.zip.xlsx'
    with ZipFile(foreign, 'w') as archive:
        archive.writestr('readme.txt', 'Not an XLSX')
    check('foreign-zip-not-xlsx', container_guard(foreign) == 'missing_parts')

    # X06-X11 / X15-X16: real typed values and metadata, without value guessing.
    metadata = stage / 'metadata.xlsx'
    wb = load_workbook(baseline)
    ws = wb.active
    ws['B2'] = 'Позвонить'
    ws['B2'].hyperlink = 'tel:+79174005219'
    ws['B2'].hyperlink.tooltip = 'Телефон только в метаданных'
    ws['C3'] = True
    ws['C4'] = 'ошибка'
    dv = DataValidation(type='date', operator='between', formula1='1', formula2='999999')
    ws.add_data_validation(dv)
    dv.add(ws['C4'])
    ws['C5'] = '#VALUE!'
    ws['C5'].data_type = 'e'
    ws['D6'] = '=HYPERLINK("https://example.invalid","сайт")'
    ws['D6'].data_type = 's'
    ws['D7'] = '=1+1'
    ws['B8'] = 89174005219
    ws['B9'] = '+7\u202e9174005219\u202c'
    ws.merge_cells('B10:B11')
    ws['A12'] = CellRichText('Иван ', TextBlock(InlineFont(b=True), 'Иванов'))
    ws.protection.sheet = True
    wb.save(metadata)
    wb.close()
    wb = load_workbook(metadata, data_only=False, rich_text=True)
    try:
        ws = wb.active
        check('X07-hyperlink-value-not-a-phone', ws['B2'].value == 'Позвонить' and candidate_shape(ws['B2'].value) is None)
        check('X07-hyperlink-metadata-preserved', ws['B2'].hyperlink.target == 'tel:+79174005219')
        check('X08-bool-is-not-date', type(ws['C3'].value) is bool and ws['C3'].data_type == 'b')
        check('X10-validation-does-not-fix-value', ws['C4'].value == 'ошибка' and len(ws.data_validations.dataValidation) == 1)
        check('X11-error-isolated-other-rows-readable', ws['C5'].data_type == 'e' and ws.max_row == 49 and ws['A49'].value is not None)
        check('X06-formula-looking-text-safe', ws['D6'].data_type == 's' and ws['D6'].value.startswith('=HYPERLINK'))
        check('formula-typed-value-preserved', ws['D7'].data_type == 'f' and ws['D7'].value == '=1+1')
        check('X15-numeric-phone-type-preserved', ws['B8'].data_type == 'n' and ws['B8'].value == 89174005219)
        check('X09-bidi-preserved-and-not-rescued', ws['B9'].value == '+7\u202e9174005219\u202c' and candidate_shape(ws['B9'].value) is None)
        check('X16-merged-not-forward-filled', ws['B11'].value is None and 'B10:B11' in str(ws.merged_cells))
        check('rich-text-is-not-plain-str', isinstance(ws['A12'].value, CellRichText) and str(ws['A12'].value) == 'Иван Иванов')
        check('sheet-protection-not-encryption', ws.protection.sheet and ws['A49'].value is not None)
        observations['hyperlink'] = {'value': ws['B2'].value, 'target': ws['B2'].hyperlink.target, 'tooltip': ws['B2'].hyperlink.tooltip}
    finally:
        wb.close()
    wb = load_workbook(metadata, data_only=True)
    try:
        check('data-only-loses-formula-not-literal', wb.active['D7'].value is None and wb.active['D6'].value.startswith('=HYPERLINK'))
    finally:
        wb.close()
    original = metadata.read_bytes()
    encoded = base64.b64encode(original).decode('ascii')
    chunk_book = stage / 'metadata_byte_audit.xlsx'
    wb = Workbook()
    wb.active.title = 'Аудит'
    wb.active.append(['index', 'payload'])
    for index, start in enumerate(range(0, len(encoded), 30000)):
        wb.active.append([index, encoded[start:start + 30000]])
    wb.save(chunk_book)
    wb.close()
    wb = load_workbook(chunk_book)
    try:
        restored_bytes = base64.b64decode(''.join(row[1] for row in list(wb.active.values)[1:]), validate=True)
    finally:
        wb.close()
    check('metadata-entire-source-byte-roundtrip', restored_bytes == original)
    restored = stage / 'metadata_restored.xlsx'
    restored.write_bytes(restored_bytes)
    wb = load_workbook(restored, rich_text=True)
    try:
        check('restored-hyperlink-and-richtext-readable', wb.active['B2'].hyperlink.target == 'tel:+79174005219' and isinstance(wb.active['A12'].value, CellRichText))
    finally:
        wb.close()
    for label, value in [('illegal-xml-control', 'Иван\x0bИванов'),
                         ('timezone-aware-native-date', datetime(2026, 9, 6, tzinfo=timezone.utc))]:
        path = stage / (label + '.xlsx')
        try:
            write_workbook(path, [[value, '+79174005219', '2026-09-06', 'сайт']])
        except Exception as exc:
            observations[label] = {'exception': type(exc).__name__, 'partial_file_exists': path.exists()}
            check(label + '-refused', type(exc).__name__ in {'IllegalCharacterError', 'ValueError'} and not path.exists())
        else:
            check(label + '-refused', False)

    # X05: failure injection into the actual generator, rather than a toy writer.
    failed = stage / 'save_failed'
    def fail_save(self, filename):
        Path(filename).write_bytes(b'partial ZIP')
        raise OSError('Injected disk error')
    with patch.object(Workbook, 'save', fail_save):
        try:
            generate(failed, mode='clean', rows=2)
        except OSError:
            pass
        else:
            raise AssertionError('Expected injected save error')
    check('X05-save-failure-no-final-output', not failed.exists())
    check('X05-save-failure-stage-cleaned', not list(stage.glob('.fixture-stage-*')))
    existing = stage / 'existing'
    existing.mkdir()
    marker = existing / 'important.txt'
    marker.write_bytes(b'original')
    try:
        generate(existing)
    except FileExistsError:
        pass
    check('X12-existing-result-untouched', marker.read_bytes() == b'original' and list(existing.iterdir()) == [marker])
    # Native process death bypasses finally: no final directory; orphan stage remains.
    crash_parent = stage / 'crash'
    crash_parent.mkdir()
    crash_code = ('from pathlib import Path; import os, sys; from openpyxl import Workbook; '
                  'from tools.generate_data import generate; '
                  'exec("def crash(self, filename):\\n Path(filename).write_bytes(b\'partial\')\\n os._exit(9)"); '
                  'Workbook.save=crash; generate(Path(sys.argv[1]), mode="clean", rows=2)')
    result = subprocess.run([sys.executable, '-c', crash_code, str(crash_parent / 'final')], cwd=ROOT,
                            capture_output=True, timeout=30,
                            creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
    check('X05-killed-process-no-final-output', result.returncode == 9 and not (crash_parent / 'final').exists(), result.stderr.decode(errors='replace'))
    observations['hard_crash'] = {'returncode': result.returncode, 'orphan_stages': [p.name for p in crash_parent.glob('.fixture-stage-*')]}

    # X04: real competing publishers on the current Windows environment.
    race = stage / 'race'
    race.mkdir()
    gate, target = race / 'gate', race / 'final'
    processes = []
    try:
        for seed in (31, 32):
            processes.append(subprocess.Popen([sys.executable, str(Path(__file__).resolve()), '--worker',
                                               str(target), str(race / f'ready-{seed}'), str(gate), str(seed)],
                                              cwd=ROOT, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                              creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0))
        deadline = time.monotonic() + 15
        while not all((race / f'ready-{seed}').exists() for seed in (31, 32)):
            if time.monotonic() > deadline:
                raise TimeoutError('Race workers not ready')
            time.sleep(0.02)
        gate.write_text('go', encoding='ascii')
        errors = [process.communicate(timeout=30)[1].decode(errors='replace') for process in processes]
        codes = [process.returncode for process in processes]
    finally:
        for process in processes:
            if process.poll() is None:
                process.kill()
                process.communicate()
    check('X04-exactly-one-publisher', sorted(codes) == [0, 2], {'codes': codes, 'stderr': errors})
    manifest = json.loads((target / 'manifest.json').read_text(encoding='utf-8'))
    check('X04-winning-manifest-matches-file', all(digest(target / item['file']) == item['sha256'] for item in manifest['files']))
    check('X04-no-mixed-generation', rows_read(target / 'clean.xlsx')[1:] == [tuple(r) for r in clean_rows(100, manifest['seed'])])
    check('X04-race-stages-cleaned', not list(race.glob('.fixture-stage-*')))

    # Gap text examples: literal oracles independent of parse-library suggestions.
    phone_oracles = {'8 (812) 123-45-67': '+78121234567', '8 (800) 707-46-19': '+78007074619',
                     '+ 7 (495) 123-45-67': '+74951234567', '8 (34786) 2-34-56': '+73478623456',
                     '０９１７４００５２１９': None, '8 (917) 40О-52-19': None,
                     '89174005219.0': None, '++79174005219': None}
    observations['phone_candidates'] = [{'input': value, 'expected': expected, 'observed': candidate_shape(value)} for value, expected in phone_oracles.items()]
    for index, record in enumerate(observations['phone_candidates'], 1):
        check(f'gap-phone-candidate-{index}', record['observed'] == record['expected'], record)
    date_oracles = {'09/03/2026': '2026-03-09', '2026-9-6': '2026-09-06', '29.02.2000': '2000-02-29',
                    '29.02.1900': None, '2026-09-06T18:00': '2026-09-06',
                    '2026-09-06T18:00+03:00': '2026-09-06', '2026-09-06T23:30:00-12:00': '2026-09-06',
                    '2026-09-06T24:00': None, '2026-09-06T18:00+03:60': None, '2026-10-01': '2026-10-01'}
    observations['dates'] = [{'input': value, 'expected': expected, 'observed': strict_calendar_probe(value)} for value, expected in date_oracles.items()]
    for index, record in enumerate(observations['dates'], 1):
        check(f'gap-calendar-{index}', record['observed']['value'] == record['expected'], record)
    check('ISO-time-context-preserved-in-probe', strict_calendar_probe('2026-09-06T18:00+03:00').get('time_context') == '18:00+03:00')
    check('UTC-original-example-does-not-cross-day', datetime.fromisoformat('2026-09-06T00:30:00-12:00').astimezone(timezone.utc).day == 6)
    check('UTC-corrected-example-crosses-day', datetime.fromisoformat('2026-09-06T23:30:00-12:00').astimezone(timezone.utc).day == 7)
    hours = [hour for hour in range(24) if re.sub(r'\D', '', f'8 917 400 52 19 (после {hour:02}:00)') != '89174005219']
    check('all-24-hours-defeat-whole-cell-digit-strip', hours == list(range(24)))

    # X13, X17, X18: limited model, no automatic promotion to production tests.
    unique = [('r2', '+79174005219'), ('r3', '+79174005220')]
    check('X17-unique-permutation-same-identities', set(first_model(unique)) == set(first_model(unique[::-1])))
    dup = [('r2', '+79174005219'), ('r3', '+79174005219'), ('r4', None), ('r5', None)]
    check('X18-duplicate-permutation-changes-first', first_model(dup) == ['r2', 'r4', 'r5'] and first_model(dup[::-1]) == ['r5', 'r4', 'r3'])
    first_bad = [('r2-invalid-date', '+79174005219'), ('r3-valid-date', '+79174005219')]
    check('X13-first-record-not-quality-ranked-model', first_model(first_bad) == ['r2-invalid-date'])
    raw_ids, lost_but_same_count = ['r2', 'r3', 'r4'], ['r2', 'r2', 'r4']
    check('inventory-detects-loss-with-equal-count', len(raw_ids) == len(lost_but_same_count) and Counter(raw_ids) != Counter(lost_but_same_count))
    check('raw-source-bytes-unchanged', digest(baseline) == baseline_hash)

    # Explicit manual mutations: catch errors with literal independent examples.
    mutants = [
        {'mutation': 'lstrip(8) instead of removing one prefix', 'caught': '+7' + '88121234567'.lstrip('8') != '+78121234567'},
        {'mutation': 'DMY interpreted as MDY', 'caught': datetime.strptime('09/03/2026', '%m/%d/%Y').date().isoformat() != '2026-03-09'},
        {'mutation': 'UTC conversion before date extraction', 'caught': datetime.fromisoformat('2026-09-06T23:30:00-12:00').astimezone(timezone.utc).date().isoformat() != '2026-09-06'},
        {'mutation': 'last duplicate selected', 'caught': list(reversed(first_model(list(reversed(first_bad))))) != ['r2-invalid-date']},
        {'mutation': 'same-size lost row replaced by another row', 'caught': Counter(raw_ids) != Counter(lost_but_same_count)},
    ]
    source = (ROOT / 'tools/generate_data.py').read_text(encoding='utf-8')
    for label, before, after, value in [
        ('allow silent truncation', 'len(actual_value) > 32767', 'len(actual_value) > 32768', 'А' * 32768),
        ('allow formula inference', 'cell.data_type = "s"  # Preserve text beginning with =, +, -, @.', 'pass', '=1+1'),
        ('allow nonfinite blank', 'not math.isfinite(actual_value)', 'False', float('nan')),
    ]:
        if source.count(before) != 1:
            raise AssertionError('Mutation target must be unique: ' + label)
        namespace = {'__file__': str(ROOT / 'tools/generate_data.py'), '__name__': 'experimental_mutant'}
        exec(compile(source.replace(before, after), '<experimental_mutant>', 'exec'), namespace)
        path = stage / ('mutant_' + str(len(mutants)) + '.xlsx')
        namespace['write_workbook'](path, [[value, '+79174005219', '2026-09-06', 'сайт']])
        wb = load_workbook(path, data_only=False)
        try:
            cell = wb.active['A2']
            if label == 'allow nonfinite blank':
                caught = cell.value is None
            elif label == 'allow formula inference':
                caught = cell.data_type != 's'
            else:
                caught = isinstance(cell.value, str) and len(cell.value) == 32767 and cell.value != value
        finally:
            wb.close()
        mutants.append({'mutation': label, 'caught': caught})
    observations['manual_mutations'] = mutants
    check('targeted-manual-mutants-detected', all(m['caught'] for m in mutants), mutants)

    # 100k is opt-in, timing includes only fixture generation and read-only counting.
    started = time.perf_counter()
    scale = stage / 'scale'
    generate(scale, mode='clean', rows=scale_rows)
    write_seconds = time.perf_counter() - started
    started = time.perf_counter()
    tracemalloc.start()
    wb = load_workbook(scale / 'clean.xlsx', read_only=True, data_only=False)
    try:
        count = sum(1 for _ in wb.active.iter_rows(values_only=True)) - 1
    finally:
        wb.close()
    _, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    observations['scale'] = {'rows': scale_rows, 'read_rows': count, 'write_seconds': round(write_seconds, 3),
                             'read_seconds_with_tracemalloc': round(time.perf_counter() - started, 3),
                             'reader_peak_python_bytes': peak, 'xlsx_bytes': (scale / 'clean.xlsx').stat().st_size,
                             'limitation': 'Not process RSS; excludes writer and future normalizer/audit memory'}
    check('scale-all-generated-rows-read', count == scale_rows)
    check('catalog-approved-new-cases-present', all(any(c['id'] == case_id for c in CASES) for case_id in ['P54', 'P57', 'D59', 'D61']))
    coverage = gap_coverage()
    dump(stage / 'gap_coverage.json', coverage)
    check('gap-all-domain-table-rows-inventoried', len(coverage) == 44)
    observations['gap_coverage'] = {'table_rows': len(coverage), 'kinds': dict(Counter(c['kind'] for c in coverage)),
                                    'concrete_probes': sum(len(c['probes']) for c in coverage),
                                    'limitation': 'Library observations, not assertions that every Research expectation is correct'}
    results = {'scope': 'Library / actual fixture generator / explicitly marked models; no production normalizer',
               'versions': {'python': sys.version, 'platform': sys.platform, 'openpyxl': importlib.metadata.version('openpyxl')},
               'source_sha256': {str(p.relative_to(ROOT)): digest(p) for p in [ROOT / 'experiments/gap-research.md', ROOT / 'experiments/additional-research.md', ROOT / 'docs/SPECIFICATION.md', ROOT / 'experiments/PLAN.md']},
               'checks': checks, 'observations': observations,
               'passed': sum(c['passed'] for c in checks), 'failed': sum(not c['passed'] for c in checks),
               'not_tested': ['Production four-sheet output and restoration CLI (absent)', 'Actual encrypted XLSX',
                              'Zip bomb limits / complete OPC relationship validation', 'Alternative reader engines',
                              'Real filesystem disk-full/permission failures; save error injected', 'POSIX publication race semantics']}
    dump(stage / 'results.json', results)
    inventory = [{'file': str(p.relative_to(stage)), 'sha256': digest(p)} for p in sorted(stage.rglob('*.xlsx'))]
    dump(stage / 'artifacts.json', inventory)
    return results


def main():
    if len(sys.argv) > 1 and sys.argv[1] == '--worker':
        return worker(*sys.argv[2:5], int(sys.argv[5]))
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path, required=True)
    parser.add_argument('--scale-rows', type=int, default=1000)
    args = parser.parse_args()
    if not 0 <= args.scale_rows <= 100000:
        parser.error('scale-rows must be between 0 and 100000')
    destination = args.output_dir.resolve()
    if destination.exists():
        parser.error('output-dir must be new')
    destination.parent.mkdir(parents=True, exist_ok=True)
    stage = Path(tempfile.mkdtemp(prefix='.gap-stage-', dir=destination.parent)).resolve()
    try:
        result = run(stage, args.scale_rows)
        if destination.exists():
            raise FileExistsError(destination)
        os.rename(stage, destination)
    finally:
        if stage.exists() and stage.parent == destination.parent and stage.name.startswith('.gap-stage-'):
            shutil.rmtree(stage)
    print(json.dumps({'output': str(destination), 'passed': result['passed'], 'failed': result['failed'],
                      'scale': result['observations']['scale']}, ensure_ascii=False))
    return 1 if result['failed'] else 0


if __name__ == '__main__':
    raise SystemExit(main())
