"""Independent P0 properties, combinations and full CLI reconciliation."""
from collections import Counter
from datetime import date
from io import BytesIO
import itertools
import atexit
import hashlib
import json
import os
from pathlib import Path
import random
import subprocess
import sys
import tempfile
import unittest
import unicodedata
from zipfile import ZipFile

from allpairspy import AllPairs
from hypothesis import given, seed, settings, strategies as st
from openpyxl import Workbook, load_workbook
from normalizer import dates, names, phone
from normalizer.models import HEADERS, RawCell
from normalizer.pipeline import normalize
from normalizer.preservation import restore_bytes

ROOT = Path(__file__).resolve().parents[1]
EXAMPLES = int(os.environ.get('QA_EXAMPLES', '300'))
SEED = int(os.environ.get('QA_SEED', '42'))
settings.register_profile('qa', max_examples=EXAMPLES, deadline=None)
settings.load_profile('qa')
digits9 = st.text(alphabet='0123456789', min_size=9, max_size=9).map(lambda tail: '9' + tail)
personal_names = st.sampled_from(['Андрей','Влад','Кирилл','Пётр','Семён','Анна','Любовь'])
calendar = st.dates(min_value=date(1,1,1),max_value=date(9999,12,31))
missing = st.sampled_from([None, '', ' ', '\u00a0', 'нет', 'NULL', 'N/A'])
CALLS=Counter()
UNIQUE={}
EXTRA={}


def record(label,*values):
    CALLS[label]+=1
    UNIQUE.setdefault(label,set()).add(hashlib.sha256(repr(values).encode('utf-8')).hexdigest())


@atexit.register
def save_statistics():
    target=os.environ.get('QA_STATS_PATH')
    if target:
        Path(target).write_text(json.dumps({'seed':SEED,'calls':dict(CALLS),
                               'unique_examples':{k:len(v) for k,v in UNIQUE.items()},'extra':EXTRA},indent=2),encoding='utf-8')


def book_bytes(rows):
    book=Workbook(); book.active.append(HEADERS)
    for row in rows:
        book.active.append(row)
    stream=BytesIO(); book.save(stream); book.close()
    return stream.getvalue()


def read_bytes(payload, sheet):
    book=load_workbook(BytesIO(payload),read_only=True,data_only=False)
    try:
        return list(book[sheet].values)
    finally:
        book.close()


class GeneratedProperties(unittest.TestCase):
    @seed(SEED)
    @given(digits9, st.sampled_from([' ', '-', '.', '\u00a0', '\u202f']))
    def test_phone_format_metamorphism_and_idempotence(self, national, sep):
        record('valid_phone',national,sep)
        expected='+7'+national
        groups=[national[:3],national[3:6],national[6:8],national[8:]]
        for variant in (expected,national, '8'+sep+sep.join(groups), '+7 ('+groups[0]+') '+sep.join(groups[1:])):
            result=phone.analyze(RawCell(variant))
            self.assertEqual(result.canonical,expected)
            self.assertEqual(phone.analyze(RawCell(result.canonical)).canonical,expected)

    @seed(SEED)
    @given(digits9, st.integers(min_value=2,max_value=11), st.sampled_from(['О','O','I','l']))
    def test_letter_mutation_cannot_create_phone(self,national,position,letter):
        record('invalid_phone_letters',national,position,letter)
        good='+7'+national
        bad=good[:position]+letter+good[position+1:]
        result=phone.analyze(RawCell(bad))
        self.assertIsNone(result.canonical)
        self.assertEqual(result.display,bad)

    @seed(SEED)
    @given(digits9,digits9)
    def test_multiple_and_unsupported_prefixes(self,first,second):
        record('multiple_and_prefix',first,second)
        if first!=second:
            self.assertIsNone(phone.analyze(RawCell('+7'+first+'; +7'+second)).canonical)
        for value in ('00'+first, '810'+first, '+7'+first+'0'):
            self.assertIsNone(phone.analyze(RawCell(value)).canonical)

    @seed(SEED)
    @given(calendar,st.sampled_from(['/', '.', '-']))
    def test_date_format_metamorphism_and_idempotence(self,value,sep):
        record('valid_dates',value,sep)
        expected=f'{value.year:04d}-{value.month:02d}-{value.day:02d}'
        for variant in (expected,f'{value.day}{sep}{value.month}{sep}{value.year:04d}',
                        expected+'T23:59:59-12:00'):
            result=dates.analyze(RawCell(variant))
            self.assertEqual(result.canonical,expected)
            self.assertEqual(dates.analyze(RawCell(result.canonical)).canonical,expected)

    @seed(SEED)
    @given(st.integers(min_value=1,max_value=9999),st.sampled_from([0,32,40,99]),st.integers(min_value=1,max_value=12))
    def test_invalid_calendar_never_rolls(self,year,day,month):
        record('invalid_dates',year,day,month)
        raw=f'{day:02d}.{month:02d}.{year:04d}'
        self.assertIsNone(dates.analyze(RawCell(raw)).canonical)

    @seed(SEED)
    @given(personal_names,st.sampled_from([' ', '\u00a0','\t','\n']),st.booleans())
    def test_name_case_spaces_and_idempotence(self,value,space,upper):
        record('valid_names',value,space,upper)
        raw=space+(value.upper() if upper else value.lower())+space
        result=names.analyze(RawCell(raw))
        self.assertEqual(result.canonical,value)
        self.assertEqual(names.analyze(RawCell(result.canonical)).canonical,value)

    @seed(SEED)
    @given(personal_names)
    def test_nfc_nfd_name_equivalence(self,value):
        record('name_nfc_nfd',value)
        self.assertEqual(names.analyze(RawCell(unicodedata.normalize('NFD',value))).canonical,value)

    @seed(SEED)
    @given(personal_names,st.sampled_from(['123','😀','\u200b',' (менеджер)']))
    def test_name_mutations_are_not_accepted(self,name,suffix):
        record('invalid_names',name,suffix)
        self.assertIsNone(names.analyze(RawCell(name+suffix)).canonical)

    @seed(SEED)
    @given(missing,st.sampled_from(['s','n']))
    def test_missing_required_fields_preserve_value(self,value,kind):
        record('missing',value,kind)
        for analyze in (names.analyze,phone.analyze,dates.analyze):
            result=analyze(RawCell(value,kind))
            self.assertIsNone(result.canonical)
            self.assertEqual(result.display,value)
            self.assertEqual(result.status,'MISSING')

    @seed(SEED)
    @given(st.sampled_from([True,False,'=1+1','#VALUE!']),st.sampled_from(['PHONE','DATE','NAME']))
    def test_excel_types_never_coerce_to_valid_field(self,value,field):
        record('excel_types',value,field)
        kind='b' if isinstance(value,bool) else 'f' if value.startswith('=') else 'e'
        analyze={'PHONE':phone.analyze,'DATE':dates.analyze,'NAME':names.analyze}[field]
        result=analyze(RawCell(value,kind))
        self.assertIsNone(result.canonical)
        self.assertEqual(result.display,value)

    @seed(SEED)
    @given(st.lists(st.tuples(personal_names,digits9,st.sampled_from(['сайт','email','custom',None])),min_size=1,max_size=8,unique_by=lambda row:row[1]))
    @settings(max_examples=min(EXAMPLES,30),deadline=None)
    def test_sequence_integrity_independent_additions_duplicates_and_order(self,records):
        record('sequences',records)
        rows=[[name,'+7'+national,'2026-09-06',source] for name,national,source in records]
        # Expected inventory and representatives calculated directly from construction.
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder)
            occupied={row[1] for row in rows}
            unique=next(f'+79{i:09d}' for i in range(10) if f'+79{i:09d}' not in occupied)
            extra=['Кирилл',unique,'2026-09-06',None]
            changed_sources=[list(r) for r in rows]
            changed_sources[0][3]='changed source'
            variants=[rows,rows+[list(rows[0])],rows+[[None,'не знаю','хх.09.2026',None]],list(reversed(rows)),rows+[extra],changed_sources]
            outputs=[]
            for index,variant in enumerate(variants):
                source,output=folder/f'in{index}.xlsx',folder/f'out{index}.xlsx'
                source.write_bytes(book_bytes(variant)); before=source.read_bytes()
                normalize(source,output)
                payload=output.read_bytes()
                original=read_bytes(payload,'Оригинал')[1:]
                self.assertEqual([r[0] for r in original],[f's1:r{i+2}' for i in range(len(variant))])
                self.assertEqual([r[3:7] for r in original],[tuple(r) for r in variant])
                self.assertEqual(restore_bytes(payload)[0],before)
                self.assertEqual(source.read_bytes(),before)
                outputs.append(read_bytes(payload,'Данные')[1:])
            self.assertEqual(outputs[0],[tuple(r) for r in rows])
            self.assertEqual(outputs[0],outputs[1])
            self.assertEqual(outputs[0],outputs[2])
            self.assertEqual(set(outputs[0]),set(outputs[3]))
            self.assertEqual(outputs[4],outputs[0]+[tuple(extra)])
            self.assertEqual([r[1] for r in outputs[5]],[r[1] for r in outputs[0]])


class CombinatorialCLI(unittest.TestCase):
    def test_pairwise_coverage_and_dangerous_three_way_cli_inventory(self):
        factors=[['Иван','Зайцев Кирилл',None,'К. Белоусов'],
                 ['+79174005219','не знаю',None,'+79174005219 доб. 123'],
                 ['2026-09-06','09/25/2026','хх.09.2026',None],
                 ['сайт',None,42], [False,True]]
        combinations=[list(row) for row in AllPairs(factors)]
        observed={(a,row[a],b,row[b]) for row in combinations for a,b in itertools.combinations(range(len(factors)),2)}
        required={(a,x,b,y) for a,b in itertools.combinations(range(len(factors)),2) for x in factors[a] for y in factors[b]}
        self.assertEqual(observed,required)
        triples=list(itertools.product(factors[0][2:],factors[1][1:3],factors[2][2:]))
        combinations.extend([[a,b,c,'сайт',True] for a,b,c in triples])
        EXTRA.update({'pairwise_rows':len(combinations)-len(triples),'covered_pairs':len(observed),'dangerous_triples':len(triples)})
        rows=[]
        expected_data=[]
        for index,(name,phone_value,date_value,source,duplicate) in enumerate(combinations):
            unique=f'+79{index:09d}'
            if phone_value and phone_value.startswith('+7'):
                phone_value=unique+(' доб. 123' if 'доб.' in phone_value else '')
            row=[name,phone_value,date_value,source]
            rows.append(row)
            if duplicate:
                rows.append(list(row))
            if name in ('Иван','Зайцев Кирилл') and phone_value==unique and date_value in ('2026-09-06','09/25/2026'):
                expected_data.append(('Иван' if name=='Иван' else 'Кирилл',unique,
                                      '2026-09-06' if date_value=='2026-09-06' else '2026-09-25',source))
        # Different reproducible seeds exercise physical row orders and blank/source values.
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder)
            for run_seed in (SEED,SEED+1,SEED+17):
                shuffled=list(rows); random.Random(run_seed).shuffle(shuffled)
                original=book_bytes(shuffled)
                source,output=folder/f'in{run_seed}.xlsx',folder/f'out{run_seed}.xlsx'
                source.write_bytes(original)
                proc=subprocess.run([sys.executable,str(ROOT/'normalize.py'),'--input',str(source),'--output',str(output)],
                                    cwd=ROOT,capture_output=True,text=True,encoding='utf-8',timeout=60)
                self.assertEqual(proc.returncode,0,proc.stderr)
                payload=output.read_bytes()
                self.assertEqual(set(read_bytes(payload,'Данные')[1:]),set(expected_data))
                raw=read_bytes(payload,'Оригинал')[1:]
                self.assertEqual([r[3:7] for r in raw],[tuple(r) for r in shuffled])
                self.assertEqual(len({r[0] for r in raw}),len(shuffled))
                self.assertEqual(restore_bytes(payload)[0],original)
                self.assertEqual(source.read_bytes(),original)
                run=json.loads(proc.stdout)
                self.assertEqual(run['input_rows'],len(shuffled))
                self.assertEqual(run['data_rows'],len(expected_data))
            EXTRA.update({'cli_combination_runs':3,'rows_per_cli':len(rows)})

    def test_nonfinite_cells_are_preserved_without_losing_valid_rows(self):
        from xml.etree import ElementTree as ET
        ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder)
            good=book_bytes([['Иван','+79174005219','2026-09-06','сайт'], ['Кирилл','+79174005220','2026-09-06','email']])
            for i,value in enumerate(('NaN','INF','-INF','1e9999')):
                source,output=folder/f'bad{i}.xlsx',folder/f'out{i}.xlsx'
                with ZipFile(BytesIO(good)) as old,ZipFile(source,'w') as new:
                    for part in old.infolist():
                        payload=old.read(part.filename)
                        if part.filename=='xl/worksheets/sheet1.xml':
                            root=ET.fromstring(payload)
                            node=root.find(f'.//{ns}c[@r="B2"]'); node.clear(); node.set('r','B2'); node.set('t','n')
                            ET.SubElement(node,ns+'v').text=value
                            payload=ET.tostring(root)
                        new.writestr(part.filename,payload)
                proc=subprocess.run([sys.executable,str(ROOT/'normalize.py'),'--input',str(source),'--output',str(output)],
                                    cwd=ROOT,capture_output=True,text=True,encoding='utf-8',timeout=30)
                self.assertEqual(proc.returncode,0,proc.stderr)
                payload=output.read_bytes()
                self.assertEqual(read_bytes(payload,'Данные')[1:],[('Кирилл','+79174005220','2026-09-06','email')])
                self.assertEqual(read_bytes(payload,'Оригинал')[1][4],value)
                self.assertEqual(read_bytes(payload,'Проверить')[1][1],value)
                self.assertEqual(json.loads(proc.stdout)['codes']['CELL_NUMERIC_INVALID'],1)
                self.assertEqual(restore_bytes(payload)[0],source.read_bytes())
            EXTRA['cli_nonfinite_runs']=4


if __name__=='__main__':
    unittest.main()
