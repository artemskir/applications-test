"""Generate labeled XLSX fixtures; deliberately does not clean user data."""
import argparse
import hashlib
import json
import math
import os
from pathlib import Path
import random
import shutil
import sys
import tempfile
from datetime import date, datetime, timedelta

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from fixtures.scenarios import CASES, HEADERS, materialize
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904, to_excel

SPEC_VERSION = "1.4"


def clean_rows(count, seed):
    rng = random.Random(seed)
    names = ["Андрей", "Егор", "Анастасия", "Ирина", "Влад"]
    sources = ["сайт", "звонок", "telegram", "email", "партнёр"]
    return [[rng.choice(names), f"+79{100000000 + i:09}",
             (date(2026, 9, 1) + timedelta(days=rng.randrange(30))).isoformat(),
             rng.choice(sources)] for i in range(count)]


def fixture_rows(count, seed, mode):
    clean = clean_rows(count, seed)
    rows, labels = [], []
    for i, base in enumerate(clean):
        row = list(base)
        selected = None
        if mode == "dirty" or (mode == "mixed" and i % 3 != 0):
            selected = CASES[i % len(CASES)]
            row[HEADERS.index(selected["field"])] = selected["input"]
        labels.append({"excel_row": i + 2, "case_id": selected["id"] if selected else "CLEAN",
                       "input": row, "field_expectation": selected})
        rows.append(row)
    # Stable order: catalog coverage is independent of random data values.
    return rows, labels


def write_workbook(path, rows, epoch=1900, headers=None, merges=(), second_sheet=False):
    if Path(path).exists():
        raise FileExistsError(f"Output workbook already exists: {path}")
    workbook = Workbook()
    if epoch == 1904:
        workbook.epoch = CALENDAR_MAC_1904
    sheet = workbook.active
    sheet.title = "Заявки"
    sheet.append(headers or HEADERS)
    for row_number, row in enumerate(rows, 2):
        for column_number, value in enumerate(row, 1):
            cell = sheet.cell(row_number, column_number)
            actual_value = materialize(value)
            if isinstance(actual_value, str) and len(actual_value) > 32767:
                raise ValueError(f"Cell {cell.coordinate} exceeds XLSX limit: {len(actual_value)} > 32767; no truncation allowed")
            if isinstance(actual_value, float) and not math.isfinite(actual_value):
                raise ValueError(f"Cell {cell.coordinate}: non-finite number cannot be preserved in XLSX; no silent blank allowed")
            if isinstance(actual_value, datetime) and actual_value.utcoffset() is not None:
                raise ValueError(f"Cell {cell.coordinate}: native timezone-aware datetime cannot be preserved in XLSX; use literal ISO text")
            cell.value = actual_value
            kind = value.get("type") if isinstance(value, dict) else None
            if kind == "formula":
                cell.data_type = "f"
            elif kind == "error":
                cell.data_type = "e"
            elif isinstance(cell.value, str):
                cell.data_type = "s"  # Preserve text beginning with =, +, -, @.
            if isinstance(value, dict) and value.get("number_format"):
                cell.number_format = value["number_format"]
    for merged_range in merges:
        sheet.merge_cells(merged_range)
    if second_sheet:
        other = workbook.create_sheet("Другие заявки")
        for row in sheet.iter_rows(values_only=True):
            other.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)
    workbook.close()


def suite(stage, count, seed):
    files = []
    for mode in ("clean", "dirty", "mixed"):
        rows, labels = fixture_rows(count if mode != "dirty" else max(count, len(CASES)), seed, mode)
        name = f"{mode}.xlsx"
        write_workbook(stage / name, rows)
        files.append({"file": name, "rows": len(rows), "labels": labels})
    a = ["Андрей Волков", "+79174005219", "09/03/2026", "сайт"]
    duplicates = [a, ["андрей волков", "8 (917) 400-52-19", "2026-03-09", "сайт"],
                  ["Андрей Волков", "+79174005219", "09.03.2026", "email"],
                  ["Егор Волков", "+79174005219", "10.03.2026", "звонок"],
                  [None, None, "bad", "сайт"], [None, None, "bad", "сайт"],
                  [None, "+79174005220", "31.04.2026", "сайт"],
                  ["Ирина", "+79174005220", "2026-04-30", "email"],
                  [None, None, None, None], ["Егор", "+79174005221", "2026-09-06", "сайт"],
                  ["Пётр", "+79174005222", "2026-09-06", "сайт"],
                  ["Пётр", "+79174005222", "2026-09-06", "email"]]
    write_workbook(stage / "duplicates.xlsx", duplicates)
    files.append({"file": "duplicates.xlsx", "rows": 12,
                  "expected": {"data_input_rows": [11, 12], "representative_or_ungrouped_rows": [2, 6, 7, 8, 10, 11, 12], "duplicate_rows": [3, 4, 5, 9, 13],
                               "duplicate_conflict_rows": [2, 3, 4, 5, 8, 9],
                               "problem_rows": [2, 3, 4, 5, 6, 7, 8, 9, 10]},
                  "note": "Rows 12/13 differ only by source and are audit-only. All participants of conflicting groups are problematic; empty phones never deduplicate."})
    baseline = clean_rows(2, seed)
    for epoch in (1900, 1904):
        rows = [list(baseline[0])]
        rows[0][2] = {"type": "serial", "value": to_excel(datetime(2026, 9, 6),
                                    CALENDAR_MAC_1904 if epoch == 1904 else Workbook().epoch),
                      "number_format": "yyyy-mm-dd"}
        name = f"epoch_{epoch}.xlsx"
        write_workbook(stage / name, rows, epoch=epoch)
        files.append({"file": name, "expected_date": "2026-09-06"})
    for name, options, code in [
        ("missing_header.xlsx", {"headers": HEADERS[:3]}, "SCHEMA_INVALID"),
        ("duplicate_header.xlsx", {"headers": ["Имя", "Телефон", "Телефон", "Источник"]}, "SCHEMA_INVALID"),
        ("extra_header.xlsx", {"headers": HEADERS + ["Примечание"]}, "SCHEMA_INVALID"),
        ("merged.xlsx", {"merges": ["A2:A3"]}, "MERGED_INPUT"),
        ("multiple_sheets.xlsx", {"second_sheet": True}, "SHEET_AMBIGUOUS"),
    ]:
        rows = [row[:3] for row in baseline] if name == "missing_header.xlsx" else baseline
        write_workbook(stage / name, rows, **options)
        files.append({"file": name, "fatal_code": code})
    write_workbook(stage / "headers_only.xlsx", [])
    files.append({"file": "headers_only.xlsx", "expected_rows": 0})
    wb = Workbook()
    wb.save(stage / "empty_book.xlsx")
    wb.close()
    files.append({"file": "empty_book.xlsx", "fatal_code": "SCHEMA_INVALID"})
    (stage / "corrupt.xlsx").write_bytes(b"This is not a ZIP/XLSX file.\n")
    files.append({"file": "corrupt.xlsx", "fatal_code": "INPUT_INVALID"})
    return files


def generate(destination, mode="suite", rows=48, seed=20261008, epoch=1900):
    destination = Path(destination).resolve()
    if destination.exists():
        raise FileExistsError(f"Output directory already exists: {destination}")
    if not 0 <= rows <= 100000:
        raise ValueError("rows must be between 0 and 100000")
    destination.parent.mkdir(parents=True, exist_ok=True)
    stage = Path(tempfile.mkdtemp(prefix=".fixture-stage-", dir=destination.parent)).resolve()
    try:
        if mode == "suite":
            files = suite(stage, rows, seed)
        else:
            values, labels = fixture_rows(rows, seed, mode)
            # A date-system-specific serial in the catalog is adjusted for 1904.
            if epoch == 1904:
                for row, label in zip(values, labels):
                    value = row[2]
                    if isinstance(value, dict) and value.get("type") == "serial":
                        if value["value"] == 46271:
                            row[2] = dict(value, value=44809)
                        elif value["value"] == 60:
                            expectation = dict(label["field_expectation"])
                            expectation.update(expected="1904-03-01", canonical="1904-03-01", codes=[],
                                               diagnostic_codes=[], problem_list=False, status="NORMALIZED",
                                               note="Serial 60 is valid in the 1904 system")
                            label["field_expectation"] = expectation
                    label["input"] = row
            write_workbook(stage / f"{mode}.xlsx", values, epoch=epoch)
            files = [{"file": f"{mode}.xlsx", "rows": rows, "labels": labels}]
        for entry in files:
            entry["sha256"] = hashlib.sha256((stage / entry["file"]).read_bytes()).hexdigest()
        manifest = {"schema_version": "1.1", "spec_version": SPEC_VERSION, "seed": seed,
                    "mode": mode, "epoch": epoch if mode != "suite" else "per workbook",
                    "files": files,
                    "note": "Logical values and labels are reproducible; XLSX ZIP timestamps are not."}
        (stage / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        (stage / "cases.json").write_text(json.dumps(CASES, ensure_ascii=False, indent=2), encoding="utf-8")
        # Windows rename refuses an existing target; never replace a user's directory.
        if destination.exists():
            raise FileExistsError(f"Output directory appeared during generation: {destination}")
        os.rename(stage, destination)
    finally:
        if stage.exists() and stage.parent == destination.parent and stage.name.startswith(".fixture-stage-"):
            shutil.rmtree(stage)
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--mode", choices=["suite", "clean", "dirty", "mixed"], default="suite")
    parser.add_argument("--rows", type=int, default=48)
    parser.add_argument("--seed", type=int, default=20261008)
    parser.add_argument("--epoch", type=int, choices=[1900, 1904], default=1900)
    args = parser.parse_args()
    try:
        manifest = generate(args.output_dir, args.mode, args.rows, args.seed, args.epoch)
    except (ValueError, FileExistsError, OSError) as exc:
        print(f"Generation failed: {exc}", file=sys.stderr)
        return 2
    print(f"Generated {len(manifest['files'])} XLSX files in {args.output_dir.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
