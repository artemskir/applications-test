"""Typed cells plus raw OOXML from the same immutable source bytes."""
from io import BytesIO
from copy import copy
import math
from pathlib import PurePosixPath
import posixpath
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile, BadZipFile, ZIP_DEFLATED
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries, get_column_letter
from .models import HEADERS, NormalizationError, RawCell, RawRow, jsonable

MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS = {'m': MAIN}


def invalid(message):
    return NormalizationError('INPUT_INVALID', message)


def xml(data):
    if re.search(br'<!\s*(?:DOCTYPE|ENTITY)\b', data.replace(b'\x00', b''), re.I):
        raise invalid('Недопустимый XML DTD/ENTITY.')
    try:
        return ET.fromstring(data)
    except ET.ParseError as exc:
        raise invalid('Повреждён XML книги.') from exc


def part_target(owner, target):
    if '\\' in target or ':' in target:
        raise invalid('Недопустимый путь OOXML-части.')
    name = posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(posixpath.dirname(owner), target))
    if name.startswith('../') or name in ('.', '..'):
        raise invalid('OOXML-отношение выходит за границы книги.')
    return name


def relationships(archive, owner):
    path = str(PurePosixPath(owner).parent / '_rels' / (PurePosixPath(owner).name + '.rels')) if owner else '_rels/.rels'
    if path not in archive.namelist():
        return {}
    result = {}
    for item in xml(archive.read(path)):
        ident = item.get('Id')
        if not ident or ident in result:
            raise invalid('Повторяющееся или пустое Id отношения OOXML.')
        result[ident] = dict(item.attrib)
    return result


def validate_archive(snapshot, max_uncompressed_bytes):
    try:
        archive = ZipFile(BytesIO(snapshot))
        infos = archive.infolist()
        names = [info.filename for info in infos]
        if len(names) != len(set(names)):
            raise invalid('Повторяющиеся имена ZIP-частей.')
        if len(infos) > 10000 or sum(i.file_size for i in infos) > max_uncompressed_bytes:
            raise NormalizationError('INPUT_LIMIT', 'Превышен лимит распакованной книги или ZIP-частей.', 5)
        if any(i.flag_bits & 1 for i in infos):
            raise invalid('Зашифрованная книга не поддерживается.')
        if not {'[Content_Types].xml', '_rels/.rels'}.issubset(names):
            raise invalid('ZIP не является книгой XLSX.')
        if archive.testzip() is not None:
            raise invalid('Ошибка CRC в ZIP-части.')
        for name in names:
            if name.endswith(('.xml', '.rels')):
                xml(archive.read(name))
        types = xml(archive.read('[Content_Types].xml'))
        if not any(item.get('ContentType') == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml' for item in types):
            raise invalid('Требуется обычная книга XLSX без макросов.')
        root_rels = relationships(archive, '')
        offices = [r for r in root_rels.values() if r.get('Type', '').endswith('/officeDocument') and r.get('TargetMode') != 'External']
        if len(offices) != 1:
            raise invalid('Отсутствует однозначное отношение к workbook.')
        workbook_part = part_target('', offices[0]['Target'])
        book = xml(archive.read(workbook_part))
        rels = relationships(archive, workbook_part)
        sheets = []
        for index, item in enumerate(book.findall('m:sheets/m:sheet', NS), 1):
            relation = rels.get(item.get('{' + REL + '}id'))
            if not relation or relation.get('TargetMode') == 'External':
                raise invalid('Недопустимое отношение листа.')
            target = part_target(workbook_part, relation['Target'])
            if target not in names:
                raise invalid('Отсутствует XML листа.')
            if relation.get('Type', '').endswith('/worksheet'):
                sheets.append((index, item.get('name'), target))
        return archive, sheets
    except NormalizationError:
        if 'archive' in locals():
            archive.close()
        raise
    except (BadZipFile, KeyError, ValueError, RuntimeError, OSError) as exc:
        if 'archive' in locals():
            archive.close()
        raise invalid('Не удалось прочитать контейнер XLSX.') from exc


def sheet_inventory(archive, part):
    root = xml(archive.read(part))
    if root.tag != '{' + MAIN + '}worksheet':
        raise invalid('Неподдерживаемое пространство имён листа.')
    cells, coordinates = {}, set()
    last, max_column = 1, 1
    ordered_cells = []
    previous_row = 0
    for row_node in root.findall('m:sheetData/m:row', NS):
        try:
            physical_row = int(row_node.get('r', previous_row + 1))
        except ValueError as exc:
            raise invalid('Недопустимый номер строки XML.') from exc
        if physical_row <= previous_row or physical_row > 1048576:
            raise invalid('Повторяющиеся или неупорядоченные строки XML.')
        previous_row, previous_column = physical_row, 0
        for node in row_node.findall('m:c', NS):
            ref = node.get('r', f'{get_column_letter(previous_column + 1)}{physical_row}').upper()
            ordered_cells.append((node, ref, physical_row))
            try:
                previous_column = coordinate_to_tuple(ref)[1]
            except (ValueError, TypeError, AttributeError) as exc:
                raise invalid('Недопустимая координата ячейки.') from exc
    for node, ref, physical_row in ordered_cells:
        try:
            row, column = coordinate_to_tuple(ref)
        except (ValueError, TypeError, AttributeError) as exc:
            raise invalid('Недопустимая координата ячейки.') from exc
        if row != physical_row or row < 1 or column < 1 or row > 1048576 or column > 16384 or ref in coordinates:
            raise invalid('Повторяющаяся или выходящая за пределы Excel ячейка.')
        coordinates.add(ref)
        if any(node.find('m:' + tag, NS) is not None for tag in ('v', 'f', 'is')):
            cells[ref] = node
            last, max_column = max(last, row), max(max_column, column)
    merges = [node.get('ref') for node in root.findall('m:mergeCells/m:mergeCell', NS)]
    rels = relationships(archive, part)
    links = []
    for node in root.findall('m:hyperlinks/m:hyperlink', NS):
        details = dict(node.attrib)
        relationship = rels.get(node.get('{' + REL + '}id'), {})
        if relationship:
            details['target'] = relationship.get('Target')
            details['target_mode'] = relationship.get('TargetMode')
        try:
            bounds = range_boundaries(node.get('ref'))
        except (ValueError, TypeError) as exc:
            raise invalid('Недопустимый диапазон гиперссылки.') from exc
        links.append((bounds, details))
    return cells, last, max_column, merges, links


def metadata(cell, node, epoch, links):
    details = {'coordinate': cell.coordinate, 'ooxml_type': node.get('t', 'n') if node is not None else None,
               'xml_value': node.findtext('m:v', namespaces=NS) if node is not None else None,
               'xml_formula': node.findtext('m:f', namespaces=NS) if node is not None else None,
               'raw_xml': ET.tostring(node, encoding='unicode') if node is not None else None,
               'decoded_value': jsonable(cell.value), 'python_type': type(cell.value).__name__,
               'style_id': getattr(cell, 'style_id', getattr(cell, '_style_id', 0)), 'number_format': cell.number_format, 'epoch': epoch}
    for (left, top, right, bottom), link in links:
        if left <= cell.column <= right and top <= cell.row <= bottom:
            details['hyperlink'] = link
    value, kind = cell.value, cell.data_type
    if node is not None and node.get('t', 'n') == 'n' and node.find('m:f', NS) is None:
        lexical = node.findtext('m:v', namespaces=NS)
        if lexical is not None:
            try:
                value = int(lexical) if re.fullmatch(r'[+-]?[0-9]+', lexical) else float(lexical)
            except ValueError as exc:
                raise invalid('Некорректное числовое значение XML.') from exc
            if isinstance(value, float) and not math.isfinite(value):
                details['parser_shadow_value'] = jsonable(cell.value)
                details['decoded_value'] = lexical
                details['python_type'] = 'non_finite_numeric_lexeme'
                details['parse_issue'] = 'CELL_NUMERIC_INVALID'
                value = lexical
            kind = 'n'
    return RawCell(value, kind, cell.coordinate, cell.number_format, epoch, details)


def safe_parser_snapshot(snapshot, archive, sheets):
    replacements = {}
    for _, _, part in sheets:
        root = xml(archive.read(part))
        changed = False
        for node in root.findall('.//m:sheetData/m:row/m:c', NS):
            if node.get('t', 'n') != 'n' or node.find('m:f', NS) is not None:
                continue
            v = node.find('m:v', NS)
            if v is None or v.text is None:
                continue
            try:
                value = float(v.text)
            except ValueError:
                continue
            if not math.isfinite(value):
                v.text = '0'
                changed = True
        if changed:
            replacements[part] = ET.tostring(root, encoding='utf-8')
    if not replacements:
        return snapshot
    buffer = BytesIO()
    with ZipFile(buffer, 'w', compression=ZIP_DEFLATED) as shadow:
        for item in archive.infolist():
            # writestr updates ZipInfo size/CRC; do not mutate the source archive inventory.
            shadow.writestr(copy(item), replacements[item.filename] if item.filename in replacements else archive.read(item.filename))
    return buffer.getvalue()


def read_rows(snapshot, sheet_name=None, max_uncompressed_bytes=536870912, max_rows=1048575):
    archive, sheets = validate_archive(snapshot, max_uncompressed_bytes)
    workbook = None
    try:
        parser_snapshot = safe_parser_snapshot(snapshot, archive, sheets)
        workbook = load_workbook(BytesIO(parser_snapshot), read_only=True, data_only=False, keep_links=False, rich_text=False)
        epoch = 1904 if workbook.epoch.year == 1904 else 1900
        candidates = []
        for index, name, part in sheets:
            if sheet_name is not None and name != sheet_name:
                continue
            ws = workbook[name]
            ws.reset_dimensions()
            inventory = sheet_inventory(archive, part)
            raw, last, max_column, merges, links = inventory
            header_cells = next(ws.iter_rows(min_row=1, max_row=1, min_col=1, max_col=max_column), ())
            mapping = {}
            valid = True
            for column, cell in enumerate(header_cells, 1):
                value = cell.value
                if value is None or isinstance(value, str) and not value.strip():
                    continue
                header = value.strip() if isinstance(value, str) and cell.data_type in ('s', 'str', 'inlineStr') else None
                if header not in HEADERS or header in mapping:
                    valid = False
                    break
                mapping[header] = column
            if valid and set(mapping) == set(HEADERS):
                if any(coordinate_to_tuple(ref)[1] not in mapping.values() for ref in raw):
                    valid = False
            if valid and set(mapping) == set(HEADERS):
                candidates.append((index, name, ws, inventory, mapping))
            elif sheet_name is not None:
                raise NormalizationError('SCHEMA_INVALID', 'Выбранный лист не соответствует четырём колонкам.')
        if not candidates:
            raise NormalizationError('SCHEMA_INVALID', 'Не найден лист с колонками Имя, Телефон, Дата заявки, Источник.')
        if len(candidates) != 1:
            raise NormalizationError('SHEET_AMBIGUOUS', 'Несколько листов заявок: укажите --sheet.')
        index, name, ws, (raw, last, max_column, merges, links), mapping = candidates[0]
        if last - 1 > max_rows:
            raise NormalizationError('INPUT_LIMIT', 'Превышен --max-rows.', 5)
        for merge in merges:
            left, top, right, bottom = range_boundaries(merge)
            if top <= last and bottom >= 1 and any(left <= column <= right for column in mapping.values()):
                raise NormalizationError('MERGED_INPUT', 'Объединённые ячейки пересекают заголовок или заявки.')
        rows = []
        for physical, cells in enumerate(ws.iter_rows(min_row=2, max_row=last, min_col=1, max_col=max_column), 2):
            selected = []
            for header in HEADERS:
                column = mapping[header]
                cell = cells[column - 1]
                if cell.value is None and not hasattr(cell, 'coordinate'):
                    ref = f'{get_column_letter(column)}{physical}'
                    details = {'coordinate': ref, 'ooxml_type': None, 'epoch': epoch, 'decoded_value': None,
                               'python_type': 'NoneType', 'number_format': 'General', 'style_id': 0,
                               'raw_xml': None, 'xml_value': None, 'xml_formula': None}
                    for (left, top, right, bottom), link in links:
                        if left <= column <= right and top <= physical <= bottom:
                            details['hyperlink'] = link
                    selected.append(RawCell(None, 'n', ref, epoch=epoch, metadata=details))
                else:
                    selected.append(metadata(cell, raw.get(cell.coordinate), epoch, links))
            rows.append(RawRow(f's{index}:r{physical}', name, physical, selected))
        if len(rows) != max(0, last - 1):
            raise NormalizationError('INVARIANT_FAILED', 'Не совпало число прочитанных исходных строк.', 4)
        return rows, {'sheet': name, 'sheet_index': index, 'epoch': epoch}
    except NormalizationError:
        raise
    except (KeyError, ValueError, TypeError, IndexError, OverflowError, BadZipFile) as exc:
        raise invalid('Не удалось разобрать XLSX и типы ячеек.') from exc
    finally:
        if workbook is not None:
            workbook.close()
        archive.close()
