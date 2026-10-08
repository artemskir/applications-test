import base64
import hashlib
from io import BytesIO
import json
from pathlib import Path
from openpyxl import load_workbook
from .models import SHEETS, NormalizationError
from .report import AUDIT_HEADERS, KINDS


def error(message):
    return NormalizationError('INVARIANT_FAILED', message, 4)


def restore_bytes(path_or_bytes):
    source = BytesIO(path_or_bytes if isinstance(path_or_bytes, bytes) else Path(path_or_bytes).read_bytes())
    workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    try:
        if workbook.sheetnames not in (SHEETS, SHEETS[:4]):
            raise error('Неверный список листов результата; требуется пять листов, либо исторический результат 1.0–1.3.')
        ws = workbook['Аудит']
        ws.reset_dimensions()
        iterator = ws.iter_rows(min_col=1, max_col=len(AUDIT_HEADERS), values_only=False)
        if [cell.value for cell in next(iterator, [])] != AUDIT_HEADERS:
            raise error('Неверный заголовок Аудита.')
        chunks, run, values = {}, None, {}
        for cells in iterator:
            row = [cell.value for cell in cells]
            if not row or row[0] not in KINDS or len(row) != len(AUDIT_HEADERS):
                raise error('Неизвестная или повреждённая запись аудита.')
            if any(cell.data_type in ('f', 'e') for cell in cells):
                raise error('В аудите обнаружены активная формула или ошибка Excel.')
            kind = row[0]
            if kind == 'RUN':
                if run is not None:
                    raise error('Повторяющаяся запись RUN.')
                run = json.loads(row[9])
            if kind in ('INPUT_CHUNK', 'CELL_METADATA_CHUNK', 'VALUE_CHUNK'):
                index, payload = row[10:12]
                if type(index) is not int or index < 0 or not isinstance(payload, str) or len(payload) > 30000:
                    raise error('Недопустимая часть аудита.')
                if kind == 'INPUT_CHUNK':
                    group = chunks
                else:
                    info = json.loads(row[9])
                    ident = info['value_id']
                    if ident not in values:
                        values[ident] = (info, {})
                    elif values[ident][0] != info:
                        raise error('Расходятся метаданные частей JSON.')
                    group = values[ident][1]
                if index in group:
                    raise error('Повторяющийся индекс части аудита.')
                group[index] = payload
        if not isinstance(run, dict) or run.get('encoding') != 'base64':
            raise error('Отсутствует корректная запись RUN.')
        if workbook.sheetnames == SHEETS[:4] and run.get('spec_version') not in ('1.0', '1.1', '1.2', '1.3'):
            raise error('У результата новой версии отсутствует лист Проверить.')
        count = run.get('chunk_count')
        if type(count) is not int or count < 1 or count != len(chunks) or any(i not in chunks for i in range(count)):
            raise error('Не совпали индексы частей исходника.')
        original = base64.b64decode(''.join(chunks[i] for i in range(count)), validate=True)
        if len(original) != run.get('byte_count') or hashlib.sha256(original).hexdigest() != run.get('sha256'):
            raise error('Не совпали размер или SHA-256 исходника.')
        for info, pieces in values.values():
            count = info['chunk_count']
            if type(count) is not int or count < 1 or count != len(pieces) or any(i not in pieces for i in range(count)):
                raise error('Не совпали индексы частей JSON.')
            value = ''.join(pieces[i] for i in range(count))
            if hashlib.sha256(value.encode('utf-8')).hexdigest() != info['sha256']:
                raise error('Не совпал SHA-256 частей JSON.')
        return original, run
    except NormalizationError:
        raise
    except (ValueError, KeyError, TypeError, OSError) as exc:
        raise error('Не удалось проверить аудит и восстановить исходник.') from exc
    finally:
        workbook.close()
