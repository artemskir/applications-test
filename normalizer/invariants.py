from collections import Counter
from datetime import date
from io import BytesIO
import json
import re
from openpyxl import load_workbook
from .models import HEADERS, SHEETS, NormalizationError, jsonable
from .report import ORIGINAL_HEADERS, PROBLEM_HEADERS, REVIEW_HEADERS, visible
from .preservation import restore_bytes


def require(condition, message):
    if not condition:
        raise NormalizationError('INVARIANT_FAILED', message, 4)


def validate_rows(rows):
    identities = {row.raw.row_id: row for row in rows}
    require(len(identities) == len(rows), 'Неуникальные исходные row_id.')
    phones = []
    for row in rows:
        require(row.representative_id in identities, 'Нет представителя строки.')
        representative = identities[row.representative_id]
        require(representative.disposition != 'DUPLICATE', 'Представитель оказался дублем.')
        if row.disposition == 'DUPLICATE':
            require(row.output_row is None and representative.raw.number < row.raw.number, 'Нарушен выбор первого дубля.')
        else:
            require(row.representative_id == row.raw.row_id, 'Неверная ссылка представителя.')
        require((row.output_row is not None) == row.admitted, 'Не совпал допуск в Данные.')
        phone = row.fields[1].canonical
        if phone is not None:
            require(isinstance(phone, str) and re.fullmatch(r'\+7[0-9]{10}', phone) is not None, 'Неверный canonical телефона.')
            if row.disposition != 'DUPLICATE':
                phones.append(phone)
        else:
            require(row.disposition == 'UNGROUPED', 'Строка без телефона попала в дубли.')
        normalized_date = row.fields[2].canonical
        if normalized_date is not None:
            require(re.fullmatch(r'[0-9]{4}-[0-9]{2}-[0-9]{2}', normalized_date) is not None, 'Неверный формат canonical даты.')
            date.fromisoformat(normalized_date)
    require(len(phones) == len(set(phones)), 'Повторяющиеся canonical телефоны представителей.')


def comparable(value):
    value = visible(value)
    return None if value == '' else value


def same_json(left, right):
    # Python equality considers True == 1; provenance must retain JSON types.
    return json.dumps(jsonable(left), sort_keys=True, allow_nan=False) == json.dumps(jsonable(right), sort_keys=True, allow_nan=False)


def same_excel(actual, expected):
    # Excel represents int/float in the same numeric type, but never bool as a number.
    return actual == expected and all((type(a) is bool) == (type(e) is bool) and
                                     (isinstance(a, str)) == (isinstance(e, str)) for a, e in zip(actual, expected))


def verify_report(path, rows, snapshot, counts):
    output_snapshot = path.read_bytes()
    original, run = restore_bytes(output_snapshot)
    require(original == snapshot, 'Не восстановились точные байты входа.')
    require(all(run[key] == value for key, value in counts.items()), 'Не совпали счётчики RUN.')
    workbook = load_workbook(BytesIO(output_snapshot), read_only=True, data_only=False, keep_links=False)
    try:
        require(workbook.sheetnames == SHEETS, 'Неверный список выходных листов.')
        for name in SHEETS:
            workbook[name].reset_dimensions()
        def audit_records():
            return workbook['Аудит'].iter_rows(min_row=2, max_col=12, values_only=True)
        packed_values = {}
        for record in audit_records():
            if record[0] in ('CELL_METADATA_CHUNK', 'VALUE_CHUNK'):
                info = json.loads(record[9])
                packed_values.setdefault(info['value_id'], {})[record[10]] = record[11]
        def unpack(value):
            result = json.loads(value)
            if isinstance(result, dict) and 'chunk_ref' in result:
                pieces = packed_values[result['value_id']]
                result = json.loads(''.join(pieces[i] for i in range(result['chunk_count'])))
            return result
        original_rows = workbook['Оригинал'].iter_rows(max_col=len(ORIGINAL_HEADERS), values_only=True)
        require(list(next(original_rows)) == ORIGINAL_HEADERS, 'Неверный заголовок Оригинала.')
        actual_ids = []
        for actual, expected in zip(original_rows, rows, strict=True):
            actual_ids.append(actual[0])
            require(actual[:3] == (expected.raw.row_id, expected.raw.sheet, expected.raw.number), 'Неверная исходная идентичность.')
            require(same_excel(actual[3:7], tuple(comparable(c.value) for c in expected.raw.cells)), 'Изменены исходные значения/типы в Оригинале.')
            require(same_json(unpack(actual[7]), {HEADERS[i]: c.metadata for i, c in enumerate(expected.raw.cells)}),
                    'Изменены исходные метаданные в Оригинале.')
        require(actual_ids == [row.raw.row_id for row in rows], 'Нарушена инвентаризация Оригинала.')
        data_rows = workbook['Данные'].iter_rows(max_col=len(HEADERS), values_only=True)
        require(list(next(data_rows)) == HEADERS, 'Неверный заголовок Данных.')
        for actual, expected in zip(data_rows, (row for row in rows if row.admitted), strict=True):
            require(same_excel(actual, tuple(comparable(r.display) for r in expected.fields)), 'Неверное значение в Данных.')
        problem_rows = workbook['Проблемы'].iter_rows(max_col=len(PROBLEM_HEADERS), values_only=True)
        require(list(next(problem_rows)) == PROBLEM_HEADERS, 'Неверный заголовок Проблем.')
        problems = list(problem_rows)
        require([row[0] for row in problems] == [r.raw.row_id for r in rows if r.codes], 'Неполный список проблемных строк.')
        require(all(row[6] == ';'.join(expected.codes) for row, expected in zip(problems, (r for r in rows if r.codes), strict=True)), 'Не совпали коды проблем.')
        for actual, expected in zip(problems, (r for r in rows if r.codes), strict=True):
            require(actual[1:6] == (expected.raw.sheet, expected.raw.number, expected.output_row,
                    expected.representative_id, expected.disposition), 'Неверная идентичность Проблем.')
            require(same_excel(actual[8:], tuple(comparable(v) for c, r in zip(expected.raw.cells, expected.fields)
                                       for v in (c.value, r.display))), 'Изменены исходные/обработанные поля в Проблемах.')
        review = workbook['Проверить'].iter_rows(max_col=len(REVIEW_HEADERS), values_only=True)
        require(list(next(review)) == REVIEW_HEADERS, 'Неверный заголовок Проверить.')
        for actual, expected in zip(review, (row for row in rows if row.codes), strict=True):
            require(same_excel(actual[:4], tuple(comparable(c.value) for c in expected.raw.cells)), 'Потеря исходных полей в Проверить.')
            require(actual[4:6] == (expected.raw.sheet, expected.raw.number) and isinstance(actual[6], str) and bool(actual[6]), 'Неверная идентичность или причина в Проверить.')
        row_records, changes, links = {}, {}, {}
        identities = {r.raw.row_id: r for r in rows}
        for record in audit_records():
            if record[0] == 'ROW':
                require(record[1] not in row_records, 'Повторяющаяся запись ROW.')
                row_records[record[1]] = record
            elif record[0] == 'FIELD_CHANGE':
                ident = (record[1], record[5])
                require(ident not in changes, 'Повторяющееся изменение поля.')
                require(record[1] in identities and record[5] in HEADERS, 'Неизвестное исходное поле аудита.')
                changes[ident] = record
            elif record[0] == 'DUPLICATE_LINKED':
                require(record[1] not in links, 'Повторяющаяся связь дубля.')
                links[record[1]] = record
        require(Counter(row_records.keys()) == Counter(actual_ids), 'Не совпала инвентаризация ROW.')
        require(set(links) == {r.raw.row_id for r in rows if r.disposition == 'DUPLICATE'}, 'Не совпали связи дублей.')
        expected_changes = set()
        for row in rows:
            record = row_records[row.raw.row_id]
            require(record[3] == row.output_row and record[4] == row.representative_id, 'Неверная судьба ROW.')
            require(same_json(unpack(record[9]), {'disposition': row.disposition, 'canonical_phone': row.fields[1].canonical,
                                         'codes': row.codes}), 'Неверные детали ROW.')
            if row.disposition == 'DUPLICATE':
                require(row.raw.row_id in links, 'Нет связи дубля.')
                link = links[row.raw.row_id]
                representative = identities[row.representative_id]
                require(link[3] is None and link[4] == row.representative_id and same_json(unpack(link[9]), {
                    'source': row.raw.cells[3].value, 'representative_source': representative.raw.cells[3].value,
                    'canonical_phone': row.fields[1].canonical}), 'Неверные детали связи дубля.')
            for i, result in enumerate(row.fields):
                cell = row.raw.cells[i]
                if type(cell.value) is not type(visible(result.display)) or cell.value != visible(result.display) or cell.kind in ('f', 'e', 'd') or cell.metadata.get('parse_issue'):
                    require((row.raw.row_id, HEADERS[i]) in changes, 'Нет аудита изменения поля.')
                    expected_changes.add((row.raw.row_id, HEADERS[i]))
                    change = changes[(row.raw.row_id, HEADERS[i])]
                    rule = 'SAFE_LITERAL' if cell.kind in ('f', 'e', 'd') or cell.metadata.get('parse_issue') else 'NORMALIZE_' + ['NAME', 'PHONE', 'DATE', 'SOURCE'][i]
                    require(change[2:7] == (cell.coordinate, row.output_row, row.representative_id, HEADERS[i], rule), 'Неверная прослеживаемость изменения.')
                    require(same_json(unpack(change[7]), {'type': cell.kind, 'value': cell.value}) and
                            same_json(unpack(change[8]), {'type': type(visible(result.display)).__name__, 'value': visible(result.display)}),
                            'Неверное исходное/итоговое значение FIELD_CHANGE.')
        require(set(changes) == expected_changes, 'Лишние изменения полей в аудите.')
    except (ValueError, TypeError, KeyError) as exc:
        raise NormalizationError('INVARIANT_FAILED', 'Не совпали длины или значения выходных строк.', 4) from exc
    finally:
        workbook.close()
