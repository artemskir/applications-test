"""Offline, independent public-resource corpora plus real XLSX/CLI reconciliation."""
from collections import Counter
from datetime import date
from io import BytesIO
import argparse
import hashlib
import json
from pathlib import Path
import random
import re
import subprocess
import struct
import sys
import time
import unicodedata
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from allpairspy import AllPairs
from hypothesis import given, seed, settings, strategies as st
from normalizer import dates, names, phone, SPEC_VERSION
from normalizer.models import HEADERS, SHEETS, RawCell
from normalizer.name_roles import classify, key
from normalizer.preservation import restore_bytes
from openpyxl import Workbook, load_workbook
from research.resource_sources import load_sources, locked_sources, LOCK
from research.probe_resource_cases import probe

NS='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


class CorpusQA:
    def __init__(self, output, sources, random_seed, examples):
        self.output, self.sources, self.seed, self.examples=output,sources,random_seed,examples
        self.counts=Counter(); self.failures=[]; self.observations={}; self.artifacts=[]
        self.generated=Counter(); self.unique={}

    def check(self, family, condition, example, failure_type='false_acceptance', details=None):
        self.counts[family]+=1
        if not condition:
            self.failures.append({'family':family,'input':example,'type':failure_type,'details':details})

    def analyze(self, family, analyzer, raw, expected=None, reject=False, **kwargs):
        try:
            result=analyzer.analyze(RawCell(raw, **kwargs))
            ok=result.canonical is None if reject else result.canonical==expected
            self.check(family,ok,raw,'false_acceptance' if reject else 'false_rejection',
                       {'canonical':result.canonical,'codes':result.codes,'expected':expected})
            self.check(family+'-raw',result.raw.value==raw and (result.canonical is not None or result.display==raw),raw,'data_loss')
            return result
        except Exception as exc:
            self.check(family,False,raw,'crash',repr(exc))

    def unicode(self):
        if unicodedata.unidata_version!='16.0.0':
            raise ValueError('This corpus is pinned to Unicode 16.0.0; runtime has '+unicodedata.unidata_version)
        records=[line.split(';') for line in self.sources['UnicodeData.txt'].splitlines() if line]
        nd={int(r[0],16):int(r[6]) for r in records if r[2]=='Nd'}
        blocks=sorted({cp-v for cp,v in nd.items()})
        for start in blocks:
            alphabet=[chr(start+i) for i in range(10)]
            for raw,expected,analyzer in (('+79174005219','+79174005219',phone),('09/03/2026','2026-03-09',dates)):
                variant=''.join(alphabet[int(c)] if c in '0123456789' else c for c in raw)
                self.analyze('unicode-decimal-blocks',analyzer,variant,expected)
        for cp,decimal in nd.items():
            self.analyze('unicode-every-Nd',phone,'+7917'+chr(cp)+'005219','+7917'+str(decimal)+'005219')
        numeric=[chr(int(r[0],16)) for r in records if r[2] in ('Nl','No')]
        for char in numeric:
            self.analyze('unicode-nondecimal',phone,'+7917'+char+'005219',reject=True)
            self.analyze('unicode-nondecimal',dates,'09/0'+char+'/2026',reject=True)
        marks=[chr(int(r[0],16)) for r in records if r[2].startswith('M')]
        controls=[chr(int(r[0],16)) for r in records if r[2] in ('Cf','Cc') and not chr(int(r[0],16)).isspace()]
        for family,chars in (('unicode-marks',marks),('unicode-controls',controls)):
            for char in chars:
                for pos in range(13):
                    raw='+79174005219'[:pos]+char+'+79174005219'[pos:]
                    self.analyze(family,phone,raw,reject=True)
                self.analyze(family,dates,'09/03'+char+'/2026',reject=True)
        self.marks, self.alphabets=marks,blocks
        self.observations['unicode_classes']={'Nd_blocks':len(blocks),'Nd_codepoints':len(nd),'marks':len(marks),
                                               'nondecimal_numeric':len(numeric),'nonspace_controls':len(controls)}
        conformant=0; cyrillic=0
        for line in self.sources['NormalizationTest.txt'].splitlines():
            data=line.split('#',1)[0].strip()
            if not data or data.startswith('@'): continue
            c=[''.join(chr(int(x,16)) for x in part.split()) for part in data.split(';')[:5]]
            ok=all(unicodedata.normalize('NFC',x)==c[1] for x in c[:3]) and all(unicodedata.normalize('NFC',x)==c[3] for x in c[3:])
            self.check('unicode-nfc-conformance',ok,c,'runtime_conformance')
            conformant+=1
            if any('\u0400'<=char<='\u052f' for char in c[0]):
                self.check('name-key-nfc',key(c[0])==c[1].casefold(),c,'false_rejection')
                cyrillic+=1
        confusable_count=0
        for line in self.sources['confusables.txt'].splitlines():
            data=line.split('#',1)[0].strip()
            if not data: continue
            fields=data.split(';'); src=fields[0].split(); dst=fields[1].split()
            if len(src)!=1 or len(dst)!=1 or chr(int(dst[0],16)) not in '0123456789': continue
            char=chr(int(src[0],16))
            if unicodedata.category(char)=='Nd' or char.isspace() or char in '().-‐‑‒–—−': continue
            self.analyze('unicode-confusables',phone,'+7917'+char+'005219',reject=True)
            confusable_count+=1
        self.observations['normalization']={'official_NFC_vectors':conformant,'cyrillic_key_vectors':cyrillic,
                                           'digit_confusable_mutations':confusable_count,'NFKC_used_by_processor':False}

    def calendar(self):
        # Literal whitelist is a spec oracle, not imported from production MONTHS.
        supported={
            'ru': ['январь января янв','февраль февраля фев февр','март марта мар','апрель апреля апр','май мая',
                   'июнь июня июн','июль июля июл','август августа авг','сентябрь сентября сен сент',
                   'октябрь октября окт','ноябрь ноября ноя нояб','декабрь декабря дек'],
            'en': ['january jan','february feb','march mar','april apr','may','june jun','july jul',
                   'august aug','september sep sept','october oct','november nov','december dec']}
        narrow_review=[]
        for locale in ('ru','en'):
            tree=json.loads(self.sources['cldr-'+locale+'.json'])['main'][locale]['dates']['calendars']['gregorian']['months']
            seen=set()
            for context in ('format','stand-alone'):
                for width in ('wide','abbreviated','narrow'):
                    for month,text in tree[context][width].items():
                        if (month,text) in seen: continue
                        seen.add((month,text))
                        raw='6 '+text+' 2026'; expected=date(2026,int(month),6).isoformat()
                        approved=text.lower().rstrip('.') in supported[locale][int(month)-1].split()
                        self.analyze('cldr-months-'+('approved' if approved else 'unsupported'),dates,raw,expected,reject=not approved)
                        if not approved: narrow_review.append({'locale':locale,'month':month,'text':text})
        self.observations['cldr_not_supported']=narrow_review
        for year in (1,1899,1900,2000,2024,2026,2100,9999):
            for month in range(1,13):
                for day in (1,28,29,30,31):
                    try: expected=date(year,month,day).isoformat()
                    except ValueError: expected=None
                    for sep in ('/','.','-',' ','\u202f'):
                        raw=sep.join((str(day),str(month),f'{year:04d}'))
                        self.analyze('calendar-boundaries',dates,raw,expected,reject=expected is None)

    def phones(self):
        tree=ET.fromstring(self.sources['PhoneNumberMetadata.xml'])
        examples={}
        for territory in tree.findall('.//territory'):
            if territory.get('id') not in ('RU','KZ'): continue
            values=sorted({e.text for e in territory.findall('.//exampleNumber')})
            examples[territory.get('id')]=values
            for value in values:
                # Metadata supplied the input. The agreed length/prefix supplied the oracle.
                if len(value)==10 and value.isascii() and value.isdecimal():
                    canonical='+7'+value
                    for variant in (canonical, '8 ('+value[:3]+') '+value[3:6]+'-'+value[6:8]+'-'+value[8:]):
                        self.analyze('metadata-phones',phone,variant,canonical)
                else:
                    self.observations.setdefault('metadata_examples_outside_contract',[]).append(value)
        self.observations['metadata_examples']=examples
        # Six metadata groupings include short local and long special forms.
        # Generate their digits independently; leadingDigits/operator allocation is irrelevant here.
        qa=self; patterns=[]
        territory=next(t for t in tree.findall('.//territory') if t.get('id')=='RU')
        for node in territory.findall('.//numberFormat'):
            pattern=node.get('pattern').replace(r'\d','[0-9]'); template=node.findtext('format')
            patterns.append({'pattern':pattern,'format':template})
            @seed(self.seed)
            @settings(max_examples=self.examples,deadline=None)
            @given(st.from_regex(pattern,fullmatch=True))
            def grouping(national):
                qa.generated['metadata-regex-groupings']+=1
                qa.unique.setdefault('metadata-regex-groupings',set()).add((pattern,national))
                match=re.fullmatch(pattern,national)
                grouped=re.sub(r'\$([0-9]+)',lambda m:match[int(m[1])],template)
                qa.analyze('metadata-regex-groupings',phone,'+7 '+grouped,'+7'+national,reject=len(national)!=10)
            grouping()
        self.observations['metadata_groupings']=patterns

    def name_gold(self):
        gold=json.loads((ROOT/'research/name_gold.json').read_text(encoding='utf-8'))
        observations=[]; correct=0
        for item in gold:
            actual,_=classify(item['text']); correct+=actual==item['label']
            result=names.analyze(RawCell(item['text']))
            observations.append({**item,'classified':actual,'canonical':result.canonical,'codes':result.codes})
            self.check('name-gold-preservation',result.raw.value==item['text'] and
                       (result.canonical is not None or result.display==item['text']),item['text'],'data_loss')
        self.observations['name_gold']={'correct':correct,'total':len(gold),'development_not_holdout':True,
                                       'cases':observations}

    def generated_properties(self):
        qa=self
        @seed(self.seed)
        @settings(max_examples=self.examples,deadline=None)
        @given(st.sampled_from(self.alphabets), st.text(alphabet='0123456789',min_size=10,max_size=10),
               st.sampled_from(self.marks), st.integers(min_value=0,max_value=12))
        def generated(block,national,mark,pos):
            qa.generated['unicode-phone-combinations']+=1
            qa.unique.setdefault('unicode-phone-combinations',set()).add((block,national,mark,pos))
            canonical='+7'+national
            variant=''.join(chr(block+int(c)) if c in '0123456789' else c for c in canonical)
            qa.analyze('hypothesis-resource-valid',phone,variant,canonical)
            broken=variant[:pos]+mark+variant[pos:]
            qa.analyze('hypothesis-resource-broken',phone,broken,reject=True)
        generated()

    def artifact(self,path):
        self.artifacts.append({'file':str(path.relative_to(self.output)),
                               'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'bytes':path.stat().st_size})

    def workbook(self):
        book=Workbook(); book.active.title='Заявки'; book.active.append(HEADERS)
        rows=[['Пётр','+79174005219','09\u202f03\u202f2026','сайт'],
              ['Пётр','8 (917) 400-52-19','2026-03-09','email'],
              ['Пётр','+79174005219 доб. 917-400-52-20','2026-03-09','telegram'],
              ['Анна','89174005222\ufe0f','2026-03-09','звонок'],
              ['Кирилл Зайцев','+79174005223',46271,'сайт'],
              ['Иван','+79174005224','2026-03-09','=1+1'],
              [None,None,None,'пустые обязательные поля']]
        for row in rows: book.active.append(row)
        book.active['C6'].number_format='yyyy-mm-dd'
        buf=BytesIO(); book.save(buf); book.close()
        return buf.getvalue(),rows

    def rewrite(self,payload,mode):
        stream=BytesIO()
        with ZipFile(BytesIO(payload)) as source, ZipFile(stream,'w',compression=ZIP_DEFLATED) as target:
            for item in source.infolist():
                value=source.read(item.filename)
                if item.filename=='xl/worksheets/sheet1.xml':
                    tree=ET.fromstring(value)
                    if mode=='rich':
                        cell=tree.find('.//'+NS+'c[@r="A2"]'); node=cell.find(NS+'is'); node.clear()
                        for text in ('Пё','тр'):
                            ET.SubElement(ET.SubElement(node,NS+'r'),NS+'t').text=text
                    elif mode=='shared':
                        cell=tree.find('.//'+NS+'c[@r="A2"]'); cell.remove(cell.find(NS+'is'))
                        cell.set('t','s'); ET.SubElement(cell,NS+'v').text='0'
                    value=ET.tostring(tree,encoding='utf-8')
                    if mode=='utf16': value=ET.tostring(tree,encoding='utf-16',xml_declaration=True)
                    if mode=='utf8-bom': value=b'\xef\xbb\xbf'+value
                    if mode=='bad-utf8':
                        value,count=re.subn(br'(<(?:[A-Za-z0-9_]+:)?sheetData>)',lambda m:m[1]+b'\xff',value,count=1)
                        if count!=1: raise ValueError('UTF-8 mutation was not applied')
                    if mode=='wrong-encoding': value=b'<?xml version="1.0" encoding="utf-16"?>'+value
                    if mode=='bad-entity':
                        value,count=re.subn(br'(<(?:[A-Za-z0-9_]+:)?sheetData>)',br'\1&doesnotexist;',value,count=1)
                        if count!=1: raise ValueError('Entity mutation was not applied')
                    if mode=='truncated': value=value[:-20]
                    if mode=='dtd': value=b'<!DOCTYPE worksheet [<!ENTITY qa "test">]>'+value
                    if mode=='false-dimension':
                        value,count=re.subn(br'(<(?:[A-Za-z0-9_]+:)?dimension\b[^>]*ref=")[^"]+',br'\g<1>A1:A1',value,count=1)
                        if count!=1: raise ValueError('Dimension mutation was not applied')
                    if mode=='missing-sheet': continue
                    if mode=='duplicate-cell':
                        tree=ET.fromstring(value); row=tree.find('.//'+NS+'row[@r="2"]')
                        row.append(ET.fromstring(ET.tostring(row.find(NS+'c')))); value=ET.tostring(tree)
                    if mode=='invalid-shared-index':
                        tree=ET.fromstring(value); cell=tree.find('.//'+NS+'c[@r="A2"]')
                        cell.remove(cell.find(NS+'is')); cell.set('t','s'); ET.SubElement(cell,NS+'v').text='999999'
                        value=ET.tostring(tree)
                if mode=='shared' and item.filename=='[Content_Types].xml':
                    tree=ET.fromstring(value)
                    ET.SubElement(tree,'{http://schemas.openxmlformats.org/package/2006/content-types}Override',
                                  PartName='/xl/sharedStrings.xml',ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml')
                    value=ET.tostring(tree)
                if mode=='shared' and item.filename=='xl/_rels/workbook.xml.rels':
                    tree=ET.fromstring(value)
                    ET.SubElement(tree,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
                                  Id='qaShared',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings',Target='sharedStrings.xml')
                    value=ET.tostring(tree)
                target.writestr(item.filename,value)
            if mode=='shared':
                tree=ET.Element(NS+'sst',count='1',uniqueCount='1')
                ET.SubElement(ET.SubElement(tree,NS+'si'),NS+'t').text='Пётр'
                target.writestr('xl/sharedStrings.xml',ET.tostring(tree))
        payload=stream.getvalue()
        if mode in ('encrypted-local','encrypted-both'):
            data=bytearray(payload); local=data.index(b'PK\x03\x04')+6
            struct.pack_into('<H',data,local,struct.unpack_from('<H',data,local)[0] | 1)
            if mode=='encrypted-both':
                central=data.index(b'PK\x01\x02')+8
                struct.pack_into('<H',data,central,struct.unpack_from('<H',data,central)[0] | 1)
            payload=bytes(data)
        return payload

    def cli_check(self,label,payload,rows=None,exit_code=0,**options):
        source=self.output/(label+'.xlsx'); output=self.output/(label+'-result.xlsx')
        source.write_bytes(payload); self.artifact(source)
        command=[sys.executable,str(ROOT/'normalize.py'),'--input',str(source),'--output',str(output)]
        for k,v in options.items(): command.extend(('--'+k.replace('_','-'),str(v)))
        try:
            run=subprocess.run(command,cwd=ROOT,capture_output=True,encoding='utf-8',timeout=30)
        except subprocess.TimeoutExpired:
            self.check('xlsx-cli',False,label,'timeout'); return
        (self.output/(label+'.log')).write_text(run.stdout+run.stderr,encoding='utf-8')
        self.check('xlsx-cli',run.returncode==exit_code,label,'controlled_refusal' if exit_code else 'crash',
                   {'expected_exit':exit_code,'actual_exit':run.returncode,'stderr':run.stderr})
        self.check('xlsx-input-preservation',source.read_bytes()==payload,label,'data_loss')
        if exit_code:
            self.check('xlsx-no-partial-publication',not output.exists(),label,'partial_publication'); return
        if not output.exists(): return
        self.artifact(output)
        book=load_workbook(output,read_only=True,data_only=False)
        try:
            self.check('xlsx-five-sheets',book.sheetnames==SHEETS,label,'structure')
            original=list(book['Оригинал'].values)[1:]
            self.check('xlsx-row-inventory',[r[0] for r in original]==['s1:r'+str(i) for i in range(2,len(rows)+2)],label,'data_loss')
            self.check('xlsx-raw-values',[list(r[3:7]) for r in original]==rows,label,'data_loss')
            audit=[r for r in book['Аудит'].values if r[0]=='ROW']
            self.check('xlsx-audit-inventory',Counter(r[1] for r in audit)==Counter(r[0] for r in original),label,'data_loss')
            review=list(book['Проверить'].values)[1:]
            self.check('xlsx-human-review',[list(r[:4]) for r in review]==[rows[i] for i in (2,3,5,6)],label,'data_loss')
            expected=[('Пётр','+79174005219','2026-03-09','сайт'),('Кирилл','+79174005223','2026-09-06','сайт')]
            self.check('xlsx-clean-data',list(book['Данные'].values)[1:]==expected,label,'false_acceptance')
            self.check('xlsx-no-active-formulas',all(c.data_type not in ('f','e') for ws in book for row in ws for c in row),label,'unsafe_value')
        finally: book.close()
        self.check('xlsx-byte-restoration',restore_bytes(output)[0]==payload,label,'data_loss')

    def xlsx(self):
        payload,rows=self.workbook()
        for mode in ('inline','rich','shared','utf16','utf8-bom','false-dimension'):
            self.cli_check(mode,payload if mode=='inline' else self.rewrite(payload,mode),rows)
        for mode in ('bad-utf8','wrong-encoding','bad-entity','truncated','missing-sheet','duplicate-cell','invalid-shared-index',
                     'dtd','encrypted-local','encrypted-both'):
            self.cli_check(mode,self.rewrite(payload,mode),exit_code=2)
        self.cli_check('uncompressed-limit',payload,exit_code=5,max_uncompressed_bytes=64)
        # Every pair is explicitly measured; approved source values vary independently.
        factors=[['NFC','NFD'],['ordinary','Nd','NBSP'],['slash','space','month'],['1900','1904'],['inline','rich','shared']]
        combinations=list(AllPairs(factors)); covered=set()
        for index,combo in enumerate(combinations):
            normalized_rows=[list(r) for r in rows]
            if combo[0]=='NFD':
                for i in (0,1): normalized_rows[i][0]=unicodedata.normalize('NFD',normalized_rows[i][0])
            if combo[1]=='Nd': normalized_rows[0][1]='＋７９１７４００５２１９'
            if combo[1]=='NBSP': normalized_rows[0][1]='8\u00a0917\u00a0400\u00a052\u00a019'
            normalized_rows[0][2]={'slash':'09/03/2026','space':'09\u202f03\u202f2026','month':'9 марта 2026'}[combo[2]]
            book=Workbook(); book.active.append(HEADERS)
            if combo[3]=='1904':
                from openpyxl.utils.datetime import CALENDAR_MAC_1904
                book.epoch=CALENDAR_MAC_1904; normalized_rows[4][2]=44809 # Independently known 1904 serial for 2026-09-06.
            for r in normalized_rows: book.active.append(r)
            book.active['C6'].number_format='yyyy-mm-dd'
            buffer=BytesIO(); book.save(buffer); book.close()
            data=buffer.getvalue()
            # Rich/shared helper's A2 value must match the normalized name variant.
            if combo[4]!='inline' and combo[0]=='NFD':
                data=self.rewrite(data,combo[4])
                normalized_rows[0][0]='Пётр' # Helper explicitly encodes the NFC scalar.
            elif combo[4]!='inline': data=self.rewrite(data,combo[4])
            self.cli_check('pair-'+str(index),data,normalized_rows)
            for i in range(len(combo)):
                for j in range(i+1,len(combo)): covered.add((i,combo[i],j,combo[j]))
        expected=sum(len(factors[i])*len(factors[j]) for i in range(len(factors)) for j in range(i+1,len(factors)))
        self.check('pairwise-coverage',len(covered)==expected,{'expected':expected,'covered':len(covered)},'coverage')
        self.observations['pairwise']={'rows':len(combinations),'covered_pairs':len(covered),'required_pairs':expected,
                                      'limitation':'Rich/shared A2 is written in NFC; NFD survives in the duplicate A3.'}

    def run(self):
        minimal=probe()
        for item in minimal: self.check('minimal-regressions',item['passed'],item['id'],'regression',item)
        self.unicode(); self.calendar(); self.phones(); self.name_gold(); self.generated_properties(); self.xlsx()


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--sources',type=Path,required=True)
    ap.add_argument('--output-dir',type=Path,required=True)
    ap.add_argument('--seed',type=int,default=20261008)
    ap.add_argument('--examples',type=int,default=300)
    args=ap.parse_args()
    if args.examples<1: ap.error('--examples must be positive')
    try: sources=load_sources(args.sources.resolve())
    except (OSError,ValueError) as exc:
        print('Offline resource check failed: '+str(exc),file=sys.stderr); return 1
    output=args.output_dir.resolve(); output.mkdir(parents=True,exist_ok=False)
    qa=CorpusQA(output,sources,args.seed,args.examples); started=time.perf_counter()
    try: qa.run()
    except Exception as exc: qa.check('runner',False,'run','crash',repr(exc))
    paths=[*sorted((ROOT/'normalizer').glob('*.py')),ROOT/'normalizer/data/name_lexicon.json.gz',ROOT/'normalizer/data/manifest.json',
           ROOT/'docs/SPECIFICATION.md',Path(__file__).resolve(),ROOT/'fixtures/resource_cases.py',ROOT/'research/resource_sources.py',
           ROOT/'research/name_gold.json',ROOT/'research/local-qa-requirements.txt',
           ROOT/'tools/fetch_qa_resources.py',ROOT/'tools/verify_artifacts.py',
           ROOT/'tests/test_resource_regressions.py',ROOT/'tests/test_resource_dependencies.py',
           ROOT/'fixtures/scenarios.py',ROOT/'tools/generate_data.py',LOCK]
    report={'spec_version':SPEC_VERSION,'seed':args.seed,'examples_limit':args.examples,'seconds':round(time.perf_counter()-started,3),
            'passed':not qa.failures,'checks':dict(qa.counts),'failed':qa.failures,'observations':qa.observations,
            'generated_calls':dict(qa.generated),'generated_unique':{k:len(v) for k,v in qa.unique.items()},
            'resource_lock_sha256':hashlib.sha256(LOCK.read_bytes()).hexdigest(),'resources':locked_sources(),
            'source_hashes':{str(p.relative_to(ROOT)).replace('\\','/'):hashlib.sha256(p.read_bytes()).hexdigest() for p in paths},
            'artifacts':qa.artifacts}
    (output/'results.json').write_text(json.dumps(report,ensure_ascii=True,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'passed':report['passed'],'checks':sum(qa.counts.values()),'failures':len(qa.failures),
                      'seconds':report['seconds'],'output':str(output)}))
    return int(bool(qa.failures))


if __name__=='__main__':
    raise SystemExit(main())
