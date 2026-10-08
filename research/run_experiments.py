"""Offline probes of Deep Research claims. This is not the production normalizer."""
import argparse
import base64
from collections import Counter
from datetime import date, datetime
import hashlib
import importlib.metadata
import json
import math
import os
from pathlib import Path
import re
import random
import shutil
import sys
import tempfile
import unicodedata
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_STORED

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from fixtures.scenarios import CASES, HEADERS
from tools.generate_data import fixture_rows, write_workbook
import dateparser
from hypothesis import find, given, settings, strategies as st
from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, from_excel
import phonenumbers as pn

BASE = datetime(2026, 10, 8, 12)
MONTHS = {word: index for index, forms in enumerate([
    'январь января янв', 'февраль февраля фев февр', 'март марта мар',
    'апрель апреля апр', 'май мая', 'июнь июня июн', 'июль июля июл',
    'август августа авг', 'сентябрь сентября сен сент', 'октябрь октября окт',
    'ноябрь ноября ноя нояб', 'декабрь декабря дек'], 1) for word in forms.split()}
STRICT = {'RELATIVE_BASE': BASE, 'DATE_ORDER': 'DMY', 'PREFER_LOCALE_DATE_ORDER': False,
          'STRICT_PARSING': True, 'REQUIRE_PARTS': ['day', 'month', 'year'],
          'RETURN_AS_TIMEZONE_AWARE': False}
CHECKS = []


def check(name, condition, detail=None):
    CHECKS.append({'name': name, 'passed': bool(condition), 'detail': detail})


def as_text(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def extract_tables(path):
    records, domain = [], None
    for line_number, line in enumerate(path.read_text(encoding='utf-8').splitlines(), 1):
        if line.startswith('## Российские телефоны:'):
            domain = 'Телефон'
        elif line.startswith('## Даты и XLSX:'):
            domain = 'Дата заявки'
        elif line.startswith('## Имена:'):
            domain = 'Имя'
        elif line.startswith('## Дедупликация,'):
            domain = None
        if not domain or not line.startswith('|'):
            continue
        cells = [x.strip() for x in line.strip().strip('|').split('|')]
        if len(cells) < 2 or cells[0] in ('Input', 'Вход') or re.fullmatch(r'[-:]+', cells[0]):
            continue
        literals = re.findall(r'`([^`]*)`', cells[0])
        concrete = len(literals) == 1 and not any(x in literals[0] for x in ('...', '…', '<', '>'))
        if any(x in cells[0] for x in ('Python ', 'numeric ', 'NaN', 'None', 'Excel blank')):
            concrete = False
        value = literals[0] if concrete else None
        if value is not None:
            if len(value) >= 2 and value[0] == value[-1] == '"':
                value = value[1:-1]
            value = value.replace(r'\t', '\t').replace(r'\n', '\n')
        records.append({'id': f'DR-L{line_number}', 'field': domain,
                        'source_line': line_number, 'source_input': cells[0],
                        'research_expectation': ' | '.join(cells[1:]),
                        'kind': 'concrete_text' if concrete else 'template_or_typed_description',
                        'value': value,
                        'existing_exact_case_ids': [c['id'] for c in CASES if c['field'] == domain and c['input'] == value] if concrete else []})
    return records


def candidate_shape(value):
    """Only a whole pure candidate. It deliberately does not extract from arbitrary text."""
    value = ''.join(str(unicodedata.decimal(c)) if unicodedata.category(c) == 'Nd'
                    else '+' if c == '＋' else c for c in value)
    if re.search(r'[^0-9+\s().\-‐‑‒–—−]', value):
        return None
    depth = 0
    for character in value:
        if character == '(':
            if depth:
                return None
            depth += 1
        elif character == ')':
            depth -= 1
            if depth < 0:
                return None
    if depth or '()' in value:
        return None
    token = re.sub(r'[\s().\-‐‑‒–—−]', '', value)
    if re.fullmatch(r'\+7[0-9]{10}', token):
        return token
    if re.fullmatch(r'[78][0-9]{10}', token):
        return '+7' + token[1:]
    if re.fullmatch(r'[0-9]{10}', token):
        return '+7' + token
    return None


def phone_info(number):
    return {'e164': pn.format_number(number, pn.PhoneNumberFormat.E164),
            'extension': number.extension, 'possible': pn.is_possible_number(number),
            'valid_range': pn.is_valid_number(number), 'region': pn.region_code_for_number(number)}


def phone_probe(value):
    result = {'input': value, 'naive_all_ascii_digits': re.sub(r'[^0-9]', '', value),
              'pure_candidate_shape': candidate_shape(value)}
    try:
        result['raw_parse'] = phone_info(pn.parse(value, 'RU'))
    except pn.NumberParseException as exc:
        result['raw_parse'] = {'error': str(exc)}
    for name, leniency in [('matcher_possible', pn.Leniency.POSSIBLE), ('matcher_valid', pn.Leniency.VALID)]:
        result[name] = [dict(phone_info(match.number), span=[match.start, match.end], raw=match.raw_string)
                        for match in pn.PhoneNumberMatcher(value, 'RU', leniency=leniency)]
    return result


def strict_calendar_probe(value):
    """Narrow, experimental whitelist. Unsupported is not a production reason code."""
    value = value.strip()
    match = re.fullmatch(r'([0-9]{1,2})\s*([./-])\s*([0-9]{1,2})\s*\2\s*([0-9]{4})', value)
    if match:
        day, _, month, year = match.groups()
    else:
        match = re.fullmatch(r'([0-9]{4})\s*([./-])\s*([0-9]{1,2})\s*\2\s*([0-9]{1,2})', value)
        if match:
            year, _, month, day = match.groups()
        else:
            match = re.fullmatch(r'([0-9]{1,2})\s+([а-яё]+)\.?\s+([0-9]{4})(?:\s+г\.|\s+года)?', value.casefold())
            if match and match[2] in MONTHS:
                day, word, year = match.groups()
                month = MONTHS[word]
            elif re.fullmatch(r'[0-9]{4}-[0-9]{2}-[0-9]{2}T[0-9]{2}:[0-9]{2}(?::[0-9]{2}(?:\.[0-9]+)?)?(?:Z|[+-][0-9]{2}:[0-9]{2})?', value):
                try:
                    time_match = re.search(r'T([0-9]{2}):([0-9]{2})(?::([0-9]{2}))?', value)
                    if int(time_match[1]) > 23 or int(time_match[2]) > 59 or (time_match[3] and int(time_match[3]) > 59):
                        raise ValueError('Invalid time; never roll into the next day')
                    parsed = datetime.fromisoformat(value)
                    # fromisoformat accepts offsets like +03:60; the contract does not.
                    offset = re.search(r'[+-]([0-9]{2}):([0-9]{2})$', value)
                    if offset and (int(offset[1]) > 23 or int(offset[2]) > 59):
                        raise ValueError('Invalid offset')
                    return {'value': parsed.date().isoformat(), 'status': 'ok',
                            'time_context': value.split('T', 1)[1]}
                except ValueError:
                    return {'value': None, 'status': 'invalid_calendar_or_time'}
            elif re.fullmatch(r'[0-9]{1,2}[./-][0-9]{1,2}[./-][0-9]{4} [0-9]{2}:[0-9]{2}', value):
                day_part, time_part = value.split(' ')
                outcome = strict_calendar_probe(day_part)
                try:
                    datetime.strptime(time_part, '%H:%M')
                except ValueError:
                    return {'value': None, 'status': 'invalid_calendar_or_time'}
                return outcome
            else:
                return {'value': None, 'status': 'outside_probe_allowlist'}
    try:
        parsed = date(int(year), int(month), int(day))
        return {'value': f'{parsed.year:04d}-{parsed.month:02d}-{parsed.day:02d}', 'status': 'ok'}
    except ValueError:
        return {'value': None, 'status': 'invalid_calendar_or_time'}


def date_probe(value):
    normal = dateparser.parse(value, languages=['ru', 'en'], settings={'RELATIVE_BASE': BASE, 'RETURN_AS_TIMEZONE_AWARE': False})
    strict = dateparser.parse(value, languages=['ru', 'en'], settings=STRICT)
    absolute = dateparser.parse(value, languages=['ru', 'en'], settings=dict(STRICT, PARSERS=['absolute-time']))
    return {'input': value, 'default': as_text(normal), 'strict_dmy': as_text(strict),
            'strict_absolute_only': as_text(absolute), 'calendar_allowlist_probe': strict_calendar_probe(value)}


def name_probe(value):
    collapsed = ' '.join(value.split())
    return {'input': value, 'title_result': collapsed.title(), 'nfc': unicodedata.normalize('NFC', value),
            'nfkc': unicodedata.normalize('NFKC', value),
            'invisible': [f'U+{ord(c):04X}' for c in value if unicodedata.category(c) in ('Cf', 'Cc') and not c.isspace()],
            'codepoints': [f'U+{ord(c):04X}' for c in value]}


def workbook_probes(stage):
    results = []
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    for epoch in (1900, 1904):
        wb = Workbook()
        wb.epoch = CALENDAR_WINDOWS_1900 if epoch == 1900 else CALENDAR_MAC_1904
        sheet = wb.active
        cells = [(0, 'yyyy-mm-dd'), (59, 'yyyy-mm-dd'), (60, 'yyyy-mm-dd'), (60.5, 'yyyy-mm-dd hh:mm'),
                 (61, 'yyyy-mm-dd'), (46271, 'General'), (46271, 'yyyy-mm-dd'),
                 (0.75, 'hh:mm'), ('=DATE(2026,9,3)', 'yyyy-mm-dd'), (float('nan'), 'General')]
        for index, (value, fmt) in enumerate(cells, 1):
            sheet.cell(index, 1, value).number_format = fmt
        path = stage / f'probe_epoch_{epoch}.xlsx'
        wb.save(path)
        wb.close()
        with ZipFile(path) as z:
            xml = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
            raw = {c.attrib['r']: {'type': c.attrib.get('t'), 'v': c.findtext('s:v', namespaces=ns),
                                  'f': c.findtext('s:f', namespaces=ns)} for c in xml.findall('.//s:c', ns)}
        normal, cached = load_workbook(path, data_only=False), load_workbook(path, data_only=True)
        try:
            for index, (_, fmt) in enumerate(cells, 1):
                cell = normal.active.cell(index, 1)
                results.append({'epoch': epoch, 'cell': cell.coordinate, 'format': fmt, 'xml': raw.get(cell.coordinate),
                                'reader_type': type(cell.value).__name__, 'reader_value': as_text(cell.value),
                                'data_only': as_text(cached.active.cell(index, 1).value)})
            check(f'epoch-{epoch}-general-remains-number', isinstance(normal.active['A6'].value, int))
            check(f'epoch-{epoch}-formula-not-evaluated', normal.active['A9'].data_type == 'f' and cached.active['A9'].value is None)
            if epoch == 1900:
                check('serial-59-60-reader-collision', normal.active['A2'].value == normal.active['A3'].value)
                check('serial-60-raw-is-preserved', raw['A3']['v'] == '60')
            else:
                check('serial-60-1904-valid', normal.active['A3'].value == datetime(1904, 3, 1))
        finally:
            normal.close()
            cached.close()
    # Check the real writer, including limits that its initial tests did not exercise.
    lengths = []
    for length in (32767, 32768):
        path = stage / f'length_{length}.xlsx'
        try:
            write_workbook(path, [['А' * length, '+79151451780', '2026-09-03', 'сайт']])
        except ValueError as exc:
            lengths.append({'input_length': length, 'status': 'rejected', 'message': str(exc), 'file_exists': path.exists()})
        else:
            wb = load_workbook(path)
            try:
                actual = len(wb.active['A2'].value)
            finally:
                wb.close()
            lengths.append({'input_length': length, 'stored_length': actual, 'status': 'unchanged' if actual == length else 'TRUNCATED'})
    check('generator-no-silent-cell-truncation', all(item['status'] != 'TRUNCATED' for item in lengths), lengths)
    # Safe text is checked using the actual existing generator writer.
    path = stage / 'safe_text.xlsx'
    write_workbook(path, [['Андрей', '+79151451780', '2026-09-03', '=1+1']])
    wb = load_workbook(path, data_only=False)
    try:
        check('generator-formula-like-text-stays-text', wb.active['D2'].value == '=1+1' and wb.active['D2'].data_type == 's')
    finally:
        wb.close()
    return results, lengths


def typed_description_probes(stage):
    fixtures = [
        ('DR-L96', 'Телефон', None), ('DR-L108', 'Телефон', 89151451780),
        ('DR-L193', 'Дата заявки', date(2026, 9, 3)),
        ('DR-L194', 'Дата заявки', datetime(2026, 9, 3, 14, 30)),
        ('DR-L195', 'Дата заявки', {'type': 'serial', 'value': 46268, 'number_format': 'yyyy-mm-dd'}),
        ('DR-L412', 'Имя', None), ('DR-L413', 'Имя', None), ('DR-L416', 'Имя', '\u00a0'),
    ]
    rows = []
    for _, field, value in fixtures:
        row = ['Андрей Волков', '+79151451780', '2026-09-03', 'research']
        row[HEADERS.index(field)] = value
        rows.append(row)
    write_workbook(stage / 'typed_descriptions.xlsx', rows)
    wb = load_workbook(stage / 'typed_descriptions.xlsx')
    observed = []
    try:
        for row_index, (source_id, field, value) in enumerate(fixtures, 2):
            cell = wb.active.cell(row_index, HEADERS.index(field) + 1)
            observed.append({'source_id': source_id, 'reader_type': type(cell.value).__name__, 'reader_value': as_text(cell.value)})
        check('typed-phone-is-int', isinstance(wb.active['B3'].value, int))
        check('typed-date-calendar-component', wb.active['C4'].value.date() == date(2026, 9, 3))
        check('typed-datetime-time-retained', wb.active['C5'].value == datetime(2026, 9, 3, 14, 30))
        check('typed-date-serial-september-3', wb.active['C6'].value == datetime(2026, 9, 3))
    finally:
        wb.close()
    try:
        write_workbook(stage / 'generator_nan.xlsx', [['Иван', float('nan'), '2026-09-03', 'research']])
    except ValueError:
        nonfinite = 'rejected'
    else:
        nonfinite = 'written'
    check('generator-no-nan-to-blank', nonfinite == 'rejected')
    return {'records': observed, 'nan_description_source': 'DR-L417 and DR-L96',
            'nan_reader_observation': 'See probe_epoch_1900.xlsx A10: native NaN becomes empty numeric XML',
            'generator_nan_status': nonfinite}


def properties(stage):
    counters = Counter()
    @settings(max_examples=300, derandomize=True, database=None, deadline=None)
    @given(st.integers(min_value=0, max_value=9999999), st.sampled_from([' ', '-', '–', '\u00a0', '\u202f', '.']))
    def phone_format_property(tail, separator):
        nsn = f'915{tail:07d}'
        mutation = separator.join(['8', nsn[:3], nsn[3:6], nsn[6:8], nsn[8:]])
        assert candidate_shape(mutation) == '+7' + nsn
        assert candidate_shape(candidate_shape(mutation)) == '+7' + nsn
        counters['phone_format_examples'] += 1
    phone_format_property()
    @settings(max_examples=300, derandomize=True, database=None, deadline=None)
    @given(st.dates(min_value=date(1, 1, 1), max_value=date(9999, 12, 31)), st.sampled_from(['/', '.', '-']))
    def calendar_round_trip(value, separator):
        text = f'{value.day}{separator}{value.month}{separator}{value.year:04d}'
        parsed = strict_calendar_probe(text)
        assert parsed['value'] == value.isoformat()
        assert strict_calendar_probe(parsed['value'])['value'] == parsed['value']
        counters['calendar_examples'] += 1
    calendar_round_trip()
    @settings(max_examples=300, derandomize=True, database=None, deadline=None)
    @given(st.text(alphabet=['е', 'ё', '\u0308', 'И', 'В', ' ', '\u00a0'], min_size=0, max_size=60))
    def nfc_property(value):
        original = value.encode('utf-8')
        canonical = unicodedata.normalize('NFC', value)
        assert unicodedata.normalize('NFC', canonical) == canonical
        assert value.encode('utf-8') == original
        counters['nfc_examples'] += 1
    nfc_property()
    example = find(st.integers(0, 23), lambda hour: re.sub(r'\D', '', f'8 917 400 52 19 (после {hour:02}:00)') != '89174005219',
                   settings=settings(derandomize=True, database=None, deadline=None))
    check('property-phone-format', counters['phone_format_examples'] == 300)
    check('property-calendar-round-trip', counters['calendar_examples'] == 300)
    check('property-nfc-idempotence', counters['nfc_examples'] == 300)
    return dict(counters), f'8 917 400 52 19 (после {example:02}:00)'


def preservation(stage, source):
    payload = source.read_bytes()
    encoded = base64.b64encode(payload).decode('ascii')
    chunks = [encoded[i:i + 30000] for i in range(0, len(encoded), 30000)]
    path = stage / 'preservation_probe.xlsx'
    wb = Workbook()
    wb.active.title = 'Аудит'
    wb.active.append(['index', 'payload'])
    for index, chunk in enumerate(chunks):
        wb.active.append([index, chunk])
    wb.save(path)
    wb.close()
    wb = load_workbook(path)
    try:
        reconstructed = base64.b64decode(''.join(r[1] for r in list(wb.active.values)[1:]), validate=True)
    finally:
        wb.close()
    check('xlsx-base64-byte-round-trip', reconstructed == payload)
    return {'source_bytes': len(payload), 'chunk_count': len(chunks),
            'source_sha256': hashlib.sha256(payload).hexdigest(), 'restored_sha256': hashlib.sha256(reconstructed).hexdigest()}


def mutations(stage):
    rng = random.Random(42)
    base = ['Андрей Волков', '+79151451780', '2026-09-03', 'research']
    phone_formats = ['+79151451780', '89151451780', '9151451780', '8 (915) 145-17-80',
                     '8\u00a0915\u202f145–17–80', 'тел. 8 (915) 145-17-80',
                     '8 915 145 17 80 (после 18:00)', '+7 915 145-17-80 доб. 123']
    date_formats = ['2026-09-03', '03.09.2026', '3/9/2026', '3-9-2026', '3 сентября 2026',
                    '3 сен. 2026', '2026-09-03T14:00:00', '03.09.2026 14:00']
    name_formats = ['андрей волков', 'АНДРЕЙ ВОЛКОВ', 'Андрей  Волков', ' Андрей\u00a0Волков ']
    rows, labels = [], []
    for index in range(300):
        raw = [rng.choice(name_formats), rng.choice(phone_formats), rng.choice(date_formats), base[3]]
        rows.append(raw)
        labels.append({'excel_row': index + 2, 'raw': raw, 'canonical_basis': base,
                       'mutated_fields': [HEADERS[i] for i in range(3) if raw[i] != base[i]],
                       'note': 'Reversible formatting only. Oracle is canonical basis, not parser output.'})
    write_workbook(stage / 'mutations.xlsx', rows)
    (stage / 'mutations.json').write_text(json.dumps(labels, ensure_ascii=False, indent=2), encoding='utf-8')
    wb = load_workbook(stage / 'mutations.xlsx')
    try:
        check('compound-mutations-xlsx-raw-round-trip', list(wb.active.values)[1:] == [tuple(row) for row in rows])
    finally:
        wb.close()
    return {'rows': len(rows), 'seed': 42, 'compound_rows': sum(len(label['mutated_fields']) > 1 for label in labels),
            'scope': 'Generator preservation only; semantic normalization not implemented'}


def duplicate_probe():
    path = ROOT / 'generated/final-suite/manifest.json'
    manifest = json.loads(path.read_text(encoding='utf-8'))
    entry = next(e for e in manifest['files'] if e['file'] == 'duplicates.xlsx')
    wb = load_workbook(path.parent / entry['file'])
    try:
        rows = list(wb.active.values)[1:]
    finally:
        wb.close()
    groups, winners, duplicates, issues = {}, [], [], set()
    for index, row in enumerate(rows, 2):
        # These fixtures contain pure candidates; this does not test arbitrary extraction.
        key = candidate_shape(row[1]) if isinstance(row[1], str) else None
        if not row[0] or not key or strict_calendar_probe(row[2] or '')['status'] != 'ok':
            issues.add(index)
        if key is None:
            winners.append(index)
        elif key in groups:
            duplicates.append(index)
            groups[key].append((index, row))
        else:
            winners.append(index)
            groups[key] = [(index, row)]
    conflicts = set()
    for group in groups.values():
        names = {unicodedata.normalize('NFC', ' '.join((row[0] or '').split())).title() for _, row in group}
        dates = {strict_calendar_probe(row[2] or '')['value'] or row[2] for _, row in group}
        if len(names) > 1 or len(dates) > 1:
            conflicts.update(index for index, _ in group)
    issues |= conflicts
    expected = entry['expected']
    check('dedup-first-winners', winners == expected['data_input_rows'])
    check('dedup-suppressed-rows', duplicates == expected['duplicate_rows'])
    check('dedup-all-conflict-participants', sorted(conflicts) == expected['duplicate_conflict_rows'])
    check('dedup-problems-exact', sorted(issues) == expected['problem_rows'])
    check('dedup-traceability', len(winners) + len(duplicates) == len(rows))
    return {'input_count': len(rows), 'winners': winners, 'duplicates': duplicates, 'conflicts': sorted(conflicts),
            'problems': sorted(issues), 'scope': 'Experimental group model; production dedup module absent'}


def run(stage):
    source = ROOT / 'research/deep-research-report.md'
    coverage = extract_tables(source)
    supplementary = [
        ('Телефон', '+7 (495) 587-88-89, доб. 10-17, 10-13', 'Concrete extension-list fragment from Research prose; original web citation not resolved'),
        ('Телефон', '+7 495 123-45-67 x123', 'Concrete specialization of x123 template'),
        ('Телефон', '+7 915 145-17-80 после 18:00, доб. 42', 'Concrete specialization of combined-note template'),
        ('Телефон', '+7 915 111-11-11 / +7 916 222-22-22', 'Golden corpus'),
        ('Телефон', '+7 915 145-17-80; +7 915 145-17-80', 'Same distinct number, not two different numbers'),
        ('Телефон', '+7 915 145-17-80 / 123', 'One valid number plus malformed candidate'),
        ('Телефон', '+7 917 400 c5219', 'Letter deletion danger'),
        ('Имя', 'е\u0308лкина анна', 'Decomposed ё'),
        ('Имя', 'ёлкина анна', 'Composed ё'),
        ('Имя', 'Иван\u200bИванов', 'ZWSP described in prose'),
        ('Имя', '① Иван', 'Compatibility folding danger'),
    ]
    kz = pn.example_number_for_type('KZ', pn.PhoneNumberType.MOBILE)
    supplemental_kz = pn.format_number(kz, pn.PhoneNumberFormat.E164)
    supplementary.append(('Телефон', supplemental_kz, 'Metadata example for KZ; specialization of KZ template'))
    template_materializations = {
        'DR-L79': ('Телефон', '+7 915 145-17-80 ext 123'),
        'DR-L80': ('Телефон', '+7 915 145-17-80 x123'),
        'DR-L81': ('Телефон', '+7 915 111-11-11; +7 916 222-22-22'),
        'DR-L82': ('Телефон', '8 915 111-11-11 / 8 916 222-22-22'),
        'DR-L83': ('Телефон', '8 915 111-11-11 или 8 916 222-22-22'),
        'DR-L84': ('Телефон', '+7 915 145-17-80 после 18:00, доб. 42'),
        'DR-L91': ('Телефон', '++7 915 145-17-80'),
        'DR-L92': ('Телефон', '0079151451780'),
        'DR-L104': ('Телефон', '＋７９１５１４５１７８０'),
        'DR-L105': ('Телефон', '8 9О3 211-82-41'),
        'DR-L106': ('Телефон', '8 9I5 145-17-80'),
        'DR-L107': ('Телефон', '8 9З5 145-17-80'),
    }
    for source_id, (field, value) in template_materializations.items():
        supplementary.append((field, value, f'Explicit specialization of template {source_id}; digits supplied by experiment, not recovered from source'))
    for index, (field, value, note) in enumerate(supplementary, 1):
        coverage.append({'id': f'SUP-{index:02}', 'field': field, 'value': value, 'kind': 'explicit_supplement', 'note': note,
                         'source_line': None, 'existing_exact_case_ids': []})
    for item in coverage:
        if item['id'] in template_materializations:
            item['probe_reference'] = 'Explicit specialization in supplementary records, not a literal from Research'
        elif item['kind'] == 'template_or_typed_description':
            item['probe_reference'] = 'typed_description_probes or workbook_probes; NaN is a Python object, not a native XLSX scalar'
    phones, dates, names = [], [], []
    for item in coverage:
        if item['value'] is None:
            continue
        probe = phone_probe if item['field'] == 'Телефон' else date_probe if item['field'] == 'Дата заявки' else name_probe
        observed = dict(probe(item['value']), id=item['id'], source_line=item['source_line'])
        (phones if item['field'] == 'Телефон' else dates if item['field'] == 'Дата заявки' else names).append(observed)
    check('dmy-user-example', strict_calendar_probe('09/03/2026')['value'] == '2026-03-09')
    for invalid in ['31.04.2026', '31.09.2026', '29.02.2026', '2026-13-10', '00.09.2026']:
        check('calendar-reject-' + invalid, strict_calendar_probe(invalid)['status'] == 'invalid_calendar_or_time')
    check('calendar-leap-2024', strict_calendar_probe('29.02.2024')['value'] == '2024-02-29')
    check('structure-does-not-prove-RU', candidate_shape(supplemental_kz) is not None and pn.region_code_for_number(kz) == 'KZ')
    check('nfc-equivalence', unicodedata.normalize('NFC', 'е\u0308лкина') == 'ёлкина')
    rows = []
    for item in coverage:
        if item['value'] is not None:
            row = ['Андрей Волков', '+79151451780', '2026-09-03', 'research']
            row[HEADERS.index(item['field'])] = item['value']
            rows.append(row)
    write_workbook(stage / 'tables.xlsx', rows)
    excel, lengths = workbook_probes(stage)
    typed_descriptions = typed_description_probes(stage)
    props, minimized = properties(stage)
    # A real readable XLSX with an uncompressed entry exercises several chunk boundaries.
    chunk_source = stage / 'chunk_source.xlsx'
    shutil.copyfile(stage / 'tables.xlsx', chunk_source)
    with ZipFile(chunk_source, 'a', compression=ZIP_STORED) as archive:
        archive.writestr('probe-padding.bin', bytes(range(256)) * 400)
    raw_preservation = preservation(stage, chunk_source)
    check('preservation-exercises-multiple-chunks', raw_preservation['chunk_count'] > 1)
    mutation_result = mutations(stage)
    duplicates = duplicate_probe()
    extension_records = [phone_probe('+7 495 123-45-67 доб. 101'), phone_probe('+7 495 123-45-67 доб. 202')]
    check('same-main-different-extensions', extension_records[0]['raw_parse']['e164'] == extension_records[1]['raw_parse']['e164']
          and extension_records[0]['raw_parse']['extension'] != extension_records[1]['raw_parse']['extension'])
    dirty_rows, _ = fixture_rows(len(CASES), 20261008, 'dirty')
    phone_keys = [c['canonical'] for c in CASES if c['field'] == 'Телефон' and c.get('canonical') is not None]
    repeated = {key: n for key, n in Counter(phone_keys).items() if n > 1}
    results = {'experiment_date': '2026-10-08', 'timezone': 'Europe/Moscow', 'relative_base': BASE.isoformat(),
               'scope': 'Offline library and generator probes; production normalizer/restore absent',
               'research_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
               'spec_sha256': hashlib.sha256((ROOT / 'docs/SPECIFICATION.md').read_bytes()).hexdigest(),
               'versions': {p: importlib.metadata.version(p) for p in ['openpyxl', 'phonenumbers', 'dateparser', 'hypothesis']},
               'coverage_counts': dict(Counter(i['kind'] for i in coverage)),
               'phones': phones, 'dates': dates, 'names': names, 'excel': excel, 'lengths': lengths, 'typed_descriptions': typed_descriptions,
               'properties': props, 'minimized_digit_strip_counterexample': minimized,
               'preservation': raw_preservation, 'duplicates': duplicates, 'extension_records': extension_records,
               'mutations': mutation_result,
               'generator': {'dirty_count': len(dirty_rows), 'repeated_field_oracle_phone_keys': repeated,
                             'note': 'Field oracles are not whole-row/integration oracles'},
               'checks': CHECKS, 'passed': sum(c['passed'] for c in CHECKS), 'failed': sum(not c['passed'] for c in CHECKS)}
    (stage / 'coverage.json').write_text(json.dumps(coverage, ensure_ascii=False, indent=2), encoding='utf-8')
    (stage / 'results.json').write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding='utf-8')
    lines = ['# Наблюдения экспериментального стенда', '',
             'Полный анализ методологии: `research/EXPERIMENT_REPORT.md`. Эти таблицы — наблюдения библиотек, не результат рабочего нормализатора.', '',
             f"Проверки: {results['passed']} успешно, {results['failed']} неуспешно. Свойства: {props}.", '',
             '## Телефоны', '', '| ID | Вход | parse(raw) | possible matcher | valid matcher |', '|---|---|---|---|---|']
    def md(value):
        return json.dumps(value, ensure_ascii=False).replace('|', '&#124;')
    for item in phones:
        lines.append(f"| {item['id']} | {md(item['input'])} | {md(item['raw_parse'])} | {md(item['matcher_possible'])} | {md(item['matcher_valid'])} |")
    lines += ['', '## Даты', '', '| ID | Вход | default | strict DMY | strict absolute | календарный allowlist probe |', '|---|---|---|---|---|---|']
    for item in dates:
        lines.append(f"| {item['id']} | {md(item['input'])} | {md(item['default'])} | {md(item['strict_dmy'])} | {md(item['strict_absolute_only'])} | {md(item['calendar_allowlist_probe'])} |")
    lines += ['', '## Имена', '', '| ID | Вход | title() | NFC | invisible |', '|---|---|---|---|---|']
    for item in names:
        lines.append(f"| {item['id']} | {md(item['input'])} | {md(item['title_result'])} | {md(item['nfc'])} | {md(item['invisible'])} |")
    lines += ['', '## XLSX', '', '```json', json.dumps(excel, ensure_ascii=False, indent=2), '```', '', '## Пределы длины', '',
              '```json', json.dumps(lengths, ensure_ascii=False, indent=2), '```']
    (stage / 'observations.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    return results


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', required=True, type=Path)
    args = parser.parse_args()
    destination = args.output_dir.resolve()
    if destination.exists():
        parser.error('output-dir must be new')
    destination.parent.mkdir(parents=True, exist_ok=True)
    stage = Path(tempfile.mkdtemp(prefix='.research-stage-', dir=destination.parent)).resolve()
    try:
        result = run(stage)
        if destination.exists():
            raise FileExistsError(destination)
        os.rename(stage, destination)
    finally:
        if stage.exists() and stage.parent == destination.parent and stage.name.startswith('.research-stage-'):
            shutil.rmtree(stage)
    print(json.dumps({'output': str(destination), 'passed': result['passed'], 'failed': result['failed'],
                      'coverage': result['coverage_counts'], 'properties': result['properties']}, ensure_ascii=False))
    return 1 if result['failed'] else 0


if __name__ == '__main__':
    raise SystemExit(main())
