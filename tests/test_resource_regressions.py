"""Small persistent regressions; public-resource downloads are never needed here."""
from io import BytesIO
from pathlib import Path
import tempfile
import struct
import unittest
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from fixtures.resource_cases import CASES
from normalizer import dates, names, phone
from normalizer.models import HEADERS, SHEETS, RawCell, NormalizationError
from normalizer.pipeline import normalize
from normalizer.preservation import restore_bytes


class ResourceRegressions(unittest.TestCase):
    def test_encryption_flag_in_either_zip_header_rejects_without_publication(self):
        book=Workbook(); book.active.append(HEADERS)
        book.active.append(['Иван','+79174005219','2026-09-06','сайт'])
        stream=BytesIO(); book.save(stream); book.close()
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            for mode in ('local','central','both'):
                payload=bytearray(stream.getvalue())
                if mode in ('local','both'):
                    local=payload.index(b'PK\x03\x04')+6
                    struct.pack_into('<H',payload,local,struct.unpack_from('<H',payload,local)[0] | 1)
                if mode in ('central','both'):
                    central=payload.index(b'PK\x01\x02')+8
                    struct.pack_into('<H',payload,central,struct.unpack_from('<H',payload,central)[0] | 1)
                source=root/(mode+'.xlsx'); output=root/(mode+'-result.xlsx')
                source.write_bytes(payload)
                with self.subTest(mode=mode):
                    with self.assertRaises(NormalizationError) as error: normalize(source,output)
                    self.assertEqual(error.exception.code,'INPUT_INVALID')
                    self.assertFalse(output.exists())
                    self.assertEqual(source.read_bytes(),bytes(payload))

    def test_independent_agent_and_boundary_cases(self):
        for case in CASES:
            with self.subTest(case=case.ident):
                result = {'phone':phone, 'dates':dates, 'names':names}[case.field_name].analyze(RawCell(case.value, **case.kwargs))
                self.assertEqual(result.canonical, case.canonical)
                self.assertEqual(result.codes, list(case.codes))
                self.assertEqual(result.display, case.value if case.canonical is None else case.canonical)

    def test_unicode_marks_cannot_modify_digits_at_any_position(self):
        for mark in ('\u0301', '\ufe0f', '\u20e3', '\u0903', '\U000e0100'):
            for pos in range(12):
                with self.subTest(mark=ascii(mark), pos=pos):
                    raw = '89174005219'[:pos] + mark + '89174005219'[pos:]
                    result = phone.analyze(RawCell(raw))
                    self.assertIsNone(result.canonical)
                    self.assertEqual(result.codes, ['PHONE_INVALID'])
                    self.assertEqual(result.display, raw)

    def test_plain_space_dates_follow_calendar_and_order_contract(self):
        for sep in (' ', '\t', '\n', '\u00a0', '\u202f', '\u3000'):
            for tokens, expected in ((('09','03','2026'),'2026-03-09'),
                                     (('09','25','2026'),'2026-09-25'),
                                     (('2026','09','06'),'2026-09-06')):
                with self.subTest(sep=ascii(sep), tokens=tokens):
                    result=dates.analyze(RawCell(sep.join(tokens)))
                    self.assertEqual(result.canonical, expected)
            for raw in (sep.join(('31','02','2026')), sep.join(('09','03','26'))):
                self.assertIsNone(dates.analyze(RawCell(raw)).canonical)
        for raw in ('2026 25 09', '09 / 03 2026', '09 03 / 2026', '09 03 2026abc'):
            self.assertIsNone(dates.analyze(RawCell(raw)).canonical)

    def test_reject_other_full_phone_in_extension_under_every_marker(self):
        for marker in ('доб.', 'вн.', 'доп.', '*', '#', ';ext=', 'x'):
            for second in ('917-400-52-20', '8-917-400-52-20', '+79174005220'):
                with self.subTest(marker=marker,second=second):
                    result=phone.analyze(RawCell('+79174005219 '+marker+' '+second))
                    self.assertIsNone(result.canonical)
                    self.assertEqual(result.codes, ['PHONE_MULTIPLE'])
            for second in ('00-12-34-56-78', '810-123-45-678'):
                with self.subTest(marker=marker,second=second):
                    result=phone.analyze(RawCell('+79174005219 '+marker+' '+second))
                    self.assertIsNone(result.canonical)
                    self.assertEqual(result.codes, ['PHONE_UNSUPPORTED_PREFIX'])

    def test_five_sheets_do_not_hide_bad_duplicate_behind_extension(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory)/'input.xlsx', Path(directory)/'output.xlsx'
            rows = [
                ['Иван', '+79174005219', '2026-03-09', 'сайт'],
                ['Иван', '+79174005219 доб. 917-400-52-20', '2026-03-09', 'email'],
                ['Кирилл', '+79174005221', '09\u202f25\u202f2026', 'telegram'],
                ['Анна', '89174005222\ufe0f', '2026-03-09', 'звонок'],
            ]
            book=Workbook(); book.active.append(HEADERS)
            for row in rows: book.active.append(row)
            buf=BytesIO(); book.save(buf); book.close()
            before=buf.getvalue(); source.write_bytes(before)
            summary=normalize(source, output)
            self.assertEqual((summary['input_rows'],summary['data_rows'],summary['problem_rows'],summary['duplicates']), (4,2,2,0))
            result=load_workbook(output, read_only=True, data_only=False)
            try:
                self.assertEqual(result.sheetnames, SHEETS)
                self.assertEqual([list(r[3:7]) for r in list(result['Оригинал'].values)[1:]], rows)
                self.assertEqual([list(r[:4]) for r in list(result['Проверить'].values)[1:]], [rows[1],rows[3]])
                self.assertEqual(list(result['Данные'].values)[2][2], '2026-09-25')
                self.assertEqual(len([r for r in result['Аудит'].values if r[0]=='ROW']),4)
            finally: result.close()
            self.assertEqual(restore_bytes(output)[0], before)
            self.assertEqual(source.read_bytes(), before)
