"""Independent reproductions from the user's QA review and approved 1.4 decisions."""
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from normalizer import dates, phone, names
from normalizer.models import HEADERS, RawCell
from normalizer.reader import read_rows
from normalizer.pipeline import normalize
from normalizer.preservation import restore_bytes


class ReviewRegressions(unittest.TestCase):
    def test_nonfinite_numeric_lexeme_in_each_column_keeps_all_rows(self):
        ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            book=Workbook(); book.active.append(HEADERS)
            book.active.append(['Иван','+79174005219','2026-09-06','сайт'])
            book.active.append(['Кирилл','+79174005220','2026-09-06','email'])
            buf=BytesIO(); book.save(buf)
            for column in 'ABCD':
                for lexeme in ('NaN','INF','-INF','1e9999'):
                    with self.subTest(column=column,lexeme=lexeme):
                        source,output=root/f'{column}-{lexeme}.xlsx',root/f'{column}-{lexeme}-result.xlsx'
                        with ZipFile(BytesIO(buf.getvalue())) as old,ZipFile(source,'w') as new:
                            for info in old.infolist():
                                payload=old.read(info.filename)
                                if info.filename=='xl/worksheets/sheet1.xml':
                                    tree=ET.fromstring(payload)
                                    cell=tree.find(f'.//{ns}c[@r="{column}2"]'); cell.clear(); cell.set('r',column+'2'); cell.set('t','n')
                                    ET.SubElement(cell,ns+'v').text=lexeme
                                    payload=ET.tostring(tree)
                                new.writestr(info.filename,payload)
                        before=source.read_bytes()
                        summary=normalize(source,output)
                        self.assertEqual((summary['input_rows'],summary['data_rows'],summary['problem_rows']),(2,1,1))
                        self.assertEqual(summary['codes']['CELL_NUMERIC_INVALID'],1)
                        parsed=load_workbook(output,read_only=True)
                        original=list(parsed['Оригинал'].values)[1]
                        self.assertEqual(original[3+'ABCD'.index(column)],lexeme)
                        metadata=json.loads(original[7])[HEADERS['ABCD'.index(column)]]
                        self.assertEqual(metadata['ooxml_type'],'n')
                        self.assertEqual(metadata['xml_value'],lexeme)
                        self.assertEqual(metadata['parse_issue'],'CELL_NUMERIC_INVALID')
                        self.assertTrue(any(r[0]=='FIELD_CHANGE' and r[5]==HEADERS['ABCD'.index(column)] and r[6]=='SAFE_LITERAL' for r in parsed['Аудит'].values))
                        parsed.close()
                        self.assertEqual(restore_bytes(output)[0],before)
                        self.assertEqual(source.read_bytes(),before)

    def test_personal_name_extraction_and_conservative_abstention(self):
        for raw, expected in [('Кирилл','Кирилл'), ('Зайцев Кирилл','Кирилл'), ('Кирилл Зайцев','Кирилл'),
                              ('Ирина Львовна','Ирина'), ('Волков Егор','Егор'), ('ПАВЛОВА анастасия','Анастасия'),
                              ('Кирилл Зайцев Андреевич','Кирилл'), ('Жан-Поль','Жан-Поль')]:
            with self.subTest(raw=raw):
                result = names.analyze(RawCell(raw))
                self.assertEqual(result.canonical, expected)
                self.assertEqual(result.codes, [])
        for raw, code in [('К. Белоусов','NAME_INCOMPLETE'), ('Волков','NAME_AMBIGUOUS'),
                          ('Андреевич','NAME_AMBIGUOUS'), ('Влад','NAME_AMBIGUOUS'), ('Ким','NAME_AMBIGUOUS'), ('Али','NAME_AMBIGUOUS'),
                          ('Ли','NAME_AMBIGUOUS'), ('Клиент Кирилл','NAME_UNKNOWN'),
                          ('не знаю','NAME_UNKNOWN'), ('Даздраперма','NAME_UNKNOWN'),
                          ('Кирилл Ирина','NAME_AMBIGUOUS')]:
            with self.subTest(raw=raw):
                result = names.analyze(RawCell(raw))
                self.assertIsNone(result.canonical)
                self.assertEqual(result.display, raw)
                self.assertEqual(result.codes, [code])

    def test_different_surnames_cannot_hide_behind_same_extracted_name(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory)/'input.xlsx', Path(directory)/'output.xlsx'
            book = Workbook(); book.active.append(HEADERS)
            for name in ('Кирилл Зайцев','Кирилл Волков'):
                book.active.append([name,'+79174005219','2026-09-06','сайт'])
            book.save(source)
            result = normalize(source,output)
            self.assertEqual((result['data_rows'],result['duplicates'],result['problem_rows']),(1,1,2))
            self.assertEqual(result['codes']['DUPLICATE_NAME_CONFLICT'],2)
            self.assertEqual(restore_bytes(output)[0],source.read_bytes())

    def test_incomplete_serial_formats_never_invent_full_date(self):
        for fmt in ('mm', 'yyyy', 'dd', 'mmm yyyy', 'dd.mm', 'dddd mmmm yyyy',
                    'hh:mm', '[h]:mm', '"dd" mm yyyy', 'dd "mm" yyyy', 'dd mm "yyyy"',
                    r'\d\d mm yyyy', '[Red]mm', 'mm;dd.mm.yyyy',
                    'dd yyyy hh:mm', 'dd yyyy mm:ss', 'hh:mm dd yyyy'):
            with self.subTest(format=fmt):
                cell = RawCell(46271, 'n', number_format=fmt)
                result = dates.analyze(cell)
                self.assertIsNone(result.canonical)
                self.assertEqual(result.display, 46271)
                self.assertEqual(result.codes, ['DATE_INCOMPLETE'])
        for fmt in ('yyyy-mm-dd', 'dd.mm.yyyy', 'd/m/yy', 'ddd, mmm d, yyyy',
                    '[Red]dd.mm.yyyy', 'dd.mm.yyyy hh:mm'):
            with self.subTest(format=fmt):
                self.assertEqual(dates.analyze(RawCell(46271, 'n', number_format=fmt)).canonical, '2026-09-06')

    def test_real_numeric_style_is_preserved(self):
        book = Workbook()
        book.active.append(HEADERS)
        book.active.append(['Иван', '+79174005219', 46271, 'сайт'])
        book.active['C2'].number_format = 'dd.mm.yyyy'
        buffer = BytesIO()
        book.save(buffer)
        reference = load_workbook(BytesIO(buffer.getvalue()), read_only=True)
        expected_style = reference.active['C2']._style_id
        reference.close()
        self.assertNotEqual(expected_style, 0)
        rows, _ = read_rows(buffer.getvalue())
        self.assertEqual(rows[0].cells[2].metadata['style_id'], expected_style)

    def test_unambiguous_mdy_only_and_audit_action(self):
        for sep in ('/', '.', '-'):
            result = dates.analyze(RawCell(f'09{sep}25{sep}2026'))
            self.assertEqual(result.canonical, '2026-09-25')
            self.assertTrue(any(a['kind'] == 'DATE_ORDER_RESOLVED' for a in result.actions))
            self.assertEqual(dates.analyze(RawCell(f'09{sep}03{sep}2026')).canonical, '2026-03-09')
        for raw in ('02/30/2026', '13/25/2026', 'хх.09.2026', '09/25/26'):
            result = dates.analyze(RawCell(raw))
            self.assertIsNone(result.canonical)
            self.assertEqual(result.display, raw)
            self.assertTrue(result.codes)

    def test_diagnostic_codes_are_explicit_without_erasing_errors(self):
        result = phone.analyze(RawCell('+79174005219 доб. 123'))
        self.assertEqual(result.diagnostic_codes, ['PHONE_EXTRA_CONTEXT'])
        self.assertEqual(result.error_codes, [])
        result = phone.analyze(RawCell('не знаю'))
        self.assertEqual(result.diagnostic_codes, [])
        self.assertEqual(result.error_codes, ['PHONE_INVALID'])

    def test_clean_sheet_and_plain_review_preserve_every_problem(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, output = root / 'input.xlsx', root / 'result.xlsx'
            book = Workbook()
            book.active.append(HEADERS)
            raw_rows = [
                ['Дарья', '+79174005219', '09/25/2026', 'сайт'],
                ['Кирилл', 'не знаю', 'хх.09.2026', 'email'],
                [None, None, None, None],
                ['Пётр', '+79174005220 доб. 123', '2026-09-06', 'звонок'],
            ]
            for row in raw_rows:
                book.active.append(row)
            book.save(source)
            before = source.read_bytes()
            summary = normalize(source, output)
            self.assertEqual((summary['data_rows'], summary['problem_rows']), (1, 3))
            self.assertEqual(source.read_bytes(), before)
            self.assertEqual(restore_bytes(output)[0], before)
            result = load_workbook(output, read_only=True)
            try:
                self.assertEqual(result.sheetnames, ['Данные', 'Проблемы', 'Оригинал', 'Аудит', 'Проверить'])
                self.assertEqual(list(result['Данные'].values)[1:], [('Дарья', '+79174005219', '2026-09-25', 'сайт')])
                rows = list(result['Проверить'].values)
                self.assertEqual(rows[0], (*HEADERS, 'Лист входа', 'Строка входа', 'Что не так'))
                self.assertEqual([tuple(r[:4]) for r in rows[1:]], [tuple(r) for r in raw_rows[1:]])
                self.assertTrue(all(r[-1] and '{' not in r[-1] for r in rows[1:]))
                self.assertEqual(len(list(result['Оригинал'].values)) - 1, 4)
            finally:
                result.close()


if __name__ == '__main__':
    unittest.main()
