"""Independent catalog oracles and real four-sheet/CLI preservation checks."""
from datetime import date
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_STORED
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from fixtures.scenarios import CASES, materialize
from tools.generate_data import generate, write_workbook
from normalizer import phone, dates, names
from normalizer.models import HEADERS, SHEETS, RawCell, NormalizationError
from normalizer.pipeline import normalize, restore, source_result
from normalizer.preservation import restore_bytes

ROOT = Path(__file__).resolve().parents[1]


def case_cell(case):
    raw = case['input']
    types = {'number': 'n', 'serial': 'n', 'boolean': 'b', 'formula': 'f', 'error': 'e', 'datetime': 'd'}
    return RawCell(materialize(raw), types.get(raw.get('type'), 's') if isinstance(raw, dict) else 's',
                   number_format=raw.get('number_format', 'General') if isinstance(raw, dict) else 'General')


class FieldTests(unittest.TestCase):
    def test_all_catalog_oracles(self):
        analyzers = dict(zip(HEADERS, [names.analyze, phone.analyze, dates.analyze, source_result]))
        for case in CASES:
            with self.subTest(case=case['id']):
                result = analyzers[case['field']](case_cell(case))
                self.assertEqual(result.display, materialize(case['expected']))
                self.assertEqual(result.codes, case['codes'])
                self.assertEqual(result.status, case['status'])
                # Typed canonical in the generator is a JSON representation.
                self.assertEqual(result.canonical, materialize(case['canonical']))

    def test_all_month_dictionary_forms_and_existing_calendar(self):
        # Literal forms from the specification, independent of the parser dictionaries.
        ru = ['январь января янв', 'февраль февраля фев февр', 'март марта мар', 'апрель апреля апр', 'май мая',
              'июнь июня июн', 'июль июля июл', 'август августа авг', 'сентябрь сентября сен сент',
              'октябрь октября окт', 'ноябрь ноября ноя нояб', 'декабрь декабря дек']
        en = ['january jan', 'february feb', 'march mar', 'april apr', 'may', 'june jun', 'july jul', 'august aug',
              'september sep sept', 'october oct', 'november nov', 'december dec']
        for forms in (ru, en):
            for month, words in enumerate(forms, 1):
                for word in words.split():
                    for suffix in ('', '.'):
                        with self.subTest(word=word, suffix=suffix):
                            result = dates.analyze(RawCell(f'6 {word.upper()}{suffix} 2026'))
                            self.assertEqual(result.canonical, date(2026, month, 6).isoformat())
        for month, words in enumerate(en, 1):
            for word in words.split():
                self.assertEqual(dates.analyze(RawCell(f'{word} 6, 2026')).canonical, f'2026-{month:02d}-06')

    def test_valid_phone_with_damaged_or_second_extension_not_rescued(self):
        for value in ['+79174005219 доб. +79174005220', '+79174005219 доб. 123 456',
                      '+79174005219 / 123', '+79174005219 8 916 111 22 33', '++79174005219']:
            with self.subTest(value=value):
                result = phone.analyze(RawCell(value))
                self.assertIsNone(result.canonical)
                self.assertEqual(result.display, value)
                self.assertTrue(result.codes)

    def test_no_calendar_rollover_at_fractional_serial_boundary(self):
        for value, expected in [(59.99999999999, '1900-02-28'), (46271.99999999999, '2026-09-06'),
                                (0.75, '1904-01-01')]:
            result = dates.analyze(RawCell(value, 'n', number_format='yyyy-mm-dd hh:mm', epoch=1904 if value < 1 else 1900))
            self.assertEqual(result.canonical, expected)
            self.assertTrue(result.actions)


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def input(self, rows, filename='input.xlsx', **kwargs):
        path = self.root / filename
        write_workbook(path, rows, **kwargs)
        return path

    def read(self, path, sheet):
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            return list(wb[sheet].values)
        finally:
            wb.close()

    def test_user_examples_four_sheets_problem_context_and_restore(self):
        source = self.input([
            ['андрей волков', 'тел.: 8 917 400 52 19 (после 18:00)', '09/03/2026', 'сайт'],
            ['ПАВЛОВА анастасия', '8 (800) 707-46-19', '6 сентября 2026', 'telegram'],
            [None, '8 917 4OO 52 19', '31.04.2026', 'email']])
        original = source.read_bytes()
        output = self.root / 'result.xlsx'
        summary = normalize(source, output)
        wb = load_workbook(output, read_only=True)
        try:
            self.assertEqual(wb.sheetnames, SHEETS)
        finally:
            wb.close()
        values = self.read(output, 'Данные')
        self.assertEqual(values[1:], [('Анастасия', '+78007074619', '2026-09-06', 'telegram')])
        self.assertEqual(self.read(output, 'Проблемы')[1][11], '+79174005219')
        self.assertEqual(self.read(output, 'Оригинал')[3][3:7], (None, '8 917 4OO 52 19', '31.04.2026', 'email'))
        self.assertEqual(summary['problem_rows'], 2)
        self.assertEqual(source.read_bytes(), original)
        restored = self.root / 'restored.xlsx'
        restore(output, restored)
        self.assertEqual(restored.read_bytes(), original)

    def test_actual_suite_valid_inputs_and_structural_refusals(self):
        suite = self.root / 'suite'
        manifest = generate(suite)
        for entry in manifest['files']:
            with self.subTest(file=entry['file']):
                source = suite / entry['file']
                original = source.read_bytes()
                output = self.root / ('out-' + entry['file'])
                if 'fatal_code' in entry:
                    with self.assertRaises(NormalizationError) as caught:
                        normalize(source, output)
                    self.assertEqual(caught.exception.code, entry['fatal_code'])
                    self.assertFalse(output.exists())
                else:
                    summary = normalize(source, output)
                    self.assertEqual(restore_bytes(output)[0], original)
                    if entry['file'] == 'duplicates.xlsx':
                        self.assertEqual((summary['input_rows'], summary['data_rows'], summary['duplicates'], summary['problem_rows']), (12, 4, 5, 8))
                        problems = self.read(output, 'Проблемы')[1:]
                        self.assertEqual([r[2] for r in problems], entry['expected']['problem_rows'])
                        self.assertEqual([r[2] for r in problems if 'DUPLICATE_' in r[6]], entry['expected']['duplicate_conflict_rows'])
                    if entry['file'].startswith('epoch_'):
                        self.assertEqual(self.read(output, 'Данные')[1][2], '2026-09-06')
                self.assertEqual(source.read_bytes(), original)

    def test_one_phone_context_still_deduplicates_first_without_filling(self):
        source = self.input([[None, '+79174005219 доб. 123', 'bad', 'сайт'],
                             ['Иван', '8 (917) 400-52-19 вн. 456', '2026-09-06', 'email']])
        output = self.root / 'result.xlsx'
        normalize(source, output)
        self.assertEqual(len(self.read(output, 'Данные')), 1)
        self.assertEqual(self.read(output, 'Проблемы')[1][9:14:2], (None, '+79174005219', 'bad'))
        problems = self.read(output, 'Проблемы')[1:]
        self.assertEqual(len(problems), 2)
        self.assertEqual(problems[1][5], 'DUPLICATE')
        self.assertIsNone(problems[1][3])
        self.assertEqual(problems[1][4], 's1:r2')

    def test_nfc_names_do_not_create_false_conflict_latin_names_stay_raw(self):
        source = self.input([['е\u0308лкина анна', '+79174005219', '09/03/2026', 'сайт'],
                             ['Ёлкина Анна', '+79174005219', '2026-03-09', 'email'],
                             ['  McDONALD John  ', '+79174005220', '2026-09-06', 'сайт']])
        output = self.root / 'result.xlsx'
        summary = normalize(source, output)
        self.assertEqual(summary['problem_rows'], 1)
        self.assertNotIn('DUPLICATE_NAME_CONFLICT', summary['codes'])
        self.assertEqual(self.read(output, 'Данные')[1][0], 'Анна')
        self.assertEqual(self.read(output, 'Проблемы')[1][8], '  McDONALD John  ')

    def test_permuted_columns_blank_and_hidden_rows_unknown_sources(self):
        source = self.input([['custom', '2026-10-01', '+79174005219', 'кирилл'],
                             [None, None, None, None], ['new', '2026-09-06', '+79174005220', 'пётр']],
                            headers=['Источник', 'Дата заявки', 'Телефон', 'Имя'])
        wb = load_workbook(source)
        wb.active.row_dimensions[4].hidden = True
        wb.save(source)
        wb.close()
        output = self.root / 'result.xlsx'
        summary = normalize(source, output)
        self.assertEqual((summary['input_rows'], summary['data_rows'], summary['problem_rows']), (3, 2, 1))
        data = self.read(output, 'Данные')
        self.assertEqual(data[1], ('Кирилл', '+79174005219', '2026-10-01', 'custom'))
        self.assertEqual(self.read(output, 'Оригинал')[2][3:7], (None, None, None, None))
        self.assertEqual(data[2][0], 'Пётр')

    def test_hyperlink_formula_bool_serial_metadata_preserved(self):
        source = self.input([['иван', 'Позвонить', True, '=1+1'],
                             ['пётр', '+79174005219', {'type':'serial','value':60,'number_format':'yyyy-mm-dd'}, {'type':'formula','value':'=1+1'}]])
        wb = load_workbook(source)
        wb.active['B2'].hyperlink = 'tel:+79174005219'
        # A normal openpyxl read/save turns serial 60 into 59; create the actual raw 60 again.
        wb.active['C3'] = 60
        wb.active['C3'].number_format = 'yyyy-mm-dd'
        wb.save(source)
        wb.close()
        original = source.read_bytes()
        output = self.root / 'result.xlsx'
        normalize(source, output)
        data = self.read(output, 'Данные')
        self.assertEqual(len(data), 1)
        original_values = self.read(output, 'Оригинал')
        self.assertEqual(original_values[1][4], 'Позвонить')
        self.assertEqual(original_values[1][5], True)
        self.assertEqual(original_values[2][5], 60)
        wb = load_workbook(output)
        try:
            self.assertEqual(wb['Оригинал']['G2'].data_type, 's')
            self.assertEqual(wb['Оригинал']['G3'].data_type, 's')
        finally:
            wb.close()
        meta = json.loads(self.read(output, 'Оригинал')[1][7])
        self.assertEqual(meta['Телефон']['hyperlink']['target'], 'tel:+79174005219')
        self.assertEqual(restore_bytes(output)[0], original)

    def test_metadata_chunks_and_safe_long_literal_no_truncation(self):
        # An unknown name stays raw in 1.4; a long formula still requires VALUE_CHUNK
        # for its safe-literal FIELD_CHANGE, independently of name interpretation.
        source = self.input([['И' * 12000, '+79174005219', '2026-09-06T18:00',
                             {'type':'formula','value':'="' + 'И' * 12000 + '"'}]])
        output = self.root / 'result.xlsx'
        normalize(source, output)
        self.assertEqual(len(self.read(output, 'Оригинал')[1][3]), 12000)
        audit = self.read(output, 'Аудит')
        self.assertTrue(any(row[0] == 'CELL_METADATA_CHUNK' for row in audit))
        self.assertTrue(any(row[0] == 'VALUE_CHUNK' for row in audit))
        self.assertTrue(any(row[0] == 'DATE_TIME_PRESERVED' for row in audit))
        self.assertEqual(restore_bytes(output)[0], source.read_bytes())

    def test_bad_dimension_actual_rows_and_renamed_worksheet_part(self):
        source = self.input([['иван', '+79174005219', '2026-09-06', 'сайт']])
        modified = self.root / 'dimension.xlsx'
        with ZipFile(source) as old, ZipFile(modified, 'w') as new:
            for item in old.infolist():
                raw = old.read(item.filename)
                target = item.filename
                if item.filename == 'xl/worksheets/sheet1.xml':
                    raw = raw.replace(b'A1:D2', b'A1:A1')
                    target = 'xl/worksheets/custom.xml'
                if item.filename in ('xl/_rels/workbook.xml.rels', '[Content_Types].xml'):
                    raw = raw.replace(b'sheet1.xml', b'custom.xml')
                new.writestr(target, raw)
        summary = normalize(modified, self.root / 'result.xlsx')
        self.assertEqual(summary['input_rows'], 1)

    def test_failed_save_and_racing_destination_preserve_source_and_output(self):
        source = self.input([['иван', '+79174005219', '2026-09-06', 'сайт']])
        original = source.read_bytes()
        output = self.root / 'result.xlsx'
        def failed_save(self, path):
            Path(path).write_bytes(b'partial')
            raise OSError('injected')
        with patch.object(Workbook, 'save', failed_save):
            with self.assertRaises(OSError):
                normalize(source, output)
        self.assertFalse(output.exists())
        self.assertFalse(list(self.root.glob('.normalize-*')))
        import normalizer.pipeline as pipeline
        actual_publish = pipeline.publish
        def race(stage, destination):
            destination.write_bytes(b'other publisher')
            return actual_publish(stage, destination)
        with patch.object(pipeline, 'publish', race):
            with self.assertRaises(FileExistsError):
                normalize(source, output)
        self.assertEqual(output.read_bytes(), b'other publisher')
        self.assertEqual(source.read_bytes(), original)
        self.assertFalse(list(self.root.glob('.normalize-*')))

    def test_cli_codes_existing_output_limits_and_restore_corruption(self):
        source = self.input([[None, None, 'bad', None]])
        output = self.root / 'result.xlsx'
        def cli(script, *args):
            return subprocess.run([sys.executable, str(ROOT / script), *map(str, args)], cwd=ROOT,
                                  capture_output=True, text=True, encoding='utf-8', timeout=30)
        self.assertEqual(cli('normalize.py', '--input', source, '--output', output).returncode, 0)
        original_output = output.read_bytes()
        self.assertEqual(cli('normalize.py', '--input', source, '--output', output).returncode, 3)
        self.assertEqual(output.read_bytes(), original_output)
        self.assertEqual(cli('normalize.py', '--input', source, '--output', source).returncode, 3)
        limited = self.root / 'limited.xlsx'
        self.assertEqual(cli('normalize.py', '--input', source, '--output', limited, '--max-input-bytes', '10').returncode, 5)
        self.assertFalse(limited.exists())
        restored = self.root / 'restored.xlsx'
        self.assertEqual(cli('restore.py', '--input', output, '--output', restored).returncode, 0)
        self.assertEqual(restored.read_bytes(), source.read_bytes())
        wb = load_workbook(output)
        ws = wb['Аудит']
        details = json.loads(ws['J2'].value)
        details['sha256'] = '0' * 64
        ws['J2'] = json.dumps(details)
        altered = self.root / 'altered.xlsx'
        wb.save(altered)
        wb.close()
        rejected = self.root / 'rejected.xlsx'
        self.assertEqual(cli('restore.py', '--input', altered, '--output', rejected).returncode, 4)
        self.assertFalse(rejected.exists())

    def test_duplicate_zip_parts_crc_corruption_and_duplicate_xml_row_refused(self):
        import warnings
        import struct
        source = self.input([['иван', '+79174005219', '2026-09-06', 'сайт']])
        original = source.read_bytes()
        duplicate = self.root / 'duplicate.xlsx'
        shutil.copyfile(source, duplicate)
        with warnings.catch_warnings():
            warnings.simplefilter('ignore', UserWarning)
            with ZipFile(duplicate, 'a') as archive:
                archive.writestr('xl/worksheets/sheet1.xml', b'<different/>')
        corrupt = self.root / 'corrupt.xlsx'
        duplicated_row = self.root / 'duplicate_row.xlsx'
        for destination in (corrupt, duplicated_row):
            with ZipFile(source) as old, ZipFile(destination, 'w', compression=ZIP_STORED) as new:
                for info in old.infolist():
                    payload = old.read(info.filename)
                    if destination == duplicated_row and info.filename == 'xl/worksheets/sheet1.xml':
                        root = ET.fromstring(payload)
                        data = root.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheetData')
                        import copy
                        data.append(copy.deepcopy(data[-1]))
                        payload = ET.tostring(root)
                    new.writestr(info.filename, payload)
        with ZipFile(corrupt) as archive:
            info = archive.getinfo('xl/worksheets/sheet1.xml')
        payload = bytearray(corrupt.read_bytes())
        header = struct.unpack_from('<IHHHHHIIIHH', payload, info.header_offset)
        start = info.header_offset + 30 + header[-1] + header[-2]
        payload[start + info.file_size - 20] ^= 1
        corrupt.write_bytes(payload)
        for malformed in (duplicate, corrupt, duplicated_row):
            output = self.root / ('out-' + malformed.name)
            with self.subTest(file=malformed.name), self.assertRaises(NormalizationError) as error:
                normalize(malformed, output)
            self.assertEqual(error.exception.code, 'INPUT_INVALID')
            self.assertFalse(output.exists())
        self.assertEqual(source.read_bytes(), original)

    def test_rich_text_big_numeric_precision_and_other_sheet_restore(self):
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
        source = self.input([['иван', '+79174005219', '2026-09-06', 'сайт']])
        wb = load_workbook(source)
        wb.active['A2'] = CellRichText('иван ', TextBlock(InlineFont(b=True), 'иванов'))
        wb.active['D2'] = 9007199254740993
        wb.create_sheet('Прочее')['A1'] = 'другая информация'
        wb.save(source)
        wb.close()
        # Patch the actual numeric lexeme: the fixture writer itself rounds this integer.
        patched = self.root / 'exact-number.xlsx'
        with ZipFile(source) as old, ZipFile(patched, 'w') as new:
            for info in old.infolist():
                payload = old.read(info.filename)
                if info.filename == 'xl/worksheets/sheet1.xml':
                    root = ET.fromstring(payload)
                    node = root.find('.//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c[@r="D2"]')
                    node.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v').text = '9007199254740993'
                    payload = ET.tostring(root)
                new.writestr(info.filename, payload)
        output = self.root / 'result.xlsx'
        normalize(patched, output)
        self.assertEqual(self.read(output, 'Данные')[1][0], 'Иван')
        self.assertEqual(self.read(output, 'Данные')[1][3], '9007199254740993')
        original_meta = json.loads(self.read(output, 'Оригинал')[1][7])
        self.assertIn('Иванов'.lower(), original_meta['Имя']['decoded_value'])
        self.assertIn('rPr', original_meta['Имя']['raw_xml'])
        self.assertEqual(original_meta['Источник']['xml_value'], '9007199254740993')
        restored = self.root / 'restored.xlsx'
        restore(output, restored)
        self.assertEqual(restored.read_bytes(), patched.read_bytes())
        wb = load_workbook(restored, rich_text=True)
        try:
            self.assertIsInstance(wb.active['A2'].value, CellRichText)
            self.assertEqual(wb['Прочее']['A1'].value, 'другая информация')
        finally:
            wb.close()

    def test_output_overlong_xml_value_refused_and_explicit_sheet_selection(self):
        source = self.input([['иван', '+79174005219', '2026-09-06', 'сайт']], second_sheet=True)
        result = self.root / 'selected.xlsx'
        self.assertEqual(normalize(source, result, sheet='Заявки')['input_rows'], 1)
        overlong = self.root / 'overlong.xlsx'
        with ZipFile(source) as old, ZipFile(overlong, 'w') as new:
            for info in old.infolist():
                raw = old.read(info.filename)
                if info.filename == 'xl/worksheets/sheet1.xml':
                    root = ET.fromstring(raw)
                    node = root.find('.//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c[@r="D2"]')
                    node.find('.//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t').text = 'А' * 32768
                    raw = ET.tostring(root)
                new.writestr(info.filename, raw)
        output = self.root / 'refused.xlsx'
        with self.assertRaises(NormalizationError) as error:
            normalize(overlong, output, sheet='Заявки')
        self.assertEqual(error.exception.exit_code, 5)
        self.assertFalse(output.exists())
        self.assertFalse(list(self.root.glob('.normalize-*')))

    def test_two_actual_cli_publishers_only_one_output_and_exact_restore(self):
        source = self.input([['иван', f'+7917400{i:04d}', '2026-09-06', 'сайт'] for i in range(30)])
        output, gate = self.root / 'result.xlsx', self.root / 'gate'
        code = """import sys, time
from pathlib import Path
import normalizer.pipeline as pipeline
import normalize
original = pipeline.checked_paths
def gated(source, output):
    paths = original(source, output)
    Path(sys.argv[3]).write_text('ready')
    deadline = time.monotonic() + 15
    while not Path(sys.argv[4]).exists():
        if time.monotonic() > deadline: raise TimeoutError()
        time.sleep(0.02)
    return paths
pipeline.checked_paths = gated
source, output = sys.argv[1:3]
sys.argv = ['normalize.py', '--input', source, '--output', output]
raise SystemExit(normalize.main())
"""
        # Capture ready/gate names before main rewrites sys.argv.
        code = code.replace("original = pipeline.checked_paths", "original = pipeline.checked_paths\nready, gate = sys.argv[3:5]").replace('Path(sys.argv[3])', 'Path(ready)').replace('Path(sys.argv[4])', 'Path(gate)')
        processes = []
        try:
            for i in range(2):
                processes.append(subprocess.Popen([sys.executable, '-X', 'utf8', '-c', code, str(source), str(output),
                                                    str(self.root / f'ready-{i}'), str(gate)], cwd=ROOT,
                                                   stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                                   creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == 'win32' else 0))
            deadline = time.monotonic() + 15
            while not all((self.root / f'ready-{i}').exists() for i in range(2)):
                if time.monotonic() > deadline:
                    self.fail('Workers did not reach the gate')
                time.sleep(0.02)
            gate.write_text('go')
            details = [process.communicate(timeout=30) for process in processes]
            self.assertEqual(sorted(p.returncode for p in processes), [0, 3], details)
            self.assertEqual(restore_bytes(output)[0], source.read_bytes())
            self.assertFalse(list(self.root.glob('.normalize-*')))
        finally:
            for process in processes:
                if process.poll() is None:
                    process.kill()
                    process.communicate()


if __name__ == '__main__':
    unittest.main()
