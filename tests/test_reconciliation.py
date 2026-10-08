from copy import deepcopy
import gzip
from itertools import permutations
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from fixtures.reconciliation_cases import ROWS, EXPECTED
from normalizer.models import RawCell, RawRow, RowResult, SHEETS, NormalizationError
from normalizer import names, phone, dates, deduplicate, report
from normalizer.name_roles import classify
from normalizer.pipeline import normalize, source_result
from normalizer.invariants import validate_rows
from tools.generate_data import write_workbook
from verification.inventory import reconcile

ROOT = Path(__file__).resolve().parents[1]


def analyzed(values):
    rows = []
    for i, values in enumerate(values, 2):
        cells = [RawCell(v, 'n' if v is None else 's', f'{chr(65+j)}{i}') for j,v in enumerate(values)]
        rows.append(RowResult(RawRow(f's1:r{i}', 'Заявки', i, cells),
                   [f(c) for f,c in zip((names.analyze,phone.analyze,dates.analyze,source_result),cells)]))
    deduplicate.apply(rows)
    validate_rows(rows)
    return rows


class ReconciliationTests(unittest.TestCase):
    def test_full_resource_sets_enforce_single_token_intersections(self):
        reference = {k:set(v) for k,v in json.loads(gzip.decompress(
            (ROOT/'normalizer/data/name_lexicon.json.gz').read_bytes())).items()}
        first, last, middle = (reference[k] for k in ('FIRST_NAME','LAST_NAME','PATRONYMIC'))
        intersections = (first & last) | (first & middle) | (last & middle)
        for word in sorted(intersections):
            if not word.split() or len(word.split())!=1: continue
            self.assertEqual(classify(word)[0],'AMBIGUOUS',word)
            self.assertIsNone(names.analyze(RawCell(word)).canonical,word)
        for text in ('Иван','Андрей','Мария','Влад','Волков','Андреевич'):
            result = names.analyze(RawCell(text))
            self.assertEqual(result.codes,['NAME_AMBIGUOUS'])
            action = next(a for a in result.actions if a['kind']=='NAME_ROLE_EVIDENCE')
            self.assertGreater(len(action['details']['possible_roles']),1)

    def test_resource_clean_name_pool_is_valid_and_not_a_manual_short_list(self):
        from tools.generate_data import clean_name_pool
        pool = clean_name_pool()
        self.assertGreater(len(pool),100)
        for word in pool:
            result = names.analyze(RawCell(word))
            self.assertEqual((result.canonical,result.codes),(word,[]),word)
            self.assertEqual(names.analyze(RawCell(result.canonical)).canonical,word)

    def test_context_projection_cannot_assert_idempotence_without_fio(self):
        result = names.analyze(RawCell('Иван Иванов'))
        self.assertEqual((result.canonical,result.codes),('Иван',[]))
        self.assertEqual(names.analyze(RawCell(result.canonical)).codes,['NAME_AMBIGUOUS'])

    def test_context_confirms_first_name_without_inventing_other_roles(self):
        for full, personal in [('Николаев Павел','Павел'), ('Степанова Вера','Вера'),
                ('Макарова Нина','Нина'),('Григорьев Денис','Денис'), ('Романов Виктор','Виктор'),
                ('Смирнов Олег','Олег')]:
            for order in permutations(full.split()):
                result = names.analyze(RawCell(' '.join(order)))
                self.assertEqual((result.canonical,result.codes),(personal,[]))
                evidence = result.actions[0]['details']['role_assignments']
                self.assertEqual({a.index('FIRST_NAME') for a in evidence}, {order.index(personal)})
        vera = names.analyze(RawCell('Степанова Вера'))
        self.assertEqual(vera.components[0]['role'], 'LAST_NAME|PATRONYMIC')
        for text in ('Ким', 'Ким Павел', 'Иван Кирилл', 'Клиент Кирилл', 'Петрович', 'Вера'):
            self.assertTrue(names.analyze(RawCell(text)).codes, text)

    def test_invisible_removal_only_at_approved_boundaries_and_audited(self):
        for text in ('Смирнов Олег\u200b', '\ufeffСмирнов\u200b \ufeffОлег\ufeff', 'Смирнов\u00a0Олег'):
            result = names.analyze(RawCell(text))
            self.assertEqual((result.canonical,result.codes), ('Олег',[]))
            if '\u200b' in text or '\ufeff' in text:
                action = next(a for a in result.actions if a['kind']=='NAME_INVISIBLES_REMOVED')
                self.assertEqual(action['details']['original'],text)
                self.assertTrue(action['details']['removed'])
        partial = names.analyze(RawCell('Иван Иванов Петрович\u200b'))
        self.assertEqual((partial.canonical,partial.codes), ('Иван',['NAME_AMBIGUOUS']))
        self.assertTrue(any(a['kind']=='NAME_INVISIBLES_REMOVED' for a in partial.actions))
        for text in ('Ол\u200bег', 'Смирнов\u200bОлег', 'Ол\ufeffег', '\u202eОлег', 'Олег\u200d', '\u200b\ufeff'):
            self.assertIsNone(names.analyze(RawCell(text)).canonical, text)

    def test_fio_permutations_yo_and_partial_information_are_compatible(self):
        for pair in [('Королёва Дарья','Королева Дарья'), ('Борисова Алёна','Алёна Борисова'),
                     ('Тарасова Мария','Мария Тарасова'), ('Иван','Иванов Иван'),
                     ('Иван Иванов','Иванович Иванов Иван')]:
            rows = analyzed([[n,'+79174005219','2026-09-06','сайт'] for n in pair])
            self.assertTrue(rows[0].admitted)
            self.assertFalse(rows[0].conflicts)
            self.assertEqual(sum(r.admitted for r in rows), 1)
            self.assertEqual(rows[0].fields[0].raw.value, pair[0])

    def test_initial_is_not_guessed_but_group_provides_confirmed_name(self):
        rows = analyzed([['К. Белоусов','+79117703148','27.09.2026','звонок'],
                         ['Белоусов Кирилл','89117703148','09/27/2026','сайт']])
        self.assertEqual(rows[0].fields[0].canonical,None)
        self.assertEqual(rows[0].fields[0].codes,['NAME_INCOMPLETE'])
        self.assertEqual(rows[0].data_fields[0].canonical,'Кирилл')
        self.assertTrue(rows[0].admitted)
        self.assertEqual(rows[0].resolution['fields'][0]['donor_row_id'],'s1:r3')
        self.assertEqual(rows[0].representative_id,'s1:r2')
        lone = analyzed([['К. Белоусов','+79117703148','27.09.2026','звонок']])
        self.assertFalse(lone[0].admitted)

    def test_complementary_fields_and_phone_context_use_explicit_donors(self):
        rows = analyzed([[None,'+79174005219 доб. 42','2026-09-06','первый'],
                         ['Кирилл','89174005219','bad','второй']])
        self.assertTrue(rows[0].admitted)
        self.assertEqual([f.display for f in rows[0].data_fields], ['Кирилл','+79174005219','2026-09-06','первый'])
        self.assertEqual([f['donor_row_id'] for f in rows[0].resolution['fields']], ['s1:r3','s1:r3','s1:r2','s1:r2'])
        self.assertIn('NAME_MISSING',rows[0].codes)
        self.assertIn('DATE_UNSUPPORTED',rows[1].codes)
        self.assertFalse(any(r.conflicts for r in rows))

    def test_conflicts_keep_first_confirmed_values_in_data_and_all_problems(self):
        rows = analyzed([[None,'+79174005219','bad','first'],
                         ['Иван Иванов','89174005219','2026-09-06','second'],
                         ['Петров Кирилл','+79174005219','2026-09-07','third']])
        self.assertTrue(rows[0].admitted)
        self.assertEqual([f.display for f in rows[0].data_fields], ['Иван','+79174005219','2026-09-06','first'])
        for row in rows:
            self.assertIn('DUPLICATE_NAME_CONFLICT',row.codes)
            self.assertIn('DUPLICATE_DATE_CONFLICT',row.codes)
        # Pairwise compatibility is not transitive: a bare Ivan cannot hide surnames.
        rows = analyzed([[n,'+79174005219','2026-09-06','сайт'] for n in ['Иван','Иван Иванов','Иван Петров']])
        self.assertTrue(rows[0].admitted)
        self.assertIn('DUPLICATE_NAME_CONFLICT',rows[0].codes)

    def test_missing_phones_and_other_groups_never_supply_fields(self):
        rows = analyzed([[None,'+79174005219','2026-09-06','сайт'],
                         ['Кирилл','+79174005220','2026-09-06','сайт'],
                         ['Иван',None,'2026-09-06','сайт'], ['Иван',None,'2026-09-06','сайт']])
        self.assertEqual([r.admitted for r in rows],[False,True,False,False])
        self.assertEqual(rows[2].representative_id,rows[2].raw.row_id)
        self.assertEqual(rows[3].representative_id,rows[3].raw.row_id)

    def test_all_user_rows_real_cli_five_sheets_lineage_and_exact_restore(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder); source, output = folder/'input.xlsx',folder/'output.xlsx'
            write_workbook(source, ROWS)
            original = source.read_bytes()
            run = subprocess.run([sys.executable,'-X','utf8',str(ROOT/'normalize.py'),
                                  '--input',str(source),'--output',str(output)], capture_output=True,
                                  text=True,encoding='utf-8',cwd=ROOT,timeout=30)
            self.assertEqual(run.returncode,0,run.stderr)
            self.assertEqual(source.read_bytes(),original)
            evidence = reconcile(original, output.read_bytes())
            self.assertEqual((evidence['rows'],evidence['data_rows'],evidence['duplicates']), (16,10,6))
            book = load_workbook(output,read_only=True)
            try:
                self.assertEqual(book.sheetnames,SHEETS)
                self.assertEqual(list(book['Данные'].values)[1:],EXPECTED)
                problems = list(book['Проблемы'].values)[1:]
                self.assertEqual(len(problems),1)
                self.assertEqual(problems[0][6],'NAME_INCOMPLETE')
                self.assertIn('Инициалы',list(book['Проверить'].values)[1][-1])
                groups = {a[1]:json.loads(a[9]) for a in list(book['Аудит'].values)[1:] if a[0]=='GROUP_RESOLUTION'}
                self.assertEqual(groups['s1:r14']['fields'][0]['donor_row_id'],'s1:r15')
            finally:
                book.close()

    def test_wrong_or_repeated_group_audit_is_never_published(self):
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder); source=folder/'input.xlsx'
            write_workbook(source,[['К. Белоусов','+79117703148','27.09.2026','сайт'],
                                   ['Белоусов Кирилл','89117703148','09/27/2026','email']])
            original = report.Writer.audit
            for label in ('wrong-donor','wrong-value','duplicate','missing'):
                def corrupt(writer, kind, *args, **kwargs):
                    if kind=='GROUP_RESOLUTION':
                        if label=='missing': return
                        kwargs=deepcopy(kwargs)
                        if label=='wrong-donor': kwargs['details']['fields'][0]['donor_row_id']='s1:r2'
                        if label=='wrong-value': kwargs['details']['fields'][0]['value']='Иван'
                        if label=='duplicate': original(writer,kind,*args,**kwargs)
                    return original(writer,kind,*args,**kwargs)
                output=folder/(label+'.xlsx')
                with patch.object(report.Writer,'audit',corrupt):
                    with self.assertRaises(NormalizationError) as caught: normalize(source,output)
                self.assertEqual(caught.exception.code,'INVARIANT_FAILED')
                self.assertFalse(output.exists())


if __name__=='__main__':
    unittest.main()
