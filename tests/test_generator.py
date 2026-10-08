"""Checks for actual XLSX types, reproducible labels and preservation on refusal."""
from datetime import datetime, timezone
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from fixtures.scenarios import CASES, HEADERS
from tools.generate_data import generate, write_workbook


class GeneratorTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_clean_has_unique_text_phones_and_calendar_dates(self):
        result = generate(self.root / "clean", "clean", 48, 7)
        wb = load_workbook(self.root / "clean" / "clean.xlsx")
        try:
            sheet = wb.active
            self.assertEqual(list(next(sheet.values)), HEADERS)
            values = list(sheet.values)[1:]
            self.assertEqual(len(values), 48)
            self.assertEqual(len({row[1] for row in values}), 48)
            for row in values:
                self.assertRegex(row[1], r"^\+7[0-9]{10}$")
                datetime.strptime(row[2], "%Y-%m-%d")
            self.assertEqual(result["files"][0]["rows"], 48)
        finally:
            wb.close()

    def test_dirty_catalog_types_and_every_case(self):
        result = generate(self.root / "dirty", "dirty", len(CASES))
        wb = load_workbook(self.root / "dirty" / "dirty.xlsx", data_only=False)
        try:
            self.assertEqual({x["case_id"] for x in result["files"][0]["labels"]}, {x["id"] for x in CASES})
            for index, case in enumerate(CASES, 2):
                cell = wb.active.cell(index, HEADERS.index(case["field"]) + 1)
                value = case["input"]
                if isinstance(value, dict):
                    if value["type"] == "formula":
                        self.assertEqual(cell.data_type, "f")
                        self.assertEqual(cell.value, value["value"])
                    if value["type"] == "error":
                        self.assertEqual(cell.data_type, "e")
                    if value["type"] in ("datetime", "serial") and value.get("number_format", "").startswith(("dd", "yyyy")):
                        self.assertIsInstance(cell.value, datetime)
                elif isinstance(value, str) and value.startswith("="):
                    self.assertEqual(cell.data_type, "s")
        finally:
            wb.close()

    def test_seed_reproduces_values_and_labels(self):
        a = generate(self.root / "a", "mixed", len(CASES), 19)
        b = generate(self.root / "b", "mixed", len(CASES), 19)
        self.assertEqual(a["files"][0]["labels"], b["files"][0]["labels"])
        wa, wb = load_workbook(self.root / "a" / "mixed.xlsx"), load_workbook(self.root / "b" / "mixed.xlsx")
        try:
            self.assertEqual(list(wa.active.values), list(wb.active.values))
        finally:
            wa.close()
            wb.close()

    def test_existing_destination_never_changes(self):
        destination = self.root / "exists"
        destination.mkdir()
        marker = destination / "important.txt"
        marker.write_bytes(b"original")
        with self.assertRaises(FileExistsError):
            generate(destination)
        self.assertEqual(marker.read_bytes(), b"original")
        self.assertEqual(list(destination.iterdir()), [marker])

    def test_suite_epochs_structural_errors_and_duplicate_expectations(self):
        manifest = generate(self.root / "suite", "suite", 48)
        self.assertEqual(len(manifest["files"]), 14)
        for epoch in (1900, 1904):
            wb = load_workbook(self.root / "suite" / f"epoch_{epoch}.xlsx")
            try:
                self.assertEqual(wb.active["C2"].value, datetime(2026, 9, 6))
            finally:
                wb.close()
        dup = next(item for item in manifest["files"] if item["file"] == "duplicates.xlsx")
        self.assertEqual(set(dup["expected"]["representative_or_ungrouped_rows"]) | set(dup["expected"]["duplicate_rows"]), set(range(2, 14)))
        self.assertEqual(dup["expected"]["data_input_rows"], [2, 8, 11, 12])
        self.assertEqual(dup["expected"]["duplicate_conflict_rows"], [2, 3, 4, 5])
        disk = json.loads((self.root / "suite" / "manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest, disk)

    def test_1904_adjusts_date_fixture_and_expectation(self):
        manifest = generate(self.root / "mac", "dirty", len(CASES), epoch=1904)
        labels = manifest["files"][0]["labels"]
        d19 = next(x for x in labels if x["case_id"] == "D19")
        d38 = next(x for x in labels if x["case_id"] == "D38")
        self.assertEqual(d19["input"][2]["value"], 44809)
        self.assertEqual(d38["field_expectation"]["codes"], [])
        self.assertEqual(d38["field_expectation"]["expected"], "1904-03-01")

    def test_long_cell_refused_without_truncation_or_publication(self):
        destination = self.root / "too-long.xlsx"
        with self.assertRaisesRegex(ValueError, "no truncation"):
            write_workbook(destination, [['А' * 32768, '+79151451780', '2026-09-03', 'сайт']])
        self.assertFalse(destination.exists())

    def test_writer_preserves_existing_file(self):
        destination = self.root / "existing.xlsx"
        destination.write_bytes(b"original user data")
        with self.assertRaises(FileExistsError):
            write_workbook(destination, [['Иван', '+79151451780', '2026-09-03', 'сайт']])
        self.assertEqual(destination.read_bytes(), b"original user data")

    def test_nonfinite_numeric_values_refused_without_silent_blank(self):
        for index, value in enumerate((float('nan'), float('inf'), float('-inf'))):
            with self.subTest(value=value):
                destination = self.root / f"nonfinite-{index}.xlsx"
                with self.assertRaisesRegex(ValueError, "no silent blank"):
                    write_workbook(destination, [['Иван', value, '2026-09-03', 'сайт']])
                self.assertFalse(destination.exists())

    def test_native_timezone_date_refused_before_partial_save(self):
        destination = self.root / 'aware.xlsx'
        with self.assertRaisesRegex(ValueError, 'timezone-aware'):
            write_workbook(destination, [['Иван', '+79174005219', datetime(2026, 9, 6, tzinfo=timezone.utc), 'сайт']])
        self.assertFalse(destination.exists())


if __name__ == "__main__":
    unittest.main()
