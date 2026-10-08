"""Regression contracts for independent verification and DEC-29."""
from itertools import permutations
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from openpyxl import Workbook, load_workbook
from normalizer import names, phone, dates, deduplicate
from normalizer.models import RawCell
from normalizer.pipeline import normalize
from verification.reference import calendar, phone as phone_oracle, role_assignments, ALPHABET, evaluate
from verification.adapter import actual
from verification.exhaustive import check
from verification.inventory import reconcile


class FormalContracts(unittest.TestCase):
    def test_stateful_seed_controls_executed_trace(self):
        from verification.stateful import check as stateful_check
        first, same, other = [stateful_check(s, 10, 10) for s in (20261008,20261008,20261009)]
        self.assertEqual(first['trace_sha256'], same['trace_sha256'])
        self.assertEqual(first['calls'], same['calls'])
        self.assertNotEqual(first['trace_sha256'], other['trace_sha256'])

    def test_corrupted_stage_is_never_published(self):
        from verification.practical import publication_faults
        with tempfile.TemporaryDirectory() as folder:
            result = publication_faults(Path(folder)/'faults')
            self.assertTrue(result['passed'], result)
            self.assertEqual(result['detected'], 10)

    def test_bounded_exhaustive_pairs(self):
        result = check(2)
        self.assertTrue(result['passed'], result)
        self.assertEqual(result['total_sequences'], 2353)

    def test_independent_field_oracles(self):
        for text in ('09/03/2026', '09/25/2026', '12/31/2026', '31.04.2026', '29.02.2000', '29.02.1900', '2026-09-06', '09 03 2026'):
            with self.subTest(text=text):
                self.assertEqual(dates.analyze(RawCell(text)).canonical, calendar(text).value)
        for text in ('9174005219', '+79174005219', '8 (917) 400-52-19', '89174005219123', '81012345678', '0000000000'):
            self.assertEqual(phone.analyze(RawCell(text)).canonical, phone_oracle(text).value)
        self.assertEqual(len(role_assignments([{'FIRST_NAME'}, {'LAST_NAME', 'PATRONYMIC'}])), 2)

    def test_petrovich_partial_name_requires_review_and_full_conflict(self):
        for parts in permutations(('Иван', 'Иванов', 'Петрович')):
            result = names.analyze(RawCell(' '.join(parts)))
            self.assertEqual(result.canonical, 'Иван')
            self.assertIn('NAME_AMBIGUOUS', result.codes)
            self.assertEqual(result.classification, 'AMBIGUOUS')
            self.assertTrue(result.actions[0]['details']['review_required'])
        for value in ('Петрович', 'Ким Иванов Иван', 'Клиент Иван Петрович', 'Иван Кирилл Петрович'):
            self.assertIsNone(names.analyze(RawCell(value)).canonical)

    def test_independent_inventory_and_plain_review(self):
        with tempfile.TemporaryDirectory() as folder:
            input_path, output_path = Path(folder)/'in.xlsx', Path(folder)/'out.xlsx'
            book = Workbook(); ws = book.active; ws.title = 'Заявки'
            ws.append(['Имя', 'Телефон', 'Дата заявки', 'Источник'])
            ws.append(['Иван Иванов Петрович', '8 (917) 400-52-19', '09/25/2026', 'сайт'])
            ws.append(['Иван Иванов', '+79174005219', '25.09.2026', 'email'])
            ws.append(['Кирилл', None, 'хх.09.2026', '=1+1'])
            ws.append(['Дарья', '+79174005220', '2026-09-06', 'сайт'])
            book.save(input_path); book.close()
            normalize(input_path, output_path)
            evidence = reconcile(input_path.read_bytes(), output_path.read_bytes())
            self.assertEqual(evidence['rows'], 4)
            self.assertEqual(evidence['data_rows'], 2)
            result = load_workbook(output_path, read_only=True)
            try:
                problems = list(result['Проблемы'].values)
                self.assertEqual(problems[1][9], 'Иван')
                self.assertEqual(list(result['Проверить'].values)[1][0], 'Иван Иванов Петрович')
                self.assertIn('NAME_AMBIGUOUS', problems[1][6])
                self.assertNotIn('DUPLICATE_NAME_CONFLICT', problems[1][6])
            finally:
                result.close()
