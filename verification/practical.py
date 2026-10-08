"""Practical independent checks, retained artifacts, and targeted mutation experiments."""
from collections import Counter
from contextlib import ExitStack
from datetime import date
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
import subprocess
import sys
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED
from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from normalizer import dates, phone, models, deduplicate, report
from normalizer.models import RawCell
from .reference import ALPHABET, calendar_valid, calendar, phone as expected_phone, evaluate
from .adapter import actual, make_rows, observe
from .inventory import HEADERS, reconcile

ROOT = Path(__file__).resolve().parents[1]


def save_book(path, rows, epoch=1900, decorate=False):
    book = Workbook(); ws = book.active; ws.title = 'Заявки'
    if epoch == 1904: book.epoch = CALENDAR_MAC_1904
    ws.append(list(HEADERS))
    for row in rows: ws.append(list(row))
    if decorate:
        ws['B2'].hyperlink = 'tel:+79999999999'; ws['B2'].hyperlink.tooltip = 'Ссылка не источник телефона'
        ws.row_dimensions[3].hidden = True
        ws['C3'] = 46271; ws['C3'].number_format = 'yyyy-mm-dd'
        ws['C4'] = 46271; ws['C4'].number_format = 'mm'
        ws['B5'] = True; ws['D5'] = '=1+1'
        ws['B6'] = 9174005219.0; ws['B6'].number_format = '0.00E+00'
        ws['D6'] = 12345678901234567890
        ws['D7'] = '@raw'; ws['D8'] = '+raw'; ws['D9'] = '-raw'
    book.save(path); book.close()


def alter_zip(path, transform):
    stream = BytesIO()
    with ZipFile(path) as source, ZipFile(stream, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist(): target.writestr(item.filename, transform(item.filename, source.read(item.filename)))
    path.write_bytes(stream.getvalue())


def cli_preservation(folder):
    folder.mkdir()
    abstract = [(r.name, r.telephone, r.date, r.source) for r in ALPHABET]
    user = [('андрей волков','8 (915) 145-17-80','09/03/2026','сайт'),
            ('Волков Егор','+7 916 347 17 03','03.09.2026','email'),
            ('ПАВЛОВА анастасия','89032118241','05.09.2026','звонок'),
            ('Ирина Львовна','+7 (913) 211-88-41','09/05/2026','сайт'),
            ('Павлов  Андрей','7 925 601 32 12','6 сентября 2026','звонок')]
    edge = [('Кирилл',f'+7917400{i:04d}', '2026-09-06', 'сайт') for i in range(12)]
    edge[0] = ('Иван Иванов Петрович', '+79174005219', '09/25/2026', 'сайт')
    edge[1] = ('Иван Иванов', '+79174005219', '25.09.2026', 'email')
    edge[7] = (None, None, None, None)
    edge[8] = ('Иван', 'не знаю', 'хх.09.2026', '=literal')
    edge[9] = ('Жан-Поль', '8 917 400 52 19 (после 18:00)', '2026-09-06T18:00', 'звонок')
    fixtures = [('abstract-48', abstract, 1900, False), ('user-5', user, 1900, False),
                ('metadata-edge', edge, 1900, True), ('epoch-1904', edge, 1904, True),
                ('bad-dimension-48', abstract, 1900, False), ('nonfinite', edge, 1900, True)]
    chunk_rows = [('Дарья','+79174005219','2026-09-06','я'*16000),
                  ('Кирилл',None,'broken_date','я'*16000)]
    fixtures.append(('metadata-chunks',chunk_rows,1900,False))
    results = []
    for label, rows, epoch, decorate in fixtures:
        original, output = folder/(label+'-input.xlsx'), folder/(label+'-output.xlsx')
        save_book(original, rows, epoch, decorate)
        if label == 'bad-dimension-48':
            alter_zip(original, lambda p,b: re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', b) if p == 'xl/worksheets/sheet1.xml' else b)
        if label == 'nonfinite':
            def nonfinite(p,b):
                return re.sub(rb'(<c r="B6"[^>]*><v>)[^<]+', rb'\g<1>NaN', b) if p == 'xl/worksheets/sheet1.xml' else b
            alter_zip(original, nonfinite)
        before = original.read_bytes()
        command = [sys.executable, '-X', 'utf8', str(ROOT/'normalize.py'), '--input', str(original), '--output', str(output)]
        run = subprocess.run(command, capture_output=True, text=True, encoding='utf-8', timeout=60, cwd=ROOT)
        (folder/(label+'.log')).write_text(run.stdout+run.stderr, 'utf-8')
        assert run.returncode == 0 and output.exists(), (label, run.returncode, run.stderr)
        assert original.read_bytes() == before, 'input unchanged'
        result = reconcile(before, output.read_bytes())
        if label == 'user-5':
            book = load_workbook(output, read_only=True)
            try:
                expected = [('Андрей','+79151451780','2026-03-09','сайт'), ('Егор','+79163471703','2026-09-03','email'),
                            ('Анастасия','+79032118241','2026-09-05','звонок'), ('Ирина','+79132118841','2026-05-09','сайт'),
                            ('Андрей','+79256013212','2026-09-06','звонок')]
                assert list(book['Данные'].values)[1:] == expected, 'user 5 values'
            finally: book.close()
        result.update(id=label, cli_exit_code=run.returncode, input=str(original), output=str(output))
        results.append(result)
    return {'passed': True, 'books': results, 'rows':sum(r['rows'] for r in results),
            'scope':'ZIP/XML fixtures, styles/epoch/formula/links/blank/hidden/NaN/dimensions; not all OOXML dialects'}


def fields():
    calendar_cases = 0; valid = 0
    for year in range(1,401):
        for month in range(1,13):
            for day in range(35):
                expected = calendar_valid(year, month, day)
                result = dates.analyze(RawCell(f'{year:04d}-{month:02d}-{day:02d}'))
                assert (result.canonical is not None) == expected, (year,month,day,result)
                if expected: assert result.canonical == f'{year:04d}-{month:02d}-{day:02d}'
                calendar_cases += 1; valid += expected
    assert valid == 146097  # Gregorian 400-year cycle: 400*365+97 leap days.
    extra = ('9999-12-31','0000-01-01','29/02/1900','29/02/2000','09/25/2026','12/31/2026','09 03 2026')
    for text in extra: assert dates.analyze(RawCell(text)).canonical == calendar(text).value
    phone_cases = 0
    for digit in '0123456789':
        for position in range(10):
            national = '9174005219'[:position]+digit+'9174005219'[position+1:]
            for text in (national, '+7'+national, '7'+national, '8'+national, '8 ('+national[:3]+') '+national[3:], national+'123'):
                assert phone.analyze(RawCell(text)).canonical == expected_phone(text).value, text
                phone_cases += 1
    return {'passed':True, 'gregorian_cycle_calendar_cases':calendar_cases, 'valid_days':valid,
            'extra_date_cases':len(extra), 'phone_cases':phone_cases}


def critical_conditions():
    records = []
    base = [True]*5
    # Unique-cause MC/DC for the five atoms of RowResult.admitted in isolation.
    # These synthetic states are not all reachable through field analyzers.
    for false_at in [None, 0, 1, 2, 3, 4]:
        atoms = base.copy()
        if false_at is not None: atoms[false_at] = False
        r = make_rows([ALPHABET[0]])[0]
        r.disposition = 'REPRESENTATIVE' if atoms[0] else 'DUPLICATE'
        r.fields[0].codes = [] if atoms[1] else ['NAME_UNKNOWN']
        for index in range(3):
            if not atoms[index+2]: r.fields[index].canonical = None
        expected = all(atoms)
        assert r.admitted == expected
        records.append({'atoms':atoms, 'decision':r.admitted})
    return {'passed':True, 'decision':'RowResult.admitted', 'unique_cause_pairs':5, 'witnesses':records,
            'limitation':'Five unit-level atoms only; synthetic states; no whole-program MC/DC claim'}


def mutations(folder, valid_input, valid_output):
    folder.mkdir(); records = []
    # Positive controls must pass before fault injections can count as killed.
    raw = valid_input.read_bytes(); clean = valid_output.read_bytes(); reconcile(raw, clean)
    sequence = [ALPHABET[0], ALPHABET[1]]
    assert actual(sequence) == evaluate(sequence)
    original_shape = phone.shape; original_apply = deduplicate.apply
    def partial(text):
        compact = re.sub('['+phone.SEPARATORS+']', '', text)
        return original_shape(compact[:12] if compact.startswith('+') else compact[:11])
    checks = [('swap-date-order', lambda: dates.analyze(RawCell('09/03/2026')).canonical == calendar('09/03/2026').value,
               patch.object(dates, 'output', lambda cell,parsed,*a,**kw: models.success(cell,f'{parsed.year:04d}-{parsed.day:02d}-{parsed.month:02d}'))),
              ('last-duplicate', lambda: actual(sequence) == evaluate(sequence),
               patch.object(deduplicate,'apply',lambda rows: original_apply(list(reversed(rows))))),
              ('partial-phone', lambda: phone.analyze(RawCell('+79174005219123')).canonical == expected_phone('+79174005219123').value,
               patch.object(phone,'shape',partial)),
              ('disable-calendar', lambda: dates.analyze(RawCell('2026-04-31')).canonical == calendar('2026-04-31').value,
               patch.object(dates,'analyze',lambda cell: models.success(cell,'2026-04-31')))]
    for label, check, context in checks:
        assert check(), ('positive control',label)
        with context:
            caught = not check()
        assert caught, ('surviving mutant',label)
        records.append({'id':label,'detected':caught,'method':'independent semantic assertion, memory-only mutation'})
    faults = ('drop-ROW', 'duplicate-row-id', 'overwrite-original', 'corrupt-metadata', 'corrupt-change-old', 'corrupt-chunk', 'corrupt-data')
    for label in faults:
        book = load_workbook(BytesIO(clean)); audit = book['Аудит']
        rows = [r for r in audit if r[0].value == 'ROW']
        if label == 'drop-ROW': audit.delete_rows(rows[0][0].row)
        elif label == 'duplicate-row-id': rows[1][1].value = rows[0][1].value
        elif label == 'overwrite-original': book['Оригинал']['D2'] = 'overwritten'
        elif label == 'corrupt-metadata':
            metadata = json.loads(book['Оригинал']['H2'].value)
            metadata['Имя']['style_id'] = 987
            book['Оригинал']['H2'] = json.dumps(metadata)
        elif label == 'corrupt-change-old':
            row = next(r for r in audit if r[0].value == 'FIELD_CHANGE')
            row[7].value = '{"type":"s","value":"wrong"}'
        elif label == 'corrupt-chunk':
            row = next(r for r in audit if r[0].value == 'INPUT_CHUNK')
            row[11].value = 'AAAA'+row[11].value[4:]
        elif label == 'corrupt-data': book['Данные']['A2'] = 'wrong'
        stream = BytesIO(); book.save(stream); book.close(); payload = stream.getvalue()
        path = folder/(label+'.xlsx'); path.write_bytes(payload)
        try: reconcile(raw,payload)
        except AssertionError as exc: caught, reason = True, str(exc)
        else: caught, reason = False, 'survived'
        assert caught, ('surviving mutant',label)
        records.append({'id':label,'detected':caught,'method':'independent output reconciliation','reason':reason,
                        'artifact':str(path), 'sha256':sha256(payload).hexdigest()})
    return {'passed':True,'detected':len(records),'total':len(records),'mutations':records}


def publication_faults(folder):
    """Faults occur while writing a stage: a broken stage must never be published."""
    from normalizer.pipeline import normalize
    folder.mkdir(); results=[]
    original=folder/'input.xlsx'
    save_book(original, [('Иван','8 (917) 400-52-19','09/03/2026','сайт'),
                         ('Иван','+79174005219','09.03.2026','email')])
    initial=original.read_bytes()
    put, audit = report.Writer.put, report.Writer.audit
    for label in ('metadata', 'change-old', 'ROW-details', 'duplicate-link', 'problem-raw', 'raw-bool-type',
                  'group-donor', 'group-value', 'group-repeat', 'group-missing'):
        # Every fault starts from a fresh eligible phone group. Previous faults
        # intentionally use a Boolean/malformed phone and cannot be reused here.
        save_book(original, [('Кирилл','8 (917) 400-52-19','09/03/2026','сайт'),
                             ('Кирилл','+79174005219','09.03.2026','email')])
        initial=original.read_bytes()
        output=folder/(label+'.xlsx')
        injections=[]
        def broken_put(writer,sheet,values):
            values=list(values)
            if label=='metadata' and sheet=='Оригинал' and values[0]!='row_id':
                injections.append(label)
                meta=json.loads(values[7]); meta['Имя']['style_id']=987; values[7]=json.dumps(meta)
            if label=='problem-raw' and sheet=='Проблемы' and values[0]!='row_id':
                injections.append(label); values[8]='erased'
            if label=='raw-bool-type' and sheet=='Оригинал' and values[0]!='row_id':
                injections.append(label); values[4]=1
            return put(writer,sheet,values)
        def broken_audit(writer,kind,*args,**kwargs):
            if label.startswith('group-') and kind=='GROUP_RESOLUTION':
                injections.append(label)
                if label=='group-missing': return
                # Work on a copy: never corrupt the validator's expected object.
                kwargs=json.loads(json.dumps(kwargs))
                if label=='group-donor': kwargs['details']['fields'][0]['donor_row_id']='s1:r3'
                if label=='group-value': kwargs['details']['fields'][0]['value']='wrong'
                if label=='group-repeat': audit(writer,kind,*args,**kwargs)
            if label=='change-old' and kind=='FIELD_CHANGE':
                injections.append(label); kwargs['old']={'type':'s','value':'erased'}
            if label=='ROW-details' and kind=='ROW':
                injections.append(label); kwargs['details']={'disposition':'DUPLICATE','canonical_phone':None,'codes':['WRONG']}
            result=audit(writer,kind,*args,**kwargs)
            if label=='duplicate-link' and kind=='DUPLICATE_LINKED':
                injections.append(label); audit(writer,kind,*args,**kwargs)
            return result
        if label in ('problem-raw', 'raw-bool-type'):
            # Context ensures Problems is nonempty while keeping a valid phone key.
            save_book(original, [('Иван',True if label=='raw-bool-type' else '8 (917) 400-52-19 (после 18:00)','09/03/2026','сайт')])
            initial=original.read_bytes()
        with patch.object(report.Writer,'put',broken_put), patch.object(report.Writer,'audit',broken_audit):
            try: normalize(original,output)
            except models.NormalizationError as exc:
                detected=exc.exit_code==4 and not output.exists(); reason=exc.code+': '+str(exc)
            else: detected=False; reason='published corrupted stage'
        assert original.read_bytes()==initial
        results.append({'id':label,'detected':detected and bool(injections),'injections':len(injections),
                        'output_published':output.exists(),'reason':reason})
    return {'passed':all(r['detected'] for r in results),'faults':results,'detected':sum(r['detected'] for r in results),'total':len(results)}
