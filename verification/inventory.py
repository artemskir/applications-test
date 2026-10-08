"""Independent XLSX reconciliation: ZIP/XML input, openpyxl output.

Supports ordinary Transitional OOXML fixtures with one selected four-column sheet.
No normalizer imports. Checks physical rows, not worksheet dimensions or row counts alone.
"""
import base64
from collections import Counter
from datetime import datetime, timedelta, time
from hashlib import sha256
from io import BytesIO
import json
import math
import posixpath
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from openpyxl import load_workbook
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
from .reference import calendar, phone, calendar_valid

HEADERS = ('Имя', 'Телефон', 'Дата заявки', 'Источник')
SHEETS = ['Данные', 'Проблемы', 'Оригинал', 'Аудит', 'Проверить']
M = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


def canonical(value):
    return json.dumps(value, sort_keys=True, ensure_ascii=True, allow_nan=False, separators=(',', ':'))


def digest(value):
    return sha256(canonical(value).encode('ascii')).hexdigest()


def same_visible(left, right):
    return left == right and all((type(a) is bool)==(type(b) is bool) and isinstance(a,str)==isinstance(b,str)
                                 for a,b in zip(left,right))


def relationships(z, part):
    relpath = posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part) + '.rels')
    if relpath not in z.namelist():
        return {}
    return {n.attrib['Id']: dict(n.attrib) for n in ET.fromstring(z.read(relpath))}


def resolve(part, target):
    return target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(part), target))


def column(ref):
    letters = ''.join(c for c in ref if c.isalpha())
    number = 0
    for c in letters:
        number = number * 26 + ord(c) - 64
    return number


def decoded_number(value, fmt, epoch):
    # Parser presentation only. Raw XML numeric values and source bytes are checked separately.
    if not isinstance(value, (float, int)) or isinstance(value, bool) or not is_date_format(fmt):
        return value
    day = math.floor(value); millis = round((value-day)*86400000)
    if 0 <= value < 1:
        stamp = datetime(1, 1, 1) + timedelta(milliseconds=millis)
        return {'type':'time', 'value':stamp.time().isoformat()}
    base = datetime(1904, 1, 1) if epoch == 1904 else datetime(1899, 12, 30)
    if epoch == 1900 and 0 < value < 60: day += 1
    stamp = base + timedelta(days=day, milliseconds=millis)
    return {'type':'datetime', 'value':stamp.isoformat()}


def inventory(payload, selected=None):
    with ZipFile(BytesIO(payload)) as z:
        root_rels = ET.fromstring(z.read('_rels/.rels'))
        workbook = next(n.attrib['Target'].lstrip('/') for n in root_rels if n.attrib['Type'].endswith('/officeDocument'))
        root = ET.fromstring(z.read(workbook))
        rels = relationships(z, workbook)
        epoch_node = root.find(M+'workbookPr')
        epoch = 1904 if epoch_node is not None and epoch_node.get('date1904') in ('1', 'true') else 1900
        shared, formats = [], []
        for rel in rels.values():
            part = resolve(workbook, rel['Target'])
            if rel['Type'].endswith('/sharedStrings'):
                shared = [''.join(t.text or '' for t in n.iter(M+'t')) for n in ET.fromstring(z.read(part))]
            if rel['Type'].endswith('/styles'):
                styles = ET.fromstring(z.read(part))
                customs = {int(n.get('numFmtId')): n.get('formatCode') for n in styles.findall(M+'numFmts/'+M+'numFmt')}
                formats = [customs.get(int(n.get('numFmtId')), BUILTIN_FORMATS.get(int(n.get('numFmtId')), 'General'))
                           for n in styles.findall(M+'cellXfs/'+M+'xf')]
        candidates = []
        for index, sheet in enumerate(root.find(M+'sheets'), 1):
            if selected is not None and sheet.get('name') != selected:
                continue
            part = resolve(workbook, rels[sheet.get(R+'id')]['Target'])
            tree = ET.fromstring(z.read(part))
            cells, last = {}, 1
            for row in tree.findall(M+'sheetData/'+M+'row'):
                for node in row.findall(M+'c'):
                    ref = node.get('r')
                    typ = node.get('t', 'n')
                    v, f = node.find(M+'v'), node.find(M+'f')
                    lexical = (v.text or '') if v is not None else None
                    formula = (f.text or '') if f is not None else None
                    if f is not None:
                        value = '=' + (formula or '')
                        kind = 'f'
                    elif typ == 'inlineStr':
                        value = ''.join(t.text or '' for t in node.iter(M+'t')) or None
                        kind = 's'
                    elif typ == 's':
                        value, kind = shared[int(lexical)], 's'
                    elif typ in ('str', 'e'):
                        value, kind = lexical, 's' if typ == 'str' else 'e'
                    elif typ == 'b':
                        value, kind = lexical == '1', 'b'
                    elif lexical is None:
                        value, kind = None, 'n'
                    elif typ == 'd':
                        # Lexical ISO native dates remain a distinct, preserved type.
                        value, kind = lexical, 'd'
                    else:
                        number = float(lexical)
                        value = lexical if not math.isfinite(number) else int(lexical) if lexical.lstrip('+-').isdigit() else number
                        kind = 'n'
                    style = int(node.get('s', '0'))
                    fmt = formats[style] if formats else 'General'
                    cells[ref] = {'coordinate': ref, 'value': value, 'kind': kind, 'ooxml_type': node.get('t', 'n'),
                                  'xml_value': lexical, 'xml_formula': formula, 'style_id': style,
                                  'number_format': fmt, 'epoch': epoch,
                                  'decoded_value': decoded_number(value, fmt, epoch) if kind == 'n' else value,
                                  'raw_xml': ET.tostring(node, encoding='unicode')}
                    if value is not None:
                        last = max(last, int(row.get('r')))
            header = {c['value'].strip(): column(ref) for ref, c in cells.items()
                      if ref.rstrip('0123456789')+'1' == ref and isinstance(c['value'], str)}
            if set(header) != set(HEADERS):
                continue
            sheet_rels = relationships(z, part)
            links = {}
            for link in tree.findall(M+'hyperlinks/'+M+'hyperlink'):
                info = dict(link.attrib)
                if link.get(R+'id') in sheet_rels:
                    info['target'] = sheet_rels[link.get(R+'id')]['Target']
                    info['target_mode'] = sheet_rels[link.get(R+'id')].get('TargetMode')
                links[link.get('ref')] = info
            rows = []
            for number in range(2, last+1):
                values = []
                for h in HEADERS:
                    col = header[h]
                    # Fixtures use A..Z; reject unsupported scope instead of silently guessing.
                    if col > 26:
                        raise ValueError('Inventory scope: column > Z')
                    ref = chr(64+col)+str(number)
                    cell = cells.get(ref, {'coordinate': ref, 'value': None, 'kind': 'n', 'ooxml_type': None,
                                          'xml_value': None, 'xml_formula': None, 'style_id': 0,
                                          'number_format': 'General', 'epoch': epoch}).copy()
                    if ref in links:
                        cell['hyperlink'] = links[ref]
                    values.append(cell)
                rows.append({'row_id': f's{index}:r{number}', 'sheet': sheet.get('name'), 'number': number, 'cells': values})
            candidates.append(rows)
        if len(candidates) != 1:
            raise ValueError('Inventory scope requires one selected input sheet')
        return candidates[0]


def literal(v):
    if isinstance(v, int) and not isinstance(v, bool) and abs(v) > 2**53-1:
        return str(v)
    if isinstance(v, float) and float(format(v, '.16g')) != v:
        return repr(v)
    return v


def reconcile(input_bytes, output_bytes, selected=None):
    expected = inventory(input_bytes, selected)
    book = load_workbook(BytesIO(output_bytes), read_only=True, data_only=False)
    try:
        assert book.sheetnames == SHEETS, 'five sheets/order'
        values = {s: list(book[s].values)[1:] for s in SHEETS}
        audit = values['Аудит']
        chunks = {}
        for row in audit:
            if row[0] in ('VALUE_CHUNK', 'CELL_METADATA_CHUNK'):
                info = json.loads(row[9])
                chunks.setdefault(info['value_id'], []).append((row[10], row[11]))
        def unpack(text):
            info = json.loads(text)
            if isinstance(info, dict) and 'chunk_ref' in info:
                items = sorted(chunks[info['value_id']])
                assert [i for i, _ in items] == list(range(info['chunk_count'])), 'chunk sequence'
                raw = ''.join(p for _, p in items)
                assert sha256(raw.encode('utf-8')).hexdigest() == info['sha256'], 'chunk checksum'
                return json.loads(raw)
            return info
        ids = [r['row_id'] for r in expected]
        originals = values['Оригинал']
        assert [r[0] for r in originals] == ids, 'Original bijection/order'
        for original, wanted in zip(originals, expected):
            assert original[1:3] == (wanted['sheet'], wanted['number']), 'source coordinates'
            assert same_visible(list(original[3:7]), [literal(c['value']) for c in wanted['cells']]), 'original raw values/types'
            meta = unpack(original[7])
            for h, cell in zip(HEADERS, wanted['cells']):
                for k in ('coordinate', 'ooxml_type', 'xml_value', 'xml_formula', 'style_id', 'number_format', 'epoch'):
                    assert meta[h][k] == cell[k], f'metadata {wanted["row_id"]} {h} {k}'
                assert canonical(meta[h]['decoded_value']) == canonical(cell.get('decoded_value', cell['value'])), 'metadata decoded value/types'
                assert meta[h]['raw_xml'] == cell.get('raw_xml'), 'raw XML preservation'
                assert meta[h].get('hyperlink') == cell.get('hyperlink'), 'hyperlink preservation'
        runs = [unpack(r[9]) for r in audit if r[0] == 'RUN']
        assert len(runs) == 1, 'single RUN'
        run = runs[0]
        pieces = sorted((r[10], r[11]) for r in audit if r[0] == 'INPUT_CHUNK')
        assert [i for i, _ in pieces] == list(range(run['chunk_count'])), 'input chunks unique/complete'
        restored = base64.b64decode(''.join(p for _, p in pieces), validate=True)
        assert restored == input_bytes, 'byte-exact restoration'
        assert run['sha256'] == sha256(input_bytes).hexdigest() and run['byte_count'] == len(input_bytes), 'RUN source checksum'
        row_audit = [r for r in audit if r[0] == 'ROW']
        assert Counter(r[1] for r in row_audit) == Counter(ids), 'ROW bijection'
        ledger = {r[1]: r for r in row_audit}
        problems = values['Проблемы']
        problem_ids = [r[0] for r in problems]
        assert len(problem_ids) == len(set(problem_ids)), 'unique problem IDs'
        assert set(problem_ids) == {i for i, r in ledger.items() if unpack(r[9])['codes']}, 'problem membership'
        for p in problems:
            original = originals[ids.index(p[0])]
            assert same_visible(list(p[8::2]), list(original[3:7])), 'problem raw preservation'
        reviews = values['Проверить']
        wanted_review = [r for r in expected if r['row_id'] in problem_ids]
        assert len(reviews) == len(wanted_review), 'review count'
        for r, e in zip(reviews, wanted_review):
            assert same_visible(list(r[:4]), [literal(c['value']) for c in e['cells']]) and r[4:6] == (e['sheet'], e['number']), 'plain review original'
            assert isinstance(r[6], str) and r[6], 'plain explanation'
        admitted = [(i, r) for i, r in ledger.items() if r[3] is not None]
        assert sorted(r[3] for _, r in admitted) == list(range(2, 2+len(values['Данные']))), 'Data bijection'
        assert len({r[1] for r in values['Данные']}) == len(values['Данные']), 'unique Data phone keys'
        for ident, r in admitted:
            assert unpack(r[9])['disposition'] == 'REPRESENTATIVE', 'Data admission representative'
        for row in values['Данные']:
            assert isinstance(row[1],str) and len(row[1])==12 and row[1].startswith('+7') and all(c in '0123456789' for c in row[1][1:]), 'canonical phone grammar'
            assert isinstance(row[2],str) and len(row[2])==10 and row[2][4]==row[2][7]=='-', 'canonical date grammar'
            assert calendar_valid(int(row[2][:4]),int(row[2][5:7]),int(row[2][8:])), 'canonical calendar validity'
        links = [r for r in audit if r[0] == 'DUPLICATE_LINKED']
        duplicate_ids = {i for i, r in ledger.items() if unpack(r[9])['disposition'] == 'DUPLICATE'}
        assert Counter(r[1] for r in links) == Counter(duplicate_ids), 'duplicate links exact'
        for ident, r in ledger.items():
            assert r[4] in ledger, 'representative exists'
            own = unpack(r[9]); representative = unpack(ledger[r[4]][9])
            if own['canonical_phone'] is None:
                assert r[4] == ident and own['disposition'] == 'UNGROUPED', 'missing phone separate'
            else:
                members = [i for i, a in ledger.items() if unpack(a[9])['canonical_phone'] == own['canonical_phone']]
                assert r[4] == min(members, key=lambda i: ids.index(i)), 'first physical representative'
                assert representative['disposition'] == 'REPRESENTATIVE', 'representative disposition'
        # Reconstruct field displays through FIELD_CHANGE; verify raw old values independently.
        changes = [r for r in audit if r[0] == 'FIELD_CHANGE']
        assert len({(r[1], r[5]) for r in changes}) == len(changes), 'unique changes'
        displays = {e['row_id']: [literal(c['value']) for c in e['cells']] for e in expected}
        for change in changes:
            assert change[1] in displays and change[5] in HEADERS, 'change source'
            field = HEADERS.index(change[5]); cell = expected[ids.index(change[1])]['cells'][field]
            old, new = unpack(change[7]), unpack(change[8])
            assert canonical(old) == canonical({'type': cell['kind'], 'value': cell['value']}), 'FIELD_CHANGE old value/types'
            assert change[2] == cell['coordinate'] and change[6], 'change provenance/rule'
            displays[change[1]][field] = new['value']
        resolutions = [r for r in audit if r[0] == 'GROUP_RESOLUTION']
        assert Counter(r[1] for r in resolutions) == Counter(i for i,r in ledger.items()
                    if unpack(r[9])['disposition'] == 'REPRESENTATIVE'), 'group resolutions exact'
        resolved = {r[1]:unpack(r[9]) for r in resolutions}
        field_issues = {(a[1],a[5]) for a in audit if a[0]=='FIELD_ISSUE' and not str(a[6]).startswith('DUPLICATE_')}
        member_ids = {}
        for ident,record in ledger.items():
            member_ids.setdefault(record[4],[]).append(ident)
        for resolution in resolutions:
            ident = resolution[1]; record = ledger[ident]; details = resolved[ident]
            assert resolution[3:7] == (record[3],ident,None,'DEC-31'), 'group provenance'
            assert details['rule_id'] == 'DEC-31' and details['admitted'] == (record[3] is not None), 'group admission'
            assert [f['field'] for f in details['fields']] == list(HEADERS), 'group field order'
            assert details['conflict_codes'] == [c for c in unpack(record[9])['codes'] if c.startswith('DUPLICATE_')], 'group conflicts'
            members = member_ids[ident]
            candidates_by_field = [[i for i in members if (i,HEADERS[field]) not in field_issues] for field in range(3)]
            eligible = all(candidates_by_field) and (ident,HEADERS[3]) not in field_issues
            assert (record[3] is not None) == eligible, 'independent group admission'
            for field, choice in enumerate(details['fields']):
                candidates = candidates_by_field[field] if field<3 else []
                donor = choice['donor_row_id']
                if field==3:
                    assert donor == ident and choice['rule_id']=='KEEP_FIRST_SOURCE', 'first source'
                elif donor is not None:
                    assert candidates and donor == candidates[0] and choice['rule_id']=='FIRST_CONFIRMED_VALUE', 'first confirmed donor'
                else:
                    assert not candidates and (ident,HEADERS[field]) in field_issues, 'unconfirmed field accounted'
                    assert choice['rule_id']=='NO_CONFIRMED_VALUE' and choice['input_cell'] is None, 'missing donor'
                source_id = donor if donor is not None else ident
                assert source_id in members, 'donor belongs to phone group'
                assert choice['value'] == displays[source_id][field], 'group value has source trail'
                if donor is not None:
                    source_cell = expected[ids.index(donor)]['cells'][field]
                    assert choice['input_cell'] == source_cell['coordinate'], 'donor cell'
                    if field<3:
                        assert choice['canonical']==choice['value'], 'confirmed canonical matches display'
            if record[3] is not None:
                assert all(c['donor_row_id'] is not None and c['canonical'] is not None for c in details['fields'][:3]), 'complete group result'
                assert not any(a[0]=='FIELD_ISSUE' and a[1]==ident and a[5]==HEADERS[3] for a in audit), 'source issue blocks Data'
        for ident, r in admitted:
            assert list(values['Данные'][r[3]-2]) == [f['value'] for f in resolved[ident]['fields']], 'Data matches group trail'
        for p in problems:
            assert list(p[9::2]) == displays[p[0]], 'problem displays match complete change trail'
        semantic_fields = 0
        for e in expected:
            for index, parser in ((1,phone),(2,calendar)):
                raw_value = e['cells'][index]['value']
                decision = parser(raw_value)
                if decision.status not in ('REVIEW','MISSING'):
                    semantic_fields += 1
                    if index==1:
                        assert unpack(ledger[e['row_id']][9])['canonical_phone'] == decision.value, 'independent phone meaning'
                    if decision.value is not None:
                        assert displays[e['row_id']][index] == decision.value, 'independent field meaning'
        # Formula-like raw text must remain text on every output sheet.
        for sheet in book:
            for row in sheet:
                assert all(c.data_type != 'f' for c in row), 'output formula execution'
        return {'rows': len(ids), 'input_inventory_sha256': digest(expected), 'output_sha256': sha256(output_bytes).hexdigest(),
                'restored_sha256': sha256(restored).hexdigest(), 'data_rows': len(admitted), 'duplicates': len(duplicate_ids),
                'metadata_cells': 4*len(ids), 'field_changes': len(changes), 'semantic_fields':semantic_fields,'passed': True}
    finally:
        book.close()
